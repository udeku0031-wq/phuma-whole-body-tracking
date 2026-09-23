#!/usr/bin/env python3
"""Build an English KBS paper-material package from frozen local evidence."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import os
from pathlib import Path
from typing import Callable, Iterable, Sequence


ROOT = Path(__file__).resolve().parents[1]
DATE = "2026-09-08"
PACKAGE = ROOT / "reports" / f"kbs_paper_package_{DATE}"
TABLES = PACKAGE / "tables"
FIGURES = PACKAGE / "figures"
DOCX_PATH = PACKAGE / "PHUMA_WBT_KBS_Paper_Materials.docx"
MD_PATH = PACKAGE / "PHUMA_WBT_KBS_Paper_Materials.md"
HIGHLIGHTS_DOCX_PATH = PACKAGE / "Highlights.docx"
HIGHLIGHTS_TXT_PATH = PACKAGE / "Highlights.txt"
CAPTIONS_PATH = PACKAGE / "Figure_Captions.txt"
CHECKLIST_PATH = PACKAGE / "KBS_Submission_Checklist.md"
README_PATH = PACKAGE / "README.md"
BIB_PATH = PACKAGE / "references_seed.bib"
MANIFEST_PATH = PACKAGE / "manifest.json"
STATISTICS_PATH = PACKAGE / "paired_statistics.json"
ALGORITHM_PATH = PACKAGE / "Algorithm_QD_HES.tex"
COVER_LETTER_PATH = PACKAGE / "Cover_Letter_Draft.md"
DECLARATIONS_PATH = PACKAGE / "Declarations_Template.md"

os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".matplotlib-cache"))

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from matplotlib.patches import FancyArrowPatch, Rectangle  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Inches, Pt, RGBColor  # noqa: E402


PAPER_TITLE = (
    "Quality-Gated Diversity-Constrained Hierarchical Error Sampling for "
    "Large-Scale Humanoid Whole-Body Tracking"
)
SHORT_METHOD_NAME = "QD-HES"

ABSTRACT = (
    "Training one humanoid policy on a large and heterogeneous motion library is not only "
    "a control problem but also a data-allocation problem: uniform sampling can overexpose "
    "frequent motion patterns while wasting updates on unreliable trajectory segments. We "
    "present quality-gated diversity-constrained hierarchical error sampling (QD-HES), a "
    "modular experience-sampling framework for whole-body motion tracking. Motions are first "
    "partitioned into fixed-duration segments. Offline trajectory-quality knowledge prevents "
    "rejected segments from becoming episode starts, unsupervised motion clusters allocate "
    "diversity budgets, and online tracking errors prioritize motions and segments within each "
    "cluster. The resulting distribution factorizes over cluster, motion, and segment levels. "
    "We implement the framework for a 29-DoF Unitree G1 policy trained with proximal policy "
    "optimization in Isaac Lab. The study uses 76,086 converted PHUMA trajectories, with 6,000 "
    "training motions, 7,636 validation motions, and 7,592 held-out test motions. Under a 34k "
    "iteration budget, the diversity-only variant obtains 0.8993 macro validation success. "
    "Validation-based budget scaling selects QD-HES at 59k iterations. On the final test set, "
    "QD-HES reaches 0.9231 micro success, 0.9079 macro success, and 0.9554 completion, while "
    "reducing failures from 731 to 584 relative to the strongest diversity-only baseline. "
    "Body and joint-position errors decrease by 5.27% and 7.50%, respectively. Dense validation "
    "also reveals non-monotonic degradation after the optimum, demonstrating that structured "
    "sampling and validation-based checkpoint selection are both important at scale."
)

HIGHLIGHTS = [
    "A hierarchical sampler allocates PPO data across clusters, motions, and segments.",
    "A trajectory-quality gate blocks unreliable segments as episode reset starts.",
    "QD-HES reaches 92.31% Test micro success on 7,592 held-out motions.",
    "QD-HES reduces Test failures by 147 versus the strong D-only baseline.",
    "Dense validation reveals a non-monotonic optimum near 59k PPO iterations.",
]

KEYWORDS = [
    "Humanoid robots",
    "Whole-body motion tracking",
    "Reinforcement learning",
    "Adaptive sampling",
    "Trajectory quality",
    "Motion diversity",
]


FORMAL_SPECS = [
    ("M0 / Uniform", "34k", "model_33999.pt", "evaluations/formal_v1/M0_seed42/validation_full/model_33999"),
    ("M1 / Quality only", "34k", "model_33999.pt", "evaluations/final_minimal/Qonly_seed42/validation_full/model_33999"),
    ("GlobalRaw", "34k", "model_33500.pt", "evaluations/final_minimal/GlobalRaw_seed42/validation_full/model_33500"),
    ("GlobalRaw-Q", "34k", "model_33999.pt", "evaluations/final_minimal/GlobalRawQ_seed42/validation_full/model_33999"),
    ("M4 / Hierarchical raw", "34k", "model_33999.pt", "evaluations/formal_v1/M4_seed42/validation_full/model_33999"),
    ("M5 / Learning gap", "34k", "model_25000.pt", "evaluations/formal_v1/M5_seed42/validation_full/model_25000"),
    ("M6 / Quality + gap", "34k", "model_33500.pt", "evaluations/formal_v1/M6_seed42_rerun1/validation_full/model_33500"),
    ("M7 / Gap + diversity", "34k", "model_33999.pt", "evaluations/formal_v1/M7_seed42/validation_full/model_33999"),
    ("D-only", "34k", "model_33500.pt", "evaluations/final_minimal/Donly_seed42/validation_full/model_33500"),
    ("M7-Raw", "34k", "model_33999.pt", "evaluations/formal_v1/M7Raw_seed42/validation_full/model_33999"),
]

BUDGET_SPECS = [
    ("D-only", 33500, "34k", "model_33500.pt", "evaluations/final_minimal/Donly_seed42/validation_full/model_33500"),
    ("D-only", 49999, "50k", "model_49999.pt", "evaluations/postfreeze_Donly_extend50k_from33999_seed42_v1/validation_full/model_49999"),
    ("D-only", 54000, "54k selected", "model_54000.pt", "evaluations/postfreeze_Donly_extend59k_from49999_seed42_v1/validation_full/model_54000"),
    ("M7-Raw", 33999, "34k", "model_33999.pt", "evaluations/formal_v1/M7Raw_seed42/validation_full/model_33999"),
    ("M7-Raw", 45000, "45k", "model_45000.pt", "evaluations/postfreeze_M7Raw_extend50k_from33999_seed42_v1/validation_full_sweep/model_45000"),
    ("M7-Raw", 49999, "50k", "model_49999.pt", "evaluations/postfreeze_M7Raw_extend50k_from33999_seed42_v1/validation_full_sweep/model_49999"),
    ("M7-Raw", 59000, "59k selected", "model_59000.pt", "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_full/model_59000"),
    ("M7-Raw", 60000, "60k", "model_60000.pt", "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_full_sweep/model_60000"),
    ("M7-Raw", 69999, "70k", "model_69999.pt", "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_full_sweep/model_69999"),
]

TEST_SPECS = [
    (
        "GlobalRaw",
        "34k",
        "model_33500.pt",
        "outputs/final_test_v2_freeze_M7Raw59_Donly54_GlobalRaw34/GlobalRaw/test_final/model_33500",
    ),
    (
        "D-only",
        "54k",
        "model_54000.pt",
        "outputs/final_test_v2_freeze_M7Raw59_Donly54_GlobalRaw34/Donly/test_final/model_54000",
    ),
    (
        "M7-Raw (QD-HES)",
        "59k",
        "model_59000.pt",
        "outputs/final_test_v2_freeze_M7Raw59_Donly54_GlobalRaw34/M7Raw/test_final/model_59000",
    ),
]

MODULE_ROWS = [
    [
        "Data conversion and multi-motion library",
        "Converts PHUMA G1 .npy trajectories into the WBT .npz schema and supports single files, directories, and ordered manifests.",
        "Creates a consistent 29-DoF/30-body reference library and assigns an independent motion and start frame to each environment.",
    ],
    [
        "Stage 0: segment infrastructure",
        "Builds fixed-duration segment indices, local/global ID mappings, shared counters, probability checks, and checkpoint state.",
        "Provides one reproducible interface for every sampler while preserving the legacy uniform random-number path.",
    ],
    [
        "Module 1: trajectory-quality gate",
        "Audits finite values, joint limits, velocity consistency, discontinuities, ground penetration, and foot sliding; labels each segment pass, borderline, or reject.",
        "Prevents rejected segments from becoming episode starts without deleting complete motions or using policy performance as a quality label.",
    ],
    [
        "Module 2: intrinsic difficulty",
        "Extracts 28 policy-independent kinematic/contact features, applies train-only robust scaling and empirical-CDF mapping, and assigns ten balanced bins.",
        "Supplies a policy-independent calibration reference for learning-gap variants; it does not directly remove or sample data.",
    ],
    [
        "Module 3: online error and learning state",
        "Attributes body, joint, orientation, completion, and termination signals to the active motion/segment and maintains exponential moving averages.",
        "Prioritizes under-learned motions and segments using raw error or difficulty-calibrated learning gaps with a uniform exploration floor.",
    ],
    [
        "Module 4: diversity constraint",
        "Aggregates 17 segment features into 30 motion features and fits train-only K-means++ clusters (K=8) for cluster-level budgets.",
        "Stops a narrow family of high-error motions from monopolizing updates and yields cluster-motion-segment sampling.",
    ],
    [
        "PPO tracking and deterministic evaluation",
        "Trains a shared 29-dimensional joint-position policy in Isaac Lab/RSL-RL and evaluates every motion from frame zero once.",
        "Keeps network, reward, termination, and evaluator settings fixed across samplers and exports motion-, category-, and corpus-level metrics.",
    ],
]

METHOD_MATRIX = [
    ["M0", "Uniform", "Uniform", "Off", "Off", "Off", "Full validation"],
    ["M1", "Uniform", "Uniform", "On", "Off", "Off", "Full validation"],
    ["M2", "Raw error", "Uniform", "Off", "Off", "Off", "Pilot only; full validation missing"],
    ["M3", "Uniform", "Raw error", "Off", "Off", "Off", "Pilot only; full validation missing"],
    ["M4", "Raw error", "Raw error", "Off", "Off", "Off", "Full validation"],
    ["M5", "Learning gap", "Relative gap", "Off", "On", "Off", "Full validation"],
    ["M6", "Learning gap", "Relative gap", "On", "On", "Off", "Full validation"],
    ["M7", "Learning gap", "Relative gap", "On", "On", "On", "Full validation"],
    ["GlobalRaw", "Flat", "Global raw error", "Off", "Off", "Off", "Full validation"],
    ["GlobalRaw-Q", "Flat", "Global raw error", "On", "Off", "Off", "Full validation"],
    ["D-only", "Raw error", "Raw error", "Off", "Off", "On", "34k/50k/54k validation"],
    ["M7-Raw / QD-HES", "Raw error", "Raw error", "On", "Off", "On", "34k-70k validation; 59k final Test"],
]

SPLIT_ROWS = [
    ["Training", "random6000_seed42.txt", 6000, "51f592792f412c5d5e31caf3621162a3b2ef356e05f98dbce43153d08a59e82d"],
    ["Validation", "validation_full.txt", 7636, "fb70e8b444b20f828226a2dea782d6cf24c2db3900a11e8c01f567f24b0bd42d"],
    ["Validation probe", "validation_probe500_seed42.txt", 500, "1e44947a99334b104d826776cf1f04a313ad96a08e43de1c21e35569076c8a4d"],
    ["Test", "test.txt", 7592, "febffa69e70ed3df966a3d2092b7c0956142babe7de53dc30dc46a8668de570b"],
]

PPO_ROWS = [
    ["Parallel environments", "3,072"],
    ["Rollout", "24 steps/environment/iteration = 73,728 transitions/iteration"],
    ["Simulation/control", "dt=0.005 s, decimation=4, control period=0.02 s (50 Hz)"],
    ["Training episode", "10 s"],
    ["Evaluation horizon", "60 s"],
    ["Actor/Critic", "[512, 256, 128], ELU; actor obs=160, critic obs=286, actions=29"],
    ["PPO", "clip=0.2, epochs=5, mini-batches=4, lr=1e-3 adaptive, gamma=0.99, lambda=0.95"],
    ["Regularization", "entropy coefficient=0.005, max gradient norm=1.0, desired KL=0.01"],
    ["Checkpoint rule", "Validation macro first; epsilon=0.002; then micro, completion, body error, earlier checkpoint"],
    ["Evaluation", "seed=42, deterministic, randomization off, Fabric off; each motion starts at frame 0 once"],
]

REWARD_ROWS = [
    ["Global anchor position", "+0.5", "Exponential error, std=0.3"],
    ["Global anchor orientation", "+0.5", "Exponential error, std=0.4"],
    ["Relative body position", "+1.0", "Exponential error, std=0.3"],
    ["Relative body orientation", "+1.0", "Exponential error, std=0.4"],
    ["Body linear velocity", "+1.0", "Exponential error, std=1.0"],
    ["Body angular velocity", "+1.0", "Exponential error, std=3.14"],
    ["Action-rate L2", "-0.1", "Action smoothness"],
    ["Joint limit", "-10.0", "Joint-limit violation"],
    ["Undesired contacts", "-0.1", "Non-wrist/non-ankle contacts"],
]

METRIC_ROWS = [
    ["Micro success (higher)", "Successful motions divided by all motions; dominated by large categories."],
    ["Macro success (higher)", "Unweighted mean of the 17 category success rates; the primary checkpoint-selection metric."],
    ["Completion (higher)", "Mean completed_frames/num_frames over motions, clamped to [0,1]."],
    ["Body error (lower)", "Mean Euclidean tracked-body position error after yaw/root alignment, in metres."],
    ["Joint L2 (lower)", "Mean L2 norm of the 29-joint position-error vector, in radians."],
    ["Joint RMS (lower)", "Joint L2 divided by sqrt(29), in radians."],
    ["Failures (lower)", "Motions that terminate before the final reference frame."],
]


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_json(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def read_csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def summarize_per_motion(path: Path) -> dict[str, object]:
    rows = read_csv_rows(path)
    if not rows:
        raise RuntimeError(f"No rows in {path}")
    by_category: dict[str, list[int]] = {}
    for row in rows:
        by_category.setdefault(row["category"], []).append(int(row["success"]))
    num_success = sum(int(row["success"]) for row in rows)
    n = len(rows)
    return {
        "num_motions": n,
        "num_success": num_success,
        "num_failure": n - num_success,
        "micro_success_rate": num_success / n,
        "macro_success_rate": sum(sum(values) / len(values) for values in by_category.values()) / len(by_category),
        "mean_completion_ratio": sum(float(row["completion_ratio"]) for row in rows) / n,
        "mean_body_position_error_m": sum(float(row["body_position_error_m"]) for row in rows) / n,
        "mean_joint_position_error_l2_rad": sum(float(row["joint_position_error_l2_rad"]) for row in rows) / n,
        "mean_joint_position_error_rms_rad": sum(float(row["joint_position_error_rms_rad"]) for row in rows) / n,
    }


def load_result(directory: str) -> dict[str, object]:
    base = ROOT / directory
    summary_path = base / "summary.json"
    payload = read_json(summary_path) if summary_path.exists() else summarize_per_motion(base / "per_motion.csv")
    per_motion = base / "per_motion.csv"
    if per_motion.exists():
        recomputed = summarize_per_motion(per_motion)
        keys = [
            "micro_success_rate",
            "macro_success_rate",
            "mean_completion_ratio",
            "mean_body_position_error_m",
            "mean_joint_position_error_l2_rad",
            "mean_joint_position_error_rms_rad",
        ]
        for key in keys:
            if abs(float(payload[key]) - float(recomputed[key])) > 7.5e-7:
                raise RuntimeError(f"Metric mismatch for {base}:{key}")
        if int(payload["num_failure"]) != int(recomputed["num_failure"]):
            raise RuntimeError(f"Failure-count mismatch for {base}")
    return payload


def metric_record(method: str, budget: str, checkpoint: str, directory: str) -> dict[str, object]:
    payload = load_result(directory)
    return {
        "method": method,
        "budget": budget,
        "checkpoint": checkpoint,
        "micro": float(payload["micro_success_rate"]),
        "macro": float(payload["macro_success_rate"]),
        "completion": float(payload["mean_completion_ratio"]),
        "body": float(payload["mean_body_position_error_m"]),
        "joint": float(payload["mean_joint_position_error_l2_rad"]),
        "rms": float(payload["mean_joint_position_error_rms_rad"]),
        "successes": int(payload["num_success"]),
        "failures": int(payload["num_failure"]),
        "directory": directory,
        "payload": payload,
    }


def build_result_sets() -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    formal = [metric_record(*spec) for spec in FORMAL_SPECS]
    budget: list[dict[str, object]] = []
    for method, iteration, budget_label, checkpoint, directory in BUDGET_SPECS:
        row = metric_record(method, budget_label, checkpoint, directory)
        row["iteration"] = iteration
        budget.append(row)
    test = [metric_record(*spec) for spec in TEST_SPECS]
    return formal, budget, test


def paired_arrays(a_path: Path, b_path: Path) -> tuple[dict[str, np.ndarray], np.ndarray, np.ndarray]:
    a_rows = read_csv_rows(a_path)
    b_rows = read_csv_rows(b_path)
    a_by_path = {row["motion_path"]: row for row in a_rows}
    b_by_path = {row["motion_path"]: row for row in b_rows}
    if len(a_by_path) != len(a_rows) or len(b_by_path) != len(b_rows):
        raise RuntimeError("Duplicate motion_path found in paired evaluation results")
    if set(a_by_path) != set(b_by_path):
        raise RuntimeError("Paired evaluation results contain different motion_path sets")
    ordered_paths = [row["motion_path"] for row in a_rows]
    aligned_a = [a_by_path[path] for path in ordered_paths]
    aligned_b = [b_by_path[path] for path in ordered_paths]
    fields = {
        "micro": "success",
        "completion": "completion_ratio",
        "body": "body_position_error_m",
        "joint": "joint_position_error_l2_rad",
    }
    diffs = {
        key: np.asarray([float(a[field]) - float(b[field]) for a, b in zip(aligned_a, aligned_b)], dtype=np.float64)
        for key, field in fields.items()
    }
    categories = np.asarray([row["category"] for row in aligned_a], dtype=object)
    groups = np.asarray([row["source_group"] for row in aligned_a], dtype=object)
    return diffs, categories, groups


def paired_bootstrap(
    a_path: Path,
    b_path: Path,
    *,
    seed: int = 42,
    resamples: int = 10000,
    batch_size: int = 128,
) -> dict[str, object]:
    diffs, categories, groups = paired_arrays(a_path, b_path)
    n = len(categories)
    rng = np.random.default_rng(seed)
    draws = {key: np.empty(resamples, dtype=np.float64) for key in diffs}
    for start in range(0, resamples, batch_size):
        stop = min(start + batch_size, resamples)
        index = rng.integers(0, n, size=(stop - start, n))
        for key, values in diffs.items():
            draws[key][start:stop] = values[index].mean(axis=1)

    category_names = sorted(set(categories.tolist()))
    macro_draws = np.empty(resamples, dtype=np.float64)
    for start in range(0, resamples, batch_size):
        stop = min(start + batch_size, resamples)
        batch = np.zeros(stop - start, dtype=np.float64)
        for category in category_names:
            values = diffs["micro"][categories == category]
            index = rng.integers(0, len(values), size=(stop - start, len(values)))
            batch += values[index].mean(axis=1)
        macro_draws[start:stop] = batch / len(category_names)

    result: dict[str, object] = {
        "seed": seed,
        "resamples": resamples,
        "num_motions": n,
        "num_source_groups": len(set(groups.tolist())),
        "num_categories": len(category_names),
    }
    for key, values in diffs.items():
        result[key] = {
            "estimate": float(values.mean()),
            "ci_low": float(np.quantile(draws[key], 0.025)),
            "ci_high": float(np.quantile(draws[key], 0.975)),
        }
    observed_macro = float(np.mean([diffs["micro"][categories == c].mean() for c in category_names]))
    result["macro"] = {
        "estimate": observed_macro,
        "ci_low": float(np.quantile(macro_draws, 0.025)),
        "ci_high": float(np.quantile(macro_draws, 0.975)),
    }
    return result


def grouped_paired_bootstrap(
    a_path: Path,
    b_path: Path,
    *,
    seed: int = 42,
    resamples: int = 10000,
    batch_size: int = 128,
) -> dict[str, object]:
    """Paired cluster bootstrap that resamples source groups, not motion chunks."""
    diffs, categories, groups = paired_arrays(a_path, b_path)
    group_names, group_inverse = np.unique(groups, return_inverse=True)
    group_counts = np.bincount(group_inverse).astype(np.float64)
    group_sums = {
        key: np.bincount(group_inverse, weights=values, minlength=len(group_names))
        for key, values in diffs.items()
    }
    rng = np.random.default_rng(seed)
    draws = {key: np.empty(resamples, dtype=np.float64) for key in diffs}
    for start in range(0, resamples, batch_size):
        stop = min(start + batch_size, resamples)
        index = rng.integers(0, len(group_names), size=(stop - start, len(group_names)))
        denominators = group_counts[index].sum(axis=1)
        for key, values in group_sums.items():
            draws[key][start:stop] = values[index].sum(axis=1) / denominators

    category_names = sorted(set(categories.tolist()))
    macro_draws = np.zeros(resamples, dtype=np.float64)
    for category in category_names:
        mask = categories == category
        category_groups, inverse = np.unique(groups[mask], return_inverse=True)
        counts = np.bincount(inverse).astype(np.float64)
        sums = np.bincount(inverse, weights=diffs["micro"][mask], minlength=len(category_groups))
        for start in range(0, resamples, batch_size):
            stop = min(start + batch_size, resamples)
            index = rng.integers(0, len(category_groups), size=(stop - start, len(category_groups)))
            macro_draws[start:stop] += sums[index].sum(axis=1) / counts[index].sum(axis=1)
    macro_draws /= len(category_names)

    result: dict[str, object] = {
        "seed": seed,
        "resamples": resamples,
        "num_motions": len(groups),
        "num_source_groups": len(group_names),
        "num_categories": len(category_names),
        "resampling_unit": "source_group",
    }
    for key, values in diffs.items():
        result[key] = {
            "estimate": float(values.mean()),
            "ci_low": float(np.quantile(draws[key], 0.025)),
            "ci_high": float(np.quantile(draws[key], 0.975)),
        }
    observed_macro = float(np.mean([diffs["micro"][categories == category].mean() for category in category_names]))
    result["macro"] = {
        "estimate": observed_macro,
        "ci_low": float(np.quantile(macro_draws, 0.025)),
        "ci_high": float(np.quantile(macro_draws, 0.975)),
    }
    return result


def exact_paired_outcomes(a_path: Path, b_path: Path) -> dict[str, object]:
    """Count paired wins/losses and compute an exact two-sided McNemar p-value."""
    a_rows = read_csv_rows(a_path)
    b_rows = read_csv_rows(b_path)
    a_by_path = {row["motion_path"]: row for row in a_rows}
    b_by_path = {row["motion_path"]: row for row in b_rows}
    if len(a_by_path) != len(a_rows) or len(b_by_path) != len(b_rows) or set(a_by_path) != set(b_by_path):
        raise RuntimeError("Paired outcome results do not have the same unique motion_path set")
    pairs = [(row, b_by_path[row["motion_path"]]) for row in a_rows]
    a_only = sum(int(a["success"]) == 1 and int(b["success"]) == 0 for a, b in pairs)
    b_only = sum(int(a["success"]) == 0 and int(b["success"]) == 1 for a, b in pairs)
    both = sum(int(a["success"]) == 1 and int(b["success"]) == 1 for a, b in pairs)
    neither = len(pairs) - a_only - b_only - both
    discordant = a_only + b_only
    if discordant:
        tail = sum(math.comb(discordant, index) for index in range(min(a_only, b_only) + 1))
        p_value = min(1.0, 2.0 * tail / (2**discordant))
    else:
        p_value = 1.0
    return {
        "a_only": a_only,
        "b_only": b_only,
        "both": both,
        "neither": neither,
        "discordant": discordant,
        "mcnemar_exact_p": p_value,
    }


def category_comparison(test: Sequence[dict[str, object]]) -> list[dict[str, object]]:
    by_method: dict[str, dict[str, dict[str, str]]] = {}
    for row in test:
        path = ROOT / str(row["directory"]) / "category_summary.csv"
        by_method[str(row["method"])] = {item["category"]: item for item in read_csv_rows(path)}
    names = sorted(next(iter(by_method.values())))
    output: list[dict[str, object]] = []
    for name in names:
        g = by_method["GlobalRaw"][name]
        d = by_method["D-only"][name]
        m = by_method["M7-Raw (QD-HES)"][name]
        output.append(
            {
                "category": name,
                "n": int(m["num_motions"]),
                "global": float(g["success_rate"]),
                "donly": float(d["success_rate"]),
                "m7raw": float(m["success_rate"]),
                "delta_d": float(m["success_rate"]) - float(d["success_rate"]),
                "delta_global": float(m["success_rate"]) - float(g["success_rate"]),
                "d_failures": int(d["num_failure"]),
                "m_failures": int(m["num_failure"]),
            }
        )
    return output


def dense_probe_rows() -> list[dict[str, object]]:
    base = ROOT / "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1"
    paths = [
        base / "validation_probe500_dense_sweep/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_50500_55000/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_55500_60000/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_60500_65000/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_65500_69999/checkpoint_comparison.csv",
    ]
    rows: dict[int, dict[str, object]] = {}
    for path in paths:
        if not path.exists():
            continue
        for row in read_csv_rows(path):
            iteration = int(row["iteration"])
            rows[iteration] = {
                "iteration": iteration,
                "checkpoint": row["checkpoint"],
                "micro": float(row["micro_success_rate"]),
                "macro": float(row["macro_success_rate"]),
                "completion": float(row["mean_completion_ratio"]),
                "body": float(row["mean_body_position_error_m"]),
                "joint": float(row["mean_joint_position_error_l2_rad"]),
                "failures": int(row["num_failure"] if "num_failure" in row else row["num_failures"]),
            }
    return [rows[key] for key in sorted(rows)]


def write_csv(path: Path, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def fmt(value: float, digits: int = 6) -> str:
    return f"{value:.{digits}f}"


def pp(value: float) -> str:
    return f"{value * 100:+.3f} pp"


def markdown_table(headers: Sequence[str], rows: Iterable[Sequence[object]]) -> str:
    output = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    output.extend("| " + " | ".join(str(value) for value in row) + " |" for row in rows)
    return "\n".join(output)


def load_module_statistics() -> dict[str, object]:
    quality = read_json(ROOT / "outputs/module1_quality_random6000_seed42_original_v1/quality_summary.json")
    difficulty = read_json(ROOT / "outputs/module2_difficulty_random6000_seed42_v1/difficulty_summary.json")
    clusters = read_json(ROOT / "outputs/module4_clusters_random6000_seed42_v1/cluster_summary.json")
    return {
        "quality": quality,
        "difficulty": difficulty,
        "clusters": clusters,
        "quality_metadata_sha256": sha256_file(
            ROOT / "outputs/module1_quality_random6000_seed42_original_v1/segment_quality_metadata.npz"
        ),
        "difficulty_metadata_sha256": sha256_file(
            ROOT / "outputs/module2_difficulty_random6000_seed42_v1/segment_difficulty_metadata.npz"
        ),
        "cluster_metadata_sha256": sha256_file(
            ROOT / "outputs/module4_clusters_random6000_seed42_v1/motion_cluster_metadata.npz"
        ),
    }


COLORS = {
    "blue": "#0072B2",
    "orange": "#D55E00",
    "green": "#009E73",
    "yellow": "#E69F00",
    "purple": "#CC79A7",
    "sky": "#56B4E9",
    "black": "#202124",
    "gray": "#6B7280",
    "light": "#E5E7EB",
}


FIGURE_CAPTIONS = {
    "Fig_1_Overall_Framework": (
        "Fig. 1. Overview of the PHUMA-WBT pipeline and the proposed quality-gated "
        "diversity-constrained hierarchical error sampler. Offline quality and motion-cluster "
        "knowledge constrain the training distribution, while online tracking error allocates "
        "budget within each cluster."
    ),
    "Fig_2_Training_Metadata": (
        "Fig. 2. Train-only metadata used by the sampling framework. (a) Segment-quality audit "
        "outcomes for the 21,575 one-second training segments. (b) Sizes of the eight unsupervised "
        "motion clusters used to allocate diversity budgets."
    ),
    "Fig_3_Validation_Ablation_34k": (
        "Fig. 3. Full-validation success rates under the common 34k-iteration budget. D-only "
        "achieves the highest macro success, whereas M7-Raw has marginally higher micro success. "
        "M2 and M3 are omitted because full-validation runs are not available."
    ),
    "Fig_4_Validation_Budget_Scaling": (
        "Fig. 4. Validation performance as the training budget increases. M7-Raw improves through "
        "59k iterations and then degrades, while the longer-budget gain of D-only is limited. "
        "Points are evaluated checkpoints rather than interpolated training measurements."
    ),
    "Fig_5_Final_Test_Comparison": (
        "Fig. 5. Final Test comparison on 7,592 held-out motions. (a) Micro success, macro success, "
        "and completion; (b) body-position error; (c) joint-position L2 error; and (d) failed "
        "motions. The M7-Raw/QD-HES checkpoint was selected on Validation before this evaluation."
    ),
    "Fig_6_Final_Test_Category_Deltas": (
        "Fig. 6. Category-wise Test success-rate difference between M7-Raw/QD-HES 59k and D-only "
        "54k. Positive values favour M7-Raw. Category sample sizes are shown in parentheses; gains "
        "for very small categories should be interpreted cautiously."
    ),
    "Fig_S1_Dense_Validation_Probe": (
        "Fig. S1. Dense 500-motion validation probe for M7-Raw between 50k and 70k iterations. "
        "The near-zero-success region from approximately 62k to 67k demonstrates severe "
        "non-monotonic training instability; tracking errors in that region are not meaningful "
        "in isolation because episodes terminate almost immediately."
    ),
    "Graphical_Abstract": (
        "Graphical abstract. Structured quality, diversity, and online-error knowledge is composed "
        "into a cluster-motion-segment sampling distribution for large-scale humanoid tracking."
    ),
}


def configure_plots() -> None:
    matplotlib.rcParams.update(
        {
            "font.family": "serif",
            "font.serif": ["Times New Roman", "Times", "Liberation Serif", "DejaVu Serif"],
            "font.size": 8.0,
            "axes.titlesize": 9.0,
            "axes.labelsize": 8.5,
            "xtick.labelsize": 7.5,
            "ytick.labelsize": 7.5,
            "legend.fontsize": 7.5,
            "axes.linewidth": 0.7,
            "lines.linewidth": 1.4,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "savefig.facecolor": "white",
            "figure.facecolor": "white",
        }
    )


def save_artwork(fig: plt.Figure, stem: str, *, tight: bool = True) -> dict[str, Path]:
    FIGURES.mkdir(parents=True, exist_ok=True)
    pdf_path = FIGURES / f"{stem}.pdf"
    png_path = FIGURES / f"{stem}_preview.png"
    save_kwargs = {"bbox_inches": "tight", "pad_inches": 0.04} if tight else {}
    fig.savefig(pdf_path, **save_kwargs)
    fig.savefig(png_path, dpi=300, **save_kwargs)
    plt.close(fig)
    return {"pdf": pdf_path, "preview": png_path}


def add_arrow(ax: plt.Axes, start: tuple[float, float], end: tuple[float, float], color: str = "#4B5563") -> None:
    ax.add_patch(
        FancyArrowPatch(
            start,
            end,
            arrowstyle="-|>",
            mutation_scale=10,
            linewidth=1.0,
            color=color,
            shrinkA=2,
            shrinkB=2,
        )
    )


def plot_framework() -> dict[str, Path]:
    fig, ax = plt.subplots(figsize=(7.48, 4.25))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 7)
    ax.axis("off")

    def box(x: float, y: float, w: float, h: float, label: str, face: str, edge: str = "#374151") -> None:
        ax.add_patch(Rectangle((x, y), w, h, facecolor=face, edgecolor=edge, linewidth=1.0))
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=7.7, color=COLORS["black"])

    box(0.15, 5.25, 1.75, 1.05, "PHUMA corpus\n76,086 motions", "#EAF2F8")
    box(2.15, 5.25, 1.75, 1.05, "G1/WBT conversion\n29 DoF, 30 bodies", "#E8F5E9")
    box(4.15, 5.25, 1.75, 1.05, "Frozen manifests\nTrain / Val / Test", "#FFF7E6")
    box(6.15, 5.25, 1.75, 1.05, "One-second\nsegment index", "#FCECE8")
    box(9.75, 5.25, 2.05, 1.05, "Isaac Lab + PPO\nshared tracking policy", "#EDE9F7")

    add_arrow(ax, (1.90, 5.78), (2.15, 5.78))
    add_arrow(ax, (3.90, 5.78), (4.15, 5.78))
    add_arrow(ax, (5.90, 5.78), (6.15, 5.78))
    add_arrow(ax, (7.90, 5.78), (9.75, 5.78))

    ax.add_patch(Rectangle((3.05, 1.20), 6.80, 3.15, facecolor="#FAFAFA", edgecolor=COLORS["blue"], linewidth=1.25))
    ax.text(6.45, 4.07, "Hierarchical training-data allocation", ha="center", va="center", fontsize=9, weight="bold")
    box(3.35, 2.75, 1.40, 0.78, "Quality gate\npass/borderline/reject", "#E8F5E9", COLORS["green"])
    box(5.05, 2.75, 1.40, 0.78, "Diversity knowledge\nK=8 clusters", "#EAF2F8", COLORS["blue"])
    box(6.75, 2.75, 1.40, 0.78, "Online competence\nraw tracking error", "#FFF7E6", COLORS["yellow"])
    box(8.45, 2.75, 1.10, 0.78, "Uniform floor\nand caps", "#F3F4F6", COLORS["gray"])
    box(4.35, 1.55, 4.20, 0.72, "P(c,m,s) = Pdiv(c) Perr(m|c) Perr(s|m)", "#FFFFFF", COLORS["orange"])
    add_arrow(ax, (4.05, 2.75), (4.95, 2.27))
    add_arrow(ax, (5.75, 2.75), (5.75, 2.27))
    add_arrow(ax, (7.45, 2.75), (7.45, 2.27))
    add_arrow(ax, (9.00, 2.75), (8.15, 2.27))
    add_arrow(ax, (7.03, 4.35), (7.03, 5.25))
    add_arrow(ax, (8.55, 1.91), (10.35, 1.91))
    box(10.35, 1.35, 1.45, 1.12, "Deterministic\nevaluation\nVal -> frozen Test", "#FCECE8", COLORS["orange"])
    ax.text(3.10, 0.58, "D-only: diversity + hierarchical raw error", fontsize=7.4, color=COLORS["blue"])
    ax.text(7.00, 0.58, "QD-HES / M7-Raw: quality + diversity + hierarchical raw error", fontsize=7.4, color=COLORS["orange"])
    fig.tight_layout(pad=0.1)
    return save_artwork(fig, "Fig_1_Overall_Framework")


def plot_training_metadata(stats: dict[str, object]) -> dict[str, Path]:
    quality = stats["quality"]
    clusters = stats["clusters"]
    q_counts = [int(quality["pass_count"]), int(quality["borderline_count"]), int(quality["reject_count"])]
    q_labels = ["Pass", "Borderline", "Reject"]
    q_colors = [COLORS["green"], COLORS["yellow"], COLORS["orange"]]
    cluster_sizes = [int(value) for value in clusters["cluster_sizes"]]

    fig, axes = plt.subplots(1, 2, figsize=(7.48, 2.75))
    ax = axes[0]
    bars = ax.bar(q_labels, q_counts, color=q_colors, edgecolor="white", linewidth=0.5)
    ax.set_ylabel("Training segments")
    ax.set_title("(a) Segment-quality audit")
    ax.grid(axis="y", alpha=0.2)
    total = sum(q_counts)
    for bar, count in zip(bars, q_counts):
        ax.text(bar.get_x() + bar.get_width() / 2, count + total * 0.015, f"{count:,}\n({count/total*100:.2f}%)", ha="center", va="bottom", fontsize=7)

    ax = axes[1]
    x = np.arange(len(cluster_sizes))
    bars = ax.bar(x, cluster_sizes, color=COLORS["blue"], edgecolor="white", linewidth=0.5)
    ax.set_xticks(x, [str(index) for index in x])
    ax.set_xlabel("Cluster ID")
    ax.set_ylabel("Training motions")
    ax.set_title("(b) Motion-cluster sizes")
    ax.grid(axis="y", alpha=0.2)
    for bar, count in zip(bars, cluster_sizes):
        ax.text(bar.get_x() + bar.get_width() / 2, count + 35, f"{count}", ha="center", va="bottom", fontsize=6.6, rotation=0)
    fig.tight_layout(pad=0.5)
    return save_artwork(fig, "Fig_2_Training_Metadata")


def plot_formal_ablation(formal: Sequence[dict[str, object]]) -> dict[str, Path]:
    ordered = sorted(formal, key=lambda row: float(row["macro"]))
    y = np.arange(len(ordered))
    height = 0.34
    fig, ax = plt.subplots(figsize=(7.48, 4.10))
    ax.barh(y - height / 2, [float(row["micro"]) * 100 for row in ordered], height, color=COLORS["blue"], label="Micro")
    ax.barh(y + height / 2, [float(row["macro"]) * 100 for row in ordered], height, color=COLORS["orange"], label="Macro")
    ax.set_yticks(y, [str(row["method"]) for row in ordered])
    ax.set_xlim(84.5, 90.6)
    ax.set_xlabel("Success rate (%)")
    ax.grid(axis="x", alpha=0.2)
    ax.legend(frameon=False, ncol=2, loc="lower right")
    for index, row in enumerate(ordered):
        ax.text(float(row["micro"]) * 100 + 0.05, index - height / 2, f"{float(row['micro'])*100:.2f}", va="center", fontsize=6.4)
        ax.text(float(row["macro"]) * 100 + 0.05, index + height / 2, f"{float(row['macro'])*100:.2f}", va="center", fontsize=6.4)
    fig.tight_layout(pad=0.5)
    return save_artwork(fig, "Fig_3_Validation_Ablation_34k")


def plot_budget_scaling(budget: Sequence[dict[str, object]]) -> dict[str, Path]:
    fig, axes = plt.subplots(2, 2, figsize=(7.48, 5.15), sharex=True)
    panels = [
        ("macro", "(a) Macro success", 100.0, "%"),
        ("micro", "(b) Micro success", 100.0, "%"),
        ("completion", "(c) Completion", 100.0, "%"),
        ("joint", "(d) Joint L2 error", 1.0, "rad"),
    ]
    styles = {
        "D-only": {"color": COLORS["blue"], "marker": "s"},
        "M7-Raw": {"color": COLORS["orange"], "marker": "o"},
    }
    for ax, (key, title, scale, unit) in zip(axes.flat, panels):
        for method in ["D-only", "M7-Raw"]:
            rows = [row for row in budget if row["method"] == method]
            ax.plot(
                [int(row["iteration"]) / 1000 for row in rows],
                [float(row[key]) * scale for row in rows],
                label=method,
                markersize=4.2,
                markeredgewidth=0.6,
                **styles[method],
            )
        ax.axvline(59, color=COLORS["gray"], linestyle="--", linewidth=0.8, alpha=0.7)
        ax.set_title(title)
        ax.set_ylabel(unit)
        ax.grid(alpha=0.2)
    axes[1, 0].set_xlabel("Checkpoint iteration (k)")
    axes[1, 1].set_xlabel("Checkpoint iteration (k)")
    axes[0, 0].legend(frameon=False, ncol=2, loc="lower right")
    axes[0, 0].annotate("selected", xy=(59, 91.3028), xytext=(60.3, 90.75), arrowprops={"arrowstyle": "->", "lw": 0.7}, fontsize=6.8)
    fig.tight_layout(pad=0.6)
    return save_artwork(fig, "Fig_4_Validation_Budget_Scaling")


def plot_final_test(test: Sequence[dict[str, object]]) -> dict[str, Path]:
    methods = [str(row["method"]).replace(" (QD-HES)", "") for row in test]
    colors = [COLORS["gray"], COLORS["blue"], COLORS["orange"]]
    x = np.arange(len(methods))
    fig, axes = plt.subplots(2, 2, figsize=(7.48, 5.15))

    ax = axes[0, 0]
    metrics = [("micro", "Micro"), ("macro", "Macro"), ("completion", "Completion")]
    width = 0.24
    for index, (key, label) in enumerate(metrics):
        positions = x + (index - 1) * width
        bars = ax.bar(positions, [float(row[key]) * 100 for row in test], width, label=label)
        for bar in bars:
            ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.25, f"{bar.get_height():.1f}", ha="center", va="bottom", fontsize=5.9)
    ax.set_xticks(x, methods)
    ax.set_ylim(84, 98)
    ax.set_ylabel("Rate (%)")
    ax.set_title("(a) Success and completion")
    ax.grid(axis="y", alpha=0.2)
    ax.legend(frameon=False, ncol=3, fontsize=6.4, loc="lower right")

    ax = axes[0, 1]
    values = [float(row["body"]) * 100 for row in test]
    bars = ax.bar(x, values, color=colors)
    ax.set_xticks(x, methods)
    ax.set_ylabel("Mean error (cm)")
    ax.set_title("(b) Body-position error")
    ax.grid(axis="y", alpha=0.2)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.04, f"{value:.3f}", ha="center", fontsize=6.5)

    ax = axes[1, 0]
    values = [float(row["joint"]) for row in test]
    bars = ax.bar(x, values, color=colors)
    ax.set_xticks(x, methods)
    ax.set_ylabel("Mean error (rad)")
    ax.set_title("(c) Joint-position L2 error")
    ax.grid(axis="y", alpha=0.2)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 0.007, f"{value:.3f}", ha="center", fontsize=6.5)

    ax = axes[1, 1]
    values = [int(row["failures"]) for row in test]
    bars = ax.bar(x, values, color=colors)
    ax.set_xticks(x, methods)
    ax.set_ylabel("Failed motions")
    ax.set_title("(d) Failure count")
    ax.grid(axis="y", alpha=0.2)
    for bar, value in zip(bars, values):
        ax.text(bar.get_x() + bar.get_width() / 2, value + 15, f"{value}", ha="center", fontsize=6.5)
    fig.tight_layout(pad=0.65)
    return save_artwork(fig, "Fig_5_Final_Test_Comparison")


def plot_category_deltas(rows: Sequence[dict[str, object]]) -> dict[str, Path]:
    ordered = sorted(rows, key=lambda row: float(row["delta_d"]))
    values = [float(row["delta_d"]) * 100 for row in ordered]
    labels = [f"{row['category']} (n={row['n']})" for row in ordered]
    colors = [COLORS["blue"] if value >= 0 else COLORS["orange"] for value in values]
    fig, ax = plt.subplots(figsize=(7.48, 4.70))
    y = np.arange(len(ordered))
    bars = ax.barh(y, values, color=colors, height=0.68)
    ax.axvline(0, color=COLORS["black"], linewidth=0.8)
    ax.set_yticks(y, labels)
    ax.set_xlabel("Success-rate difference (percentage points)")
    ax.grid(axis="x", alpha=0.2)
    for bar, value in zip(bars, values):
        offset = 0.12 if value >= 0 else -0.12
        ax.text(value + offset, bar.get_y() + bar.get_height() / 2, f"{value:+.2f}", ha="left" if value >= 0 else "right", va="center", fontsize=6.4)
    fig.tight_layout(pad=0.5)
    return save_artwork(fig, "Fig_6_Final_Test_Category_Deltas")


def plot_dense_probe(rows: Sequence[dict[str, object]]) -> dict[str, Path]:
    fig, ax = plt.subplots(figsize=(7.48, 3.15))
    x = [int(row["iteration"]) / 1000 for row in rows]
    ax.plot(x, [float(row["micro"]) * 100 for row in rows], color=COLORS["blue"], marker="o", markersize=2.3, label="Micro")
    ax.plot(x, [float(row["macro"]) * 100 for row in rows], color=COLORS["orange"], marker="o", markersize=2.3, label="Macro")
    ax.plot(x, [float(row["completion"]) * 100 for row in rows], color=COLORS["green"], marker="o", markersize=2.3, label="Completion")
    ax.axvline(59, color=COLORS["black"], linestyle="--", linewidth=0.8)
    ax.axvspan(62, 67, color=COLORS["orange"], alpha=0.12)
    ax.text(59.15, 8, "59k selected", fontsize=6.6)
    ax.text(62.3, 21, "collapse region", fontsize=6.6, color=COLORS["orange"])
    ax.set_ylim(-2, 102)
    ax.set_xlabel("Checkpoint iteration (k)")
    ax.set_ylabel("Probe metric (%)")
    ax.grid(alpha=0.2)
    ax.legend(frameon=False, ncol=3, loc="lower left")
    fig.tight_layout(pad=0.5)
    return save_artwork(fig, "Fig_S1_Dense_Validation_Probe")


def plot_graphical_abstract(test: Sequence[dict[str, object]]) -> dict[str, Path]:
    proposed = next(row for row in test if str(row["method"]).startswith("M7-Raw"))
    baseline = next(row for row in test if row["method"] == "D-only")
    fig, ax = plt.subplots(figsize=(13.28, 5.31))
    ax.set_xlim(0, 20)
    ax.set_ylim(0, 8)
    ax.axis("off")

    def ga_box(x: float, y: float, w: float, h: float, label: str, face: str, edge: str) -> None:
        ax.add_patch(Rectangle((x, y), w, h, facecolor=face, edgecolor=edge, linewidth=2.0))
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontsize=17, weight="bold", color=COLORS["black"])

    ga_box(0.35, 2.15, 3.10, 3.70, "Heterogeneous\nPHUMA motions\n\n76,086 trajectories", "#EAF2F8", COLORS["blue"])
    ga_box(4.25, 5.05, 3.20, 1.45, "Quality gate\nreliable reset starts", "#E8F5E9", COLORS["green"])
    ga_box(4.25, 3.25, 3.20, 1.45, "Diversity knowledge\n8 motion clusters", "#EAF2F8", COLORS["blue"])
    ga_box(4.25, 1.45, 3.20, 1.45, "Online tracking error\nmotion + segment", "#FFF7E6", COLORS["yellow"])
    ga_box(8.35, 2.15, 3.55, 3.70, "Hierarchical budget\n\nCluster -> Motion\n-> Segment -> Start", "#FCECE8", COLORS["orange"])
    ga_box(12.80, 2.15, 2.70, 3.70, "PPO tracking\n\nUnitree G1\n29 DoF", "#EDE9F7", COLORS["purple"])
    improvement = (int(baseline["failures"]) - int(proposed["failures"]))
    ga_box(
        16.35,
        1.45,
        3.30,
        5.05,
        f"Final Test\n7,592 held-out motions\n\nMicro  {float(proposed['micro'])*100:.2f}%\nMacro  {float(proposed['macro'])*100:.2f}%\nCompletion  {float(proposed['completion'])*100:.2f}%\n\n{improvement} fewer failures",
        "#F3F8F5",
        COLORS["green"],
    )
    add_arrow(ax, (3.45, 4.00), (4.25, 5.77), COLORS["gray"])
    add_arrow(ax, (3.45, 4.00), (4.25, 3.97), COLORS["gray"])
    add_arrow(ax, (3.45, 4.00), (4.25, 2.17), COLORS["gray"])
    add_arrow(ax, (7.45, 5.77), (8.35, 4.90), COLORS["gray"])
    add_arrow(ax, (7.45, 3.97), (8.35, 4.00), COLORS["gray"])
    add_arrow(ax, (7.45, 2.17), (8.35, 3.10), COLORS["gray"])
    add_arrow(ax, (11.90, 4.00), (12.80, 4.00), COLORS["gray"])
    add_arrow(ax, (15.50, 4.00), (16.35, 4.00), COLORS["gray"])
    fig.subplots_adjust(left=0.01, right=0.99, top=0.98, bottom=0.02)
    artwork = save_artwork(fig, "Graphical_Abstract", tight=False)

    # Elsevier's general graphical-abstract guidance requests at least 1328 x 531 px at 300 dpi.
    tiff_path = FIGURES / "Graphical_Abstract.tiff"
    preview = plt.imread(artwork["preview"])
    plt.imsave(tiff_path, preview, dpi=300, format="tiff", pil_kwargs={"compression": "tiff_lzw"})
    artwork["tiff"] = tiff_path
    return artwork


def build_paper_tables(
    formal: Sequence[dict[str, object]],
    budget: Sequence[dict[str, object]],
    test: Sequence[dict[str, object]],
    categories: Sequence[dict[str, object]],
    probe: Sequence[dict[str, object]],
    stats: dict[str, object],
) -> tuple[dict[str, tuple[list[str], list[list[object]]]], dict[str, object]]:
    test_by_method = {str(row["method"]): row for row in test}
    m7 = test_by_method["M7-Raw (QD-HES)"]
    donly = test_by_method["D-only"]
    global_raw = test_by_method["GlobalRaw"]

    m7_path = ROOT / str(m7["directory"]) / "per_motion.csv"
    donly_path = ROOT / str(donly["directory"]) / "per_motion.csv"
    global_path = ROOT / str(global_raw["directory"]) / "per_motion.csv"
    motion_bootstrap_d = paired_bootstrap(m7_path, donly_path)
    motion_bootstrap_g = paired_bootstrap(m7_path, global_path)
    bootstrap_d = grouped_paired_bootstrap(m7_path, donly_path)
    bootstrap_g = grouped_paired_bootstrap(m7_path, global_path)

    quality = stats["quality"]
    difficulty = stats["difficulty"]
    clusters = stats["clusters"]
    module_statistics = [
        ["Converted corpus", "Motions", "76,086", "PHUMA G1 trajectories converted to WBT .npz"],
        ["Training pool", "Motions", "6,000", "Fixed random seed-42 manifest"],
        ["Segment infrastructure", "Segments", f"{int(quality['segment_count']):,}", "Nominal duration: 1 s"],
        ["Quality audit", "Pass / borderline / reject", f"{int(quality['pass_count']):,} / {int(quality['borderline_count']):,} / {int(quality['reject_count']):,}", "Borderline segments remain eligible"],
        ["Quality gate", "Eligible reset-start frames", f"{float(quality['eligible_start_fraction'])*100:.4f}%", "5,998 motions retain at least one eligible start"],
        ["Intrinsic difficulty", "Bins", str(int(difficulty["num_bins"])), "Approximately equal-frequency train-only bins"],
        ["Diversity model", "Clusters", str(int(clusters["num_clusters"])), "K-means++ on 30 aggregated motion features"],
        ["Diversity model", "Silhouette score", f"{float(clusters['silhouette_score']):.4f}", "Diagnostic, sampled over 2,000 motions"],
    ]

    formal_rows = [
        [
            row["method"], row["budget"], row["checkpoint"], fmt(float(row["micro"])), fmt(float(row["macro"])),
            fmt(float(row["completion"])), fmt(float(row["body"])), fmt(float(row["joint"])), int(row["failures"]),
        ]
        for row in formal
    ]
    budget_rows = [
        [
            row["method"], row["budget"], row["checkpoint"], fmt(float(row["micro"])), fmt(float(row["macro"])),
            fmt(float(row["completion"])), fmt(float(row["body"])), fmt(float(row["joint"])), int(row["failures"]),
        ]
        for row in budget
    ]
    test_rows = [
        [
            row["method"], row["budget"], row["checkpoint"], int(row["successes"]), fmt(float(row["micro"])),
            fmt(float(row["macro"])), fmt(float(row["completion"])), fmt(float(row["body"])),
            fmt(float(row["joint"])), int(row["failures"]),
        ]
        for row in test
    ]

    metric_specs = [
        ("Micro success", "micro", "pp", True),
        ("Macro success", "macro", "pp", True),
        ("Completion", "completion", "pp", True),
        ("Body-position error", "body", "m", False),
        ("Joint-position L2", "joint", "rad", False),
    ]
    delta_rows: list[list[object]] = []
    for baseline_name, baseline, bootstrap in [
        ("D-only 54k", donly, bootstrap_d),
        ("GlobalRaw 34k", global_raw, bootstrap_g),
    ]:
        for label, key, unit, higher_is_better in metric_specs:
            estimate = float(bootstrap[key]["estimate"])
            ci_low = float(bootstrap[key]["ci_low"])
            ci_high = float(bootstrap[key]["ci_high"])
            scale = 100.0 if unit == "pp" else 1.0
            relative = estimate / float(baseline[key]) * 100
            relative_improvement = relative if higher_is_better else -relative
            delta_rows.append(
                [
                    f"M7-Raw 59k vs {baseline_name}",
                    label,
                    f"{estimate*scale:+.4f}",
                    f"[{ci_low*scale:+.4f}, {ci_high*scale:+.4f}]",
                    unit,
                    f"{relative_improvement:+.2f}%",
                ]
            )
        delta_rows.append(
            [
                f"M7-Raw 59k vs {baseline_name}",
                "Failed motions",
                f"{int(m7['failures'])-int(baseline['failures']):+d}",
                "not bootstrapped as a count",
                "motions",
                f"{(int(baseline['failures'])-int(m7['failures']))/int(baseline['failures'])*100:+.2f}% fewer",
            ]
        )

    category_rows = [
        [
            row["category"], int(row["n"]), fmt(float(row["global"])), fmt(float(row["donly"])),
            fmt(float(row["m7raw"])), pp(float(row["delta_d"])), int(row["d_failures"]), int(row["m_failures"]),
        ]
        for row in categories
    ]

    validation_lookup = {
        "GlobalRaw": next(row for row in formal if row["method"] == "GlobalRaw"),
        "D-only": next(row for row in budget if row["method"] == "D-only" and int(row["iteration"]) == 54000),
        "M7-Raw (QD-HES)": next(row for row in budget if row["method"] == "M7-Raw" and int(row["iteration"]) == 59000),
    }
    generalization_rows: list[list[object]] = []
    for row in test:
        validation = validation_lookup[str(row["method"])]
        generalization_rows.append(
            [
                row["method"], row["checkpoint"], fmt(float(validation["micro"])), fmt(float(row["micro"])),
                pp(float(row["micro"]) - float(validation["micro"])), fmt(float(validation["macro"])),
                fmt(float(row["macro"])), pp(float(row["macro"]) - float(validation["macro"])),
            ]
        )

    termination_keys = [
        "completed",
        "ee_body_pos",
        "anchor_pos",
        "anchor_ori",
        "anchor_pos+ee_body_pos",
        "anchor_pos+anchor_ori+ee_body_pos",
    ]
    termination_rows: list[list[object]] = []
    for key in termination_keys:
        termination_rows.append(
            [
                key,
                int(global_raw["payload"].get("termination_reason_counts", {}).get(key, 0)),
                int(donly["payload"].get("termination_reason_counts", {}).get(key, 0)),
                int(m7["payload"].get("termination_reason_counts", {}).get(key, 0)),
            ]
        )

    probe_rows = [
        [
            int(row["iteration"]), row["checkpoint"], fmt(float(row["micro"])), fmt(float(row["macro"])),
            fmt(float(row["completion"])), fmt(float(row["body"])), fmt(float(row["joint"])), int(row["failures"]),
        ]
        for row in probe
    ]

    supplementary_specs = [
        (
            "V2-Diag-A",
            "Quality + difficulty + diversity; motion raw / segment relative gap",
            "34k",
            "model_25000.pt",
            "evaluations/formal_v2/A_raw_motion_gap_segment_seed42/validation_full/model_25000",
        ),
        (
            "M7-v2 (lambda=0.10)",
            "Quality + difficulty + diversity; raw segment error with gap correction",
            "34k",
            "model_33500.pt",
            "evaluations/formal_v2/M7_lambda0p10_seed42/validation_full/model_33500",
        ),
    ]
    supplementary = [metric_record(name, budget_label, checkpoint, directory) for name, _, budget_label, checkpoint, directory in supplementary_specs]
    supplementary_rows = [
        [
            row["method"], supplementary_specs[index][1], row["budget"], row["checkpoint"], fmt(float(row["micro"])),
            fmt(float(row["macro"])), fmt(float(row["completion"])), fmt(float(row["body"])),
            fmt(float(row["joint"])), int(row["failures"]),
        ]
        for index, row in enumerate(supplementary)
    ]

    joint_gap = read_json(ROOT / "outputs/joint_gap_stage6_probe500/probe_summary.json")
    joint_gap_rows: list[list[object]] = []
    for method_name, key, checkpoint in [
        ("M7-Raw", "m7_raw", "model_10000.pt"),
        ("M7-JGap (lambda_joint=0.025)", "m7_jgap", "model_9999.pt"),
    ]:
        row = joint_gap[key]
        joint_gap_rows.append(
            [
                method_name, "10k Probe500", checkpoint, fmt(float(row["micro_success_rate"])),
                fmt(float(row["macro_success_rate"])), fmt(float(row["mean_completion_ratio"])),
                fmt(float(row["mean_body_position_error_m"])), fmt(float(row["mean_joint_position_error_l2_rad"])),
                int(row["num_failure"]),
            ]
        )

    tables = {
        "Table_01_Module_Design": (["Module", "Design", "Function"], MODULE_ROWS),
        "Table_02_Module_Statistics": (["Component", "Statistic", "Value", "Interpretation"], module_statistics),
        "Table_03_Data_Splits": (["Split", "Manifest", "Motions", "SHA256"], SPLIT_ROWS),
        "Table_04_Method_Matrix": (["Method", "Motion level", "Segment level", "Quality", "Difficulty", "Diversity", "Evidence"], METHOD_MATRIX),
        "Table_05_PPO_Configuration": (["Item", "Setting"], PPO_ROWS),
        "Table_06_Reward_Terms": (["Term", "Weight", "Definition"], REWARD_ROWS),
        "Table_07_Metric_Definitions": (["Metric", "Definition"], METRIC_ROWS),
        "Table_08_Validation_Ablation_34k": (["Method", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body err. (m) ↓", "Joint L2 (rad) ↓", "Failures ↓"], formal_rows),
        "Table_09_Validation_Budget_Scaling": (["Method", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body err. (m) ↓", "Joint L2 (rad) ↓", "Failures ↓"], budget_rows),
        "Table_10_Final_Test": (["Method", "Budget", "Checkpoint", "Successes ↑", "Micro ↑", "Macro ↑", "Completion ↑", "Body err. (m) ↓", "Joint L2 (rad) ↓", "Failures ↓"], test_rows),
        "Table_11_Final_Test_Paired_Deltas": (["Comparison", "Metric", "Difference", "95% source-group bootstrap CI", "Unit", "Relative improvement"], delta_rows),
        "Table_12_Validation_to_Test": (["Method", "Checkpoint", "Val micro ↑", "Test micro ↑", "Delta micro", "Val macro ↑", "Test macro ↑", "Delta macro"], generalization_rows),
        "Table_13_Final_Test_Categories": (["Category", "N", "GlobalRaw ↑", "D-only ↑", "M7-Raw ↑", "M7-Raw - D-only", "D failures ↓", "M7-Raw failures ↓"], category_rows),
        "Table_14_Termination_Reasons": (["Termination reason", "GlobalRaw", "D-only", "M7-Raw"], termination_rows),
        "Table_S1_Diagnostic_Full_Validation": (["Method", "Design", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body err. ↓", "Joint L2 ↓", "Failures ↓"], supplementary_rows),
        "Table_S2_Joint_Gap_Probe": (["Method", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body err. ↓", "Joint L2 ↓", "Failures ↓"], joint_gap_rows),
        "Table_S3_Dense_Probe": (["Iteration", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body err. ↓", "Joint L2 ↓", "Failures / 500 ↓"], probe_rows),
    }
    analysis = {
        "bootstrap_m7_vs_donly": bootstrap_d,
        "bootstrap_m7_vs_globalraw": bootstrap_g,
        "motion_bootstrap_m7_vs_donly": motion_bootstrap_d,
        "motion_bootstrap_m7_vs_globalraw": motion_bootstrap_g,
        "paired_m7_vs_donly": exact_paired_outcomes(m7_path, donly_path),
        "paired_m7_vs_globalraw": exact_paired_outcomes(m7_path, global_path),
        "category_wins_ties_losses": {
            "wins": sum(float(row["delta_d"]) > 1e-12 for row in categories),
            "ties": sum(abs(float(row["delta_d"])) <= 1e-12 for row in categories),
            "losses": sum(float(row["delta_d"]) < -1e-12 for row in categories),
        },
        "m7": m7,
        "donly": donly,
        "globalraw": global_raw,
    }
    return tables, analysis


def latex_escape(value: object) -> str:
    text = str(value)
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "↑": r"$\uparrow$",
        "↓": r"$\downarrow$",
    }
    for source, replacement in replacements.items():
        text = text.replace(source, replacement)
    return text


def latex_table(
    caption: str,
    label: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    *,
    wide: bool = False,
    bold_contains: str | None = None,
) -> str:
    environment = "table*" if wide else "table"
    align = "l" + "r" * (len(headers) - 1)
    lines = [
        f"\\begin{{{environment}}}[t]",
        "\\centering",
        f"\\caption{{{latex_escape(caption)}}}",
        f"\\label{{{label}}}",
        "\\resizebox{\\textwidth}{!}{%" if wide else "",
        f"\\begin{{tabular}}{{{align}}}",
        "\\toprule",
        " & ".join(latex_escape(value) for value in headers) + r" \\",
        "\\midrule",
    ]
    for row in rows:
        escaped = [latex_escape(value) for value in row]
        if bold_contains and any(bold_contains in str(value) for value in row):
            escaped = [f"\\textbf{{{value}}}" for value in escaped]
        lines.append(" & ".join(escaped) + r" \\")
    lines.extend(["\\bottomrule", "\\end{tabular}", "}" if wide else "", f"\\end{{{environment}}}", ""])
    return "\n".join(line for line in lines if line != "")


def write_table_artifacts(tables: dict[str, tuple[list[str], list[list[object]]]]) -> None:
    TABLES.mkdir(parents=True, exist_ok=True)
    for name, (headers, rows) in tables.items():
        write_csv(TABLES / f"{name}.csv", headers, rows)

    captions = {
        "Table_03_Data_Splits": "Dataset partitions and frozen manifest identities.",
        "Table_08_Validation_Ablation_34k": "Full-validation comparison under the common 34k-iteration budget.",
        "Table_09_Validation_Budget_Scaling": "Validation performance across extended training budgets.",
        "Table_10_Final_Test": "Final Test performance on 7,592 held-out motions.",
        "Table_11_Final_Test_Paired_Deltas": "Paired Test differences for the proposed M7-Raw/QD-HES checkpoint.",
        "Table_13_Final_Test_Categories": "Category-level Test success rates and failure counts.",
    }
    chunks = ["% Requires \\usepackage{booktabs} and \\usepackage{graphicx}.\n"]
    for name, caption in captions.items():
        headers, rows = tables[name]
        bold_contains = None
        if name == "Table_09_Validation_Budget_Scaling":
            bold_contains = "59k selected"
        elif name == "Table_10_Final_Test":
            bold_contains = "M7-Raw (QD-HES)"
        chunks.append(
            latex_table(
                caption,
                f"tab:{name.lower()}",
                headers,
                rows,
                wide=len(headers) > 5 or name == "Table_03_Data_Splits",
                bold_contains=bold_contains,
            )
        )
    (TABLES / "kbs_tables.tex").write_text("\n\n".join(chunks), encoding="utf-8")


def table_markdown(tables: dict[str, tuple[list[str], list[list[object]]]], name: str) -> str:
    headers, rows = tables[name]
    return markdown_table(headers, rows)


def bootstrap_interval(analysis: dict[str, object], key: str, *, scale: float = 1.0) -> str:
    item = analysis[key]
    return f"{float(item['estimate'])*scale:+.4f} [{float(item['ci_low'])*scale:+.4f}, {float(item['ci_high'])*scale:+.4f}]"


def build_markdown(
    tables: dict[str, tuple[list[str], list[list[object]]]],
    analysis: dict[str, object],
    categories: Sequence[dict[str, object]],
    stats: dict[str, object],
) -> str:
    m7 = analysis["m7"]
    donly = analysis["donly"]
    global_raw = analysis["globalraw"]
    boot_d = analysis["bootstrap_m7_vs_donly"]
    paired_d = analysis["paired_m7_vs_donly"]
    wtl = analysis["category_wins_ties_losses"]
    gains = sorted(categories, key=lambda row: float(row["delta_d"]), reverse=True)[:5]
    losses = sorted(categories, key=lambda row: float(row["delta_d"]))[:3]

    body_reduction = (float(donly["body"]) - float(m7["body"])) / float(donly["body"]) * 100
    joint_reduction = (float(donly["joint"]) - float(m7["joint"])) / float(donly["joint"]) * 100
    failure_reduction = (int(donly["failures"]) - int(m7["failures"])) / int(donly["failures"]) * 100
    quality = stats["quality"]
    clusters = stats["clusters"]

    sections: list[str] = [
        f"# {PAPER_TITLE}",
        f"**Target journal:** Knowledge-Based Systems (KBS)  \n"
        f"**Material freeze:** {DATE}  \n"
        f"**Proposed short name:** {SHORT_METHOD_NAME} (repository variant: M7-Raw)",
        (
            "> Evidence status: every numerical result below is recomputed from the local CSV/JSON artifacts. "
            "The Test manifest is disjoint from Train and Validation, but it is not described as a pristine, "
            "never-accessed blind set because an earlier 34k Test protocol exists in the repository. The 59k "
            "M7-Raw and 54k D-only checkpoints were selected using Validation only, and no post-Test reselection "
            "is permitted."
        ),
        "## 1. Submission-ready front matter",
        "### Proposed title",
        PAPER_TITLE,
        "### Abstract",
        ABSTRACT,
        "### Highlights",
        "\n".join(f"- {item}" for item in HIGHLIGHTS),
        "### Keywords",
        "; ".join(KEYWORDS),
        "## 2. KBS positioning",
        (
            "The paper should be positioned as knowledge-guided data allocation for reinforcement learning, "
            "not merely as a larger humanoid tracking benchmark. Its central knowledge components are: "
            "offline trajectory-quality labels, train-only motion-cluster structure, and online policy-competence "
            "estimates. Their responsibilities are explicitly separated and composed into one hierarchical "
            "sampling distribution. This framing connects the robotics contribution to KBS interests in machine "
            "learning methodology, knowledge representation, and data-driven optimization."
        ),
        "### Research questions",
        (
            "- **RQ1:** Does hierarchical error allocation with diversity budgets outperform uniform and flat "
            "error sampling under a common 34k training budget?\n"
            "- **RQ2:** Does trajectory-quality knowledge provide an additional benefit when the training budget "
            "is long enough?\n"
            "- **RQ3:** Is policy quality monotonic in PPO iterations, or is independent checkpoint selection "
            "necessary?\n"
            "- **RQ4:** Do Validation-selected policies preserve their ranking on the held-out Test split?"
        ),
        "### Claimed contributions",
        (
            "1. We formulate large-library humanoid tracking as hierarchical training-data allocation and compose "
            "quality, diversity, and policy-competence knowledge at cluster, motion, and segment levels.\n"
            "2. We introduce a trajectory-quality gate that changes episode-start eligibility without deleting "
            "complete motions, conflating artifacts with difficulty, or consuming Validation/Test information.\n"
            "3. We provide a controlled study over 76,086 converted PHUMA trajectories, including fixed-budget "
            "ablations, budget scaling, dense checkpoint diagnostics, a 7,592-motion final Test, category analysis, "
            "and source-group-aware paired uncertainty estimates.\n"
            "4. We identify a budget-dependent quality-gate benefit and a non-monotonic optimum near 59k "
            "iterations, showing that sampling design and checkpoint selection jointly determine the outcome."
        ),
        "## 3. System and module design",
        table_markdown(tables, "Table_01_Module_Design"),
        "### 3.1 Data conversion and multi-motion tracking",
        (
            "PHUMA G1 trajectories are converted from `.npy` to a common WBT `.npz` schema containing joint "
            "positions and velocities and world-frame body poses and velocities. The resulting library contains "
            "76,086 motions (approximately 19 GB), 29 actuated degrees of freedom, and 30 tracked bodies. The "
            "loader accepts one file, a directory, or an ordered manifest. Each parallel environment receives an "
            "independent motion and start frame, so one policy learns the complete training library."
        ),
        "### 3.2 Stage 0: segment infrastructure",
        (
            "For motion `m`, `segment_frames[m] = max(1, round(fps[m] * 1 s))` and "
            "`num_segments[m] = ceil(num_frames[m] / segment_frames[m])`. A global segment index maps each "
            "motion/local-segment pair to shared counters, validity masks, sampling probabilities, and checkpoint "
            "state. The 6,000-motion training manifest produces 21,575 one-second segments. The disabled-module "
            "path preserves the legacy uniform random-number sequence, enabling controlled ablations."
        ),
        "### 3.3 Module 1: trajectory-quality gate",
        (
            "The offline audit checks finite values, quaternion norms, URDF joint limits, velocity consistency, "
            "isolated acceleration/jerk peaks, pose discontinuities, ground penetration, and foot sliding. It uses "
            "reference trajectories only; policy reward, tracking error, success, and Test data are excluded. For "
            "segment `(m,s)`, the score is `Q(m,s) = clip(1 - sum_i w_i severity_i / sum_i w_i, 0, 1)`. "
            f"The audit labels {int(quality['pass_count']):,} segments pass, "
            f"{int(quality['borderline_count']):,} borderline, and {int(quality['reject_count']):,} reject. "
            f"With borderline segments allowed, {float(quality['eligible_start_fraction'])*100:.4f}% of candidate "
            "start frames remain eligible and 5,998 of 6,000 motions retain at least one eligible start. Rejected "
            "segments cannot initialize an episode, but a rollout may traverse them after a valid start."
        ),
        "### 3.4 Module 2: intrinsic difficulty",
        (
            "Twenty-eight policy-independent kinematic and contact descriptors capture root motion, accelerations, "
            "joint activity, body-height variation, limb speed, support state, flight ratio, and contact switches. "
            "Derivatives are computed on complete motions before segment aggregation. Train-only robust scaling "
            "uses the median and `1.4826 * MAD`, with a P05-P95 fallback for near-zero MAD. A weighted score is "
            "mapped through an empirical CDF into ten approximately balanced bins. Difficulty does not remove or "
            "sample data directly; it only calibrates learning-gap variants."
        ),
        "### 3.5 Module 3: online error and learning state",
        (
            "Body, joint, orientation, completion, and termination signals are attributed to the active motion and "
            "segment. Exponential moving averages (`rho=0.95`) are updated after a 1,000-iteration warm-up, and "
            "sampling probabilities are refreshed every 50 iterations. Raw-error variants prioritize current "
            "tracking errors directly. Learning-gap variants compare each segment with peers in the same intrinsic-"
            "difficulty bin. All adaptive variants retain a 0.15 uniform mixture, an undersampling bonus, and "
            "probability caps to maintain coverage."
        ),
        "### 3.6 Module 4: diversity constraint",
        (
            "Seventeen segment-level source descriptors are aggregated into 30 motion-level features. Train-only "
            "robust scaling and K-means++ produce eight frozen motion clusters with sizes "
            f"`{list(map(int, clusters['cluster_sizes']))}` (silhouette={float(clusters['silhouette_score']):.4f}). "
            "Neither source-directory labels nor evaluation categories enter clustering. With `C` clusters, "
            "`f=0.5`, and `alpha=0.5`, the cluster budget is "
            "`P_div(c) = f/C + (1-f) n_c^alpha / sum_j n_j^alpha`. Errors act only within a cluster, yielding "
            "`P(c,m,s) = P_div(c) P_err(m|c) P_err(s|m)`. This separation prevents one family of similar, "
            "high-error motions from monopolizing the rollout budget."
        ),
        "### 3.7 Proposed method contract",
        (
            "QD-HES is the repository's M7-Raw variant: hierarchical raw error at the motion and segment levels, "
            "the quality gate enabled, diversity budgets enabled, and difficulty/learning-gap calibration disabled. "
            "D-only is the strongest structural baseline and differs by exactly one switch: its quality gate is off."
        ),
        table_markdown(tables, "Table_04_Method_Matrix"),
        "## 4. Experimental protocol",
        "### 4.1 Frozen data partitions",
        table_markdown(tables, "Table_03_Data_Splits"),
        (
            "The splits are source-group disjoint: chunks from the same underlying sequence are kept together. "
            "The 500-motion probe is a subset of full Validation and is used only to screen checkpoints. Full "
            "Validation selects candidates; Test does not select a method or checkpoint."
        ),
        "### 4.2 PPO and environment",
        table_markdown(tables, "Table_05_PPO_Configuration"),
        "### 4.3 Reward and termination",
        table_markdown(tables, "Table_06_Reward_Terms"),
        (
            "Early termination is triggered by excessive anchor-height error (0.25 m), anchor-orientation error "
            "(0.8 rad), or wrist/ankle end-body height error (0.25 m); timeout is treated separately. Network, "
            "reward, termination, randomization, and evaluator settings are fixed across samplers."
        ),
        "### 4.4 Metrics and checkpoint selection",
        table_markdown(tables, "Table_07_Metric_Definitions"),
        (
            "Checkpoint selection is macro-first because the 17 categories are highly imbalanced. If macro scores "
            "differ by at most 0.002, the rule compares micro success, completion, body error, and then the earlier "
            "checkpoint. Joint L2 is reported but is not a selection criterion. Very early failures can create "
            "artificially low mean errors, so errors must always be read together with success and completion."
        ),
        "### 4.5 Statistical analysis",
        (
            "Final Test comparisons are paired by motion path. The primary 95% intervals use 10,000 paired cluster "
            "bootstrap resamples of the 3,610 source groups (seed 42), preserving within-source chunk dependence. "
            "Macro success is bootstrapped by category and source group. A motion-level paired bootstrap is also "
            "exported as a sensitivity analysis. These intervals quantify evaluation-set uncertainty, not training-"
            "seed uncertainty."
        ),
        "## 5. Results",
        "### 5.1 Common-budget 34k full Validation",
        table_markdown(tables, "Table_08_Validation_Ablation_34k"),
        (
            "At 34k, D-only has the best macro success (0.899319). Relative to GlobalRaw, it improves micro by "
            "2.095 percentage points, macro by 1.792 points, and completion by 1.151 points while reducing failures "
            "from 1,042 to 882. M7-Raw is marginally better than D-only in micro, body error, joint error, and "
            "failures, but its macro score is 0.629 points lower; therefore the original macro-first D-only choice "
            "is internally consistent. M2 and M3 are absent because only pilot evidence exists."
        ),
        "### 5.2 Budget scaling and checkpoint stability",
        table_markdown(tables, "Table_09_Validation_Budget_Scaling"),
        (
            "M7-Raw rises from 0.893026 macro at 34k to 0.913028 at 59k, with 244 fewer Validation failures. "
            "D-only reaches only 0.898413 macro at its 54k selection point. M7-Raw then drops to 0.899619 at 60k "
            "and 0.898613 at 70k. The dense probe additionally contains a near-zero-success region around 62k-67k, "
            "which demonstrates severe non-monotonic training instability rather than smooth saturation."
        ),
        "### 5.3 Final Test",
        table_markdown(tables, "Table_10_Final_Test"),
        (
            f"QD-HES/M7-Raw succeeds on {int(m7['successes']):,} of 7,592 Test motions. Against D-only, micro "
            f"success improves by {(float(m7['micro'])-float(donly['micro']))*100:.3f} points, macro success by "
            f"{(float(m7['macro'])-float(donly['macro']))*100:.3f} points, and completion by "
            f"{(float(m7['completion'])-float(donly['completion']))*100:.3f} points. Body error falls by "
            f"{body_reduction:.2f}%, joint L2 by {joint_reduction:.2f}%, and failures by "
            f"{int(donly['failures'])-int(m7['failures'])} motions ({failure_reduction:.2f}%). Relative to "
            f"GlobalRaw, M7-Raw adds {int(m7['successes'])-int(global_raw['successes'])} successful motions."
        ),
        "### 5.4 Paired uncertainty and outcome changes",
        table_markdown(tables, "Table_11_Final_Test_Paired_Deltas"),
        (
            f"For M7-Raw versus D-only, the source-group bootstrap estimates are: micro "
            f"{bootstrap_interval(boot_d, 'micro', scale=100)} percentage points, macro "
            f"{bootstrap_interval(boot_d, 'macro', scale=100)} points, completion "
            f"{bootstrap_interval(boot_d, 'completion', scale=100)} points, body error "
            f"{bootstrap_interval(boot_d, 'body') } m, and joint L2 "
            f"{bootstrap_interval(boot_d, 'joint')} rad. At the motion level, M7-Raw newly solves "
            f"{int(paired_d['a_only'])} motions that D-only fails, while D-only uniquely solves "
            f"{int(paired_d['b_only'])}; {int(paired_d['both'])} are solved by both and "
            f"{int(paired_d['neither'])} by neither."
        ),
        "### 5.5 Category and failure analysis",
        (
            f"M7-Raw improves {int(wtl['wins'])} categories, ties {int(wtl['ties'])}, and declines in "
            f"{int(wtl['losses'])}. The five largest success-rate gains over D-only are "
            + ", ".join(f"{row['category']} ({float(row['delta_d'])*100:+.2f} pp, n={row['n']})" for row in gains)
            + ". The three lowest deltas are "
            + ", ".join(f"{row['category']} ({float(row['delta_d'])*100:+.2f} pp, n={row['n']})" for row in losses)
            + ". Small categories should not support strong standalone claims."
        ),
        table_markdown(tables, "Table_14_Termination_Reasons"),
        (
            "The dominant residual failure mode is end-body position error. Its reduction explains most of the "
            "net failure improvement and is consistent with the lower body and joint tracking errors. Qualitative "
            "videos should therefore include wrist/ankle-intensive examples, not only walking motions."
        ),
        "### 5.6 Validation-to-Test consistency",
        table_markdown(tables, "Table_12_Validation_to_Test"),
        (
            "The Validation-selected ranking is preserved on Test: M7-Raw remains first, D-only second, and "
            "GlobalRaw third on both micro and macro success. This supports generalization across the fixed split, "
            "although it does not replace multi-seed training evidence."
        ),
        "## 6. Discussion material",
        "### 6.1 Why diversity helps at a short budget",
        (
            "Raw tracking error is a fast competence signal, but without a cluster budget it can repeatedly sample "
            "many similar difficult motions. Diversity limits this concentration while retaining error prioritization "
            "inside each cluster. The 34k results support this interpretation: D-only substantially exceeds flat "
            "GlobalRaw and hierarchical M4."
        ),
        "### 6.2 Why the quality gate appears budget-dependent",
        (
            "Only 90 of 21,575 training segments are rejected, so a large immediate effect is not expected. Over "
            "billions of rollout transitions, however, repeatedly avoiding unreliable reset starts can reduce wasted "
            "episodes and stabilize credit assignment. The crossover after 34k is consistent with a cumulative, "
            "slow-onset benefit. This remains a mechanism hypothesis rather than a proven causal explanation and "
            "should be tested with additional training seeds or gate-usage traces."
        ),
        "### 6.3 Why learning-gap variants are not the main method",
        (
            "Difficulty calibration, EMA statistics, and relative-gap estimation introduce delay and estimator "
            "noise. M5-M7 do not beat the simpler raw-error variants at 34k, and the joint-gap pilot is dominated. "
            "These are informative negative results: more knowledge signals are useful only when their estimation "
            "cost and reliability match the learning regime."
        ),
        "### 6.4 Why 59k is an optimum, not an arbitrary endpoint",
        (
            "The selected checkpoint is the maximum of a predeclared Validation rule over evaluated candidates, "
            "not the final training checkpoint. The sharp 60k macro decrease and later collapse show that PPO and "
            "an adaptive sampling distribution form a non-stationary coupled process. Consequently, dense "
            "checkpointing and an independent Validation set are part of the experimental method."
        ),
        "## 7. Limitations and validity threats",
        (
            "- **Training stochasticity:** all formal policies use training seed 42. Bootstrap intervals do not "
            "measure across-run variance. The highest-priority additional experiment is two or more seeds for "
            "M7-Raw 59k, D-only 54k, and GlobalRaw 34k.\n"
            "- **Incomplete named ablation:** M2 and M3 have pilot runs but no 7,636-motion full Validation. Do not "
            "claim a complete M0-M7 ablation unless those evaluations are added.\n"
            "- **Test provenance:** the split is held out from Train/Validation but was accessed by an older 34k "
            "protocol. Describe the current comparison accurately and forbid Test-driven reselection.\n"
            "- **Simulation scope:** the evidence does not establish real-robot robustness or sim-to-real transfer.\n"
            "- **Data and profile status:** quality, difficulty, and clustering profiles should be frozen with their "
            "actual hashes; any remaining `provisional` label must be resolved before submission.\n"
            "- **External baseline scope:** current quantitative tables compare internal sampling variants under a "
            "shared implementation. Claims of state of the art require compatible external baselines or careful "
            "qualification.\n"
            "- **Compute reporting:** GPU model, wall-clock duration, and energy/compute accounting still need to be "
            "recorded for the final manuscript."
        ),
        "## 8. Conclusion draft",
        (
            "This work presents a knowledge-guided hierarchical sampling framework for training one humanoid "
            "whole-body tracking policy on a large heterogeneous motion library. The framework separates offline "
            "trajectory quality, train-only motion diversity, and online policy competence, then composes them into "
            "a cluster-motion-segment distribution. Under a common 34k budget, diversity-constrained hierarchical "
            "raw-error sampling is the strongest short-budget variant. With a longer Validation-selected budget, "
            "adding the quality gate yields a higher optimum: QD-HES at 59k reaches 0.9231 micro success and 0.9079 "
            "macro success on 7,592 held-out Test motions while reducing failures from 731 to 584 relative to D-only. "
            "Performance degrades after the optimum, emphasizing that structured data allocation and independent "
            "checkpoint selection are both necessary for reliable large-scale humanoid motion tracking."
        ),
        "## 9. Recommended manuscript structure",
        (
            "1. **Introduction:** large heterogeneous motion libraries, data-allocation problem, research gap, "
            "contributions.\n"
            "2. **Related work:** physics-based humanoid imitation; scalable whole-body tracking; curriculum, hard "
            "example mining, and prioritized sampling; data quality and diversity.\n"
            "3. **Problem formulation:** tracking MDP, segmented motion library, knowledge sources, objectives.\n"
            "4. **Method:** Stage 0 infrastructure, quality gate, online error, diversity budget, QD-HES algorithm.\n"
            "5. **Experimental setup:** PHUMA conversion, splits, Unitree G1, PPO, baselines, metrics, statistics.\n"
            "6. **Results:** 34k ablation, budget scaling, final Test, category/failure analysis.\n"
            "7. **Discussion:** budget dependence, negative results, instability, limitations.\n"
            "8. **Conclusion.**"
        ),
        "## 10. Figure plan and captions",
        "\n".join(f"- **{name}:** {caption}" for name, caption in FIGURE_CAPTIONS.items()),
        "## 11. Table inventory",
        (
            "Editable CSV files for Tables 1-14 and Supplementary Tables S1-S3 are in `tables/`. The main "
            "submission tables are also supplied in `tables/kbs_tables.tex` with `booktabs` and no vertical rules."
        ),
        "## 12. Related-work evidence map",
        (
            "- **Physics-based imitation:** use DeepMimic to introduce reference-motion tracking with RL.\n"
            "- **Scalable whole-body tracking:** contrast QD-HES with BeyondMimic's compact tracker and GMT's "
            "adaptive sampling/MoE approach.\n"
            "- **Large motion data:** cite PHUMA for the physically reliable 73-hour source corpus and explain the "
            "local G1/WBT conversion.\n"
            "- **Optimization:** cite PPO for the policy update and K-means++ for frozen diversity clusters.\n"
            "- **Curriculum and prioritization:** distinguish QD-HES from curriculum learning and replay-based "
            "prioritization: it changes on-policy rollout assignment and explicitly factors diversity and quality.\n"
            "The included `.bib` is a verified seed list, not a complete literature review."
        ),
        "## 13. Reproducibility identity",
        table_markdown(tables, "Table_02_Module_Statistics"),
        (
            f"- Quality metadata SHA256: `{stats['quality_metadata_sha256']}`\n"
            f"- Difficulty metadata SHA256: `{stats['difficulty_metadata_sha256']}`\n"
            f"- Cluster metadata SHA256: `{stats['cluster_metadata_sha256']}`\n"
            f"- Final evaluator commit recorded by all three Test summaries: `{m7['payload']['git_commit']}`\n"
            f"- M7-Raw 59k checkpoint SHA256: `{m7['payload']['checkpoint_sha256']}`\n"
            f"- D-only 54k checkpoint SHA256: `{donly['payload']['checkpoint_sha256']}`\n"
            f"- GlobalRaw 34k checkpoint SHA256: `{global_raw['payload']['checkpoint_sha256']}`"
        ),
        "## 14. Claim guardrails",
        (
            "**Supported:** QD-HES is the best evaluated candidate on the current full Validation and held-out Test "
            "splits; it improves every reported aggregate Test metric over D-only and GlobalRaw; the 59k checkpoint "
            "is better than 60k and 70k on full Validation.\n\n"
            "**Not yet supported:** universal superiority across random seeds, a complete M0-M7 ablation, real-robot "
            "transfer, state-of-the-art performance against external systems, or a pristine never-accessed blind Test."
        ),
    ]
    return "\n\n".join(sections) + "\n"


def set_cell_margins(cell, top: int = 45, start: int = 55, bottom: int = 45, end: int = 55) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, edge: str, *, size: int = 8, color: str = "000000") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    tag = f"w:{edge}"
    element = borders.find(qn(tag))
    if element is None:
        element = OxmlElement(tag)
        borders.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), str(size))
    element.set(qn("w:space"), "0")
    element.set(qn("w:color"), color)


def set_run_style(run, *, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(5)
    for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string("17365D")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(4)


def configure_headers(doc: Document) -> None:
    first = doc.sections[0]
    header = first.header.paragraphs[0]
    header.text = f"PHUMA-WBT | KBS paper materials | {DATE}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        set_run_style(run, size=8, color="64748B")
    footer = first.footer.paragraphs[0]
    footer.clear()
    add_page_number(footer)
    for run in footer.runs:
        set_run_style(run, size=8, color="64748B")
    for section in doc.sections[1:]:
        section.header.is_linked_to_previous = True
        section.footer.is_linked_to_previous = True


def add_body(doc: Document, text: str, *, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.65)
    run = paragraph.add_run(text)
    set_run_style(run)
    run.italic = italic


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_run_style(run)


def add_numbered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(text)
    set_run_style(run)


def add_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.right_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_style(run, size=9.5, bold=True, color="7F6000")
    run.italic = True


def add_equation_block(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(5)
    run = paragraph.add_run(text)
    run.font.name = "Cambria Math"
    run.font.size = Pt(9.5)


def add_kbs_table(
    doc: Document,
    headers: Sequence[str],
    rows: Sequence[Sequence[object]],
    *,
    font_size: float = 8.0,
    bold_row: Callable[[Sequence[object]], bool] | None = None,
) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        cell = header_cells[index]
        cell.text = str(header)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_border(cell, "top", size=12)
        set_cell_border(cell, "bottom", size=8)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_style(run, size=font_size, bold=True)
    header_tr_pr = table.rows[0]._tr.get_or_add_trPr()
    header_tr_pr.append(OxmlElement("w:tblHeader"))
    for row_values in rows:
        cells = table.add_row().cells
        make_bold = bool(bold_row(row_values)) if bold_row else False
        for index, value in enumerate(row_values):
            cell = cells[index]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index < 2 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_style(run, size=font_size, bold=make_bold)
    if len(table.rows) > 1:
        for cell in table.rows[-1].cells:
            set_cell_border(cell, "bottom", size=12)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, preview: Path, caption: str, *, width: float = 6.55) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(preview), width=Inches(width))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    caption_paragraph.paragraph_format.space_after = Pt(7)
    run = caption_paragraph.add_run(caption)
    set_run_style(run, size=9)


def switch_orientation(doc: Document, *, landscape: bool) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    if landscape:
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        section.left_margin = Cm(1.3)
        section.right_margin = Cm(1.3)
    else:
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)


def build_docx(
    tables: dict[str, tuple[list[str], list[list[object]]]],
    analysis: dict[str, object],
    categories: Sequence[dict[str, object]],
    stats: dict[str, object],
    artwork: dict[str, dict[str, Path]],
) -> None:
    m7 = analysis["m7"]
    donly = analysis["donly"]
    global_raw = analysis["globalraw"]
    boot_d = analysis["bootstrap_m7_vs_donly"]
    paired_d = analysis["paired_m7_vs_donly"]
    wtl = analysis["category_wins_ties_losses"]

    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(38)
    run = title.add_run(PAPER_TITLE)
    set_run_style(run, size=23, bold=True, color="17365D")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("English manuscript materials for Knowledge-Based Systems")
    set_run_style(run, size=13, bold=True, color="475569")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"Project: {ROOT}\nEvidence freeze: {DATE}\nMethod: {SHORT_METHOD_NAME} / M7-Raw")
    set_run_style(run, size=10, color="64748B")
    add_figure(doc, artwork["Graphical_Abstract"]["preview"], FIGURE_CAPTIONS["Graphical_Abstract"], width=6.7)
    add_note(
        doc,
        "Evidence note: the Test split is disjoint from Train and Validation. An older 34k Test protocol exists, "
        "so the manuscript should say held-out Test, not pristine never-accessed blind Test. The 59k/54k "
        "checkpoints were selected using Validation only and are now frozen against Test-driven changes.",
    )
    doc.add_page_break()

    doc.add_heading("1. Executive Summary", level=1)
    add_body(
        doc,
        "The project now provides an end-to-end large-library humanoid whole-body tracking system, a modular "
        "knowledge-guided sampling method, complete full-Validation evidence for the principal variants, budget "
        "scaling through 70k iterations, and a final three-model Test comparison. The recommended paper method is "
        "M7-Raw, named QD-HES for publication. It combines a trajectory-quality gate, a frozen diversity budget, "
        "and hierarchical raw-error allocation.",
    )
    executive_rows = [
        ["34k fixed-budget result", "D-only, model_33500.pt", "Best macro Validation: 0.899319"],
        ["Final proposed model", "QD-HES, model_59000.pt", "Test micro 0.923077; macro 0.907906"],
        ["Strong structural baseline", "D-only, model_54000.pt", "Test micro 0.903714; macro 0.881601"],
        ["Flat baseline", "GlobalRaw, model_33500.pt", "Test micro 0.868678; macro 0.858846"],
        ["Training-limit result", "M7-Raw near 59k", "60k macro drops; 62k-67k probe collapse"],
    ]
    add_kbs_table(doc, ["Evidence layer", "Checkpoint", "Finding"], executive_rows, font_size=8.8)
    add_body(
        doc,
        f"On the 7,592-motion Test split, QD-HES solves {int(m7['successes']):,} motions and fails "
        f"{int(m7['failures']):,}. It adds {int(m7['successes'])-int(donly['successes'])} successes over D-only "
        f"and {int(m7['successes'])-int(global_raw['successes'])} over GlobalRaw. All six aggregate metrics favor "
        "QD-HES.",
    )

    doc.add_heading("2. Submission-Ready Front Matter", level=1)
    doc.add_heading("2.1 Title", level=2)
    add_body(doc, PAPER_TITLE)
    doc.add_heading("2.2 Abstract", level=2)
    add_body(doc, ABSTRACT)
    doc.add_heading("2.3 Highlights", level=2)
    for item in HIGHLIGHTS:
        add_bullet(doc, item)
    doc.add_heading("2.4 Keywords", level=2)
    add_body(doc, "; ".join(KEYWORDS))

    doc.add_heading("3. KBS Positioning and Paper Claims", level=1)
    add_body(
        doc,
        "The strongest KBS framing is knowledge-guided data allocation for on-policy reinforcement learning. "
        "QD-HES represents trajectory reliability as segment labels, motion diversity as frozen clusters, and "
        "policy competence as online tracking statistics. These knowledge sources enter different levels of one "
        "factorized sampling distribution, making the method interpretable and auditable.",
    )
    doc.add_heading("3.1 Research questions", level=2)
    for item in [
        "RQ1: Does hierarchical error allocation with diversity budgets outperform uniform and flat error sampling under a common 34k budget?",
        "RQ2: Does trajectory-quality knowledge add value when the training budget is sufficiently long?",
        "RQ3: Is performance monotonic in PPO iterations, or is independent checkpoint selection necessary?",
        "RQ4: Do Validation-selected policies preserve their ordering on the held-out Test split?",
    ]:
        add_bullet(doc, item)
    doc.add_heading("3.2 Contributions", level=2)
    contributions = [
        "A hierarchical on-policy data-allocation formulation that composes quality, diversity, and policy-competence knowledge at cluster, motion, and segment levels.",
        "A trajectory-quality gate that changes episode-start eligibility without deleting complete motions or using policy/evaluation outcomes as quality labels.",
        "A controlled study over 76,086 converted PHUMA trajectories with common-budget ablations, budget scaling, dense checkpoint diagnostics, held-out Test evaluation, and source-group-aware uncertainty.",
        "Evidence that quality gating has a budget-dependent benefit and that adaptive PPO training has a non-monotonic optimum near 59k iterations.",
    ]
    for item in contributions:
        add_numbered(doc, item)

    doc.add_heading("4. System and Module Design", level=1)
    add_figure(doc, artwork["Fig_1_Overall_Framework"]["preview"], FIGURE_CAPTIONS["Fig_1_Overall_Framework"])
    headers, rows = tables["Table_01_Module_Design"]
    add_kbs_table(doc, headers, rows, font_size=7.5)

    doc.add_heading("4.1 Data conversion and motion library", level=2)
    add_body(
        doc,
        "The conversion stage maps PHUMA G1 trajectories to a uniform WBT NPZ schema with joint states and "
        "world-frame body states. The library contains 76,086 converted motions, 29 actuated joints, and 30 "
        "bodies. MotionLoader accepts individual files, directories, or ordered manifests; MotionCommand assigns "
        "an independent reference and start frame to every parallel environment.",
    )
    doc.add_heading("4.2 Stage 0: segment infrastructure", level=2)
    add_equation_block(
        doc,
        "segment_frames[m] = max(1, round(fps[m] * 1 s));   N_seg[m] = ceil(N_frame[m] / segment_frames[m])",
    )
    add_body(
        doc,
        "The training set contains 21,575 one-second segments. Global/local IDs, eligibility masks, coverage "
        "counters, probability checks, and adaptive-sampler state share one infrastructure. The legacy uniform "
        "path preserves its random-number call order when all research modules are disabled.",
    )
    doc.add_heading("4.3 Module 1: trajectory-quality gate", level=2)
    add_body(
        doc,
        "The train-only audit checks numerical validity, joint-limit compliance, velocity consistency, temporal "
        "continuity, ground penetration, and foot sliding. It does not read policy errors, rewards, success rates, "
        "or evaluation data. Each segment is labeled pass, borderline, or reject. Rejected segments cannot become "
        "episode starts, although a rollout may later traverse them.",
    )
    add_equation_block(doc, "Q(m,s) = clip(1 - sum_i w_i severity_i(m,s) / sum_i w_i, 0, 1)")
    headers, rows = tables["Table_02_Module_Statistics"]
    add_kbs_table(doc, headers, rows, font_size=7.8)
    doc.add_heading("4.4 Module 2: intrinsic difficulty", level=2)
    add_body(
        doc,
        "Twenty-eight policy-independent kinematic and contact features are computed from complete trajectories "
        "before segment aggregation. Train-only robust scaling uses median/MAD with a P05-P95 fallback, followed "
        "by empirical-CDF mapping and ten near-balanced bins. Difficulty provides a calibration reference for "
        "learning-gap variants; it neither rejects nor directly samples trajectories.",
    )
    doc.add_heading("4.5 Module 3: online error and learning state", level=2)
    add_body(
        doc,
        "Tracking and traversal signals are attributed to the active motion and segment and maintained with "
        "exponential moving averages. Raw-error variants directly prioritize current errors. Learning-gap variants "
        "compare them with same-difficulty peers. A 0.15 uniform mixture, undersampling bonuses, and probability "
        "caps preserve exploration and coverage.",
    )
    doc.add_heading("4.6 Module 4: diversity constraint", level=2)
    add_body(
        doc,
        "Seventeen segment descriptors are aggregated into 30 motion features and clustered using train-only "
        "K-means++ (K=8). Source labels and evaluation categories are excluded. Cluster probability depends on "
        "eligible cluster size, while error operates only within a cluster.",
    )
    add_equation_block(doc, "P(c,m,s) = P_div(c) P_err(m | c) P_err(s | m)")
    add_equation_block(doc, "P_div(c) = f/C + (1-f) n_c^alpha / sum_j n_j^alpha, with f=0.5 and alpha=0.5")
    add_figure(doc, artwork["Fig_2_Training_Metadata"]["preview"], FIGURE_CAPTIONS["Fig_2_Training_Metadata"])
    doc.add_heading("4.7 Proposed method and baselines", level=2)
    add_body(
        doc,
        "QD-HES is M7-Raw: quality on, diversity on, hierarchical raw motion error, hierarchical raw segment "
        "error, and learning-gap calibration off. D-only is the strongest direct baseline and differs only in the "
        "quality switch. GlobalRaw prioritizes segments globally without a cluster budget.",
    )
    switch_orientation(doc, landscape=True)
    doc.add_heading("4.8 Method matrix", level=2)
    headers, rows = tables["Table_04_Method_Matrix"]
    add_kbs_table(doc, headers, rows, font_size=7.1, bold_row=lambda row: "QD-HES" in str(row[0]))
    switch_orientation(doc, landscape=False)

    doc.add_heading("5. Experimental Protocol", level=1)
    doc.add_heading("5.1 Fixed partitions", level=2)
    headers, rows = tables["Table_03_Data_Splits"]
    add_kbs_table(doc, headers, rows, font_size=7.1)
    add_body(
        doc,
        "Train, full Validation, and Test are source-group disjoint. The 500-motion probe is contained within "
        "Validation. It screens candidate checkpoints but does not replace full Validation. The final candidates "
        "were frozen from Validation rankings before the current Test comparison.",
    )
    doc.add_heading("5.2 PPO, observations, actions, and evaluation", level=2)
    headers, rows = tables["Table_05_PPO_Configuration"]
    add_kbs_table(doc, headers, rows, font_size=8.0)
    doc.add_heading("5.3 Reward and termination", level=2)
    headers, rows = tables["Table_06_Reward_Terms"]
    add_kbs_table(doc, headers, rows, font_size=8.0)
    add_body(
        doc,
        "Early termination thresholds are 0.25 m for anchor height, 0.8 rad for anchor orientation, and 0.25 m "
        "for wrist/ankle end-body height. All methods share the network, reward, termination logic, randomization "
        "state, deterministic evaluator, and frame-zero start protocol.",
    )
    doc.add_heading("5.4 Metrics and macro-first selection", level=2)
    headers, rows = tables["Table_07_Metric_Definitions"]
    add_kbs_table(doc, headers, rows, font_size=8.0)
    add_body(
        doc,
        "Macro success is the primary selection metric because categories are imbalanced. Within epsilon=0.002, "
        "ties are broken by micro success, completion, body error, and earlier checkpoint. Joint L2 is not used for "
        "selection. Error means are interpreted only alongside success and completion because early termination can "
        "create deceptively low errors.",
    )
    doc.add_heading("5.5 Statistical protocol", level=2)
    add_body(
        doc,
        "Every Test method is evaluated on the same ordered 7,592-motion manifest, enabling paired differences. "
        "Primary 95% intervals use 10,000 source-group cluster-bootstrap resamples (3,610 groups, seed 42). Macro "
        "success is resampled within category and source group. Motion-level bootstrap results are retained as a "
        "sensitivity analysis. These intervals do not capture training-seed variation.",
    )

    switch_orientation(doc, landscape=True)
    doc.add_heading("6. Experimental Results", level=1)
    doc.add_heading("6.1 Fixed-budget 34k full Validation", level=2)
    headers, rows = tables["Table_08_Validation_Ablation_34k"]
    add_kbs_table(doc, headers, rows, font_size=6.8, bold_row=lambda row: str(row[0]) == "D-only")
    add_figure(doc, artwork["Fig_3_Validation_Ablation_34k"]["preview"], FIGURE_CAPTIONS["Fig_3_Validation_Ablation_34k"], width=8.9)
    add_body(
        doc,
        "D-only is the 34k macro winner. Compared with GlobalRaw, it gains 2.095 micro points and 1.792 macro "
        "points and removes 160 failures. M7-Raw has slightly better micro/error values but is 0.629 macro points "
        "lower, so the original macro-first D-only decision remains valid.",
    )
    add_note(doc, "M2 and M3 have pilot evidence only. Do not label the table a complete M0-M7 ablation unless their 34k full-Validation evaluations are added.")
    doc.add_heading("6.2 Budget scaling", level=2)
    headers, rows = tables["Table_09_Validation_Budget_Scaling"]
    add_kbs_table(doc, headers, rows, font_size=6.8, bold_row=lambda row: str(row[0]) == "M7-Raw" and "59k" in str(row[1]))
    add_figure(doc, artwork["Fig_4_Validation_Budget_Scaling"]["preview"], FIGURE_CAPTIONS["Fig_4_Validation_Budget_Scaling"], width=8.9)
    add_body(
        doc,
        "The result is budget-dependent. M7-Raw improves through the 59k Validation optimum and then degrades. "
        "D-only shows little macro benefit from additional training. The comparison supports a short-budget "
        "diversity benefit and a longer-horizon quality-gate benefit, not a claim that longer training is always better.",
    )
    doc.add_heading("6.3 Final Test", level=2)
    headers, rows = tables["Table_10_Final_Test"]
    add_kbs_table(doc, headers, rows, font_size=6.8, bold_row=lambda row: "M7-Raw" in str(row[0]))
    add_figure(doc, artwork["Fig_5_Final_Test_Comparison"]["preview"], FIGURE_CAPTIONS["Fig_5_Final_Test_Comparison"], width=8.9)
    add_body(
        doc,
        f"QD-HES reaches micro={float(m7['micro']):.6f}, macro={float(m7['macro']):.6f}, and "
        f"completion={float(m7['completion']):.6f}. Relative to D-only, body error decreases by "
        f"{(float(donly['body'])-float(m7['body']))/float(donly['body'])*100:.2f}%, joint L2 by "
        f"{(float(donly['joint'])-float(m7['joint']))/float(donly['joint'])*100:.2f}%, and failures from "
        f"{int(donly['failures'])} to {int(m7['failures'])}.",
    )
    doc.add_heading("6.4 Paired Test analysis", level=2)
    headers, rows = tables["Table_11_Final_Test_Paired_Deltas"]
    add_kbs_table(doc, headers, rows, font_size=7.0)
    add_body(
        doc,
        f"Against D-only, the primary source-group bootstrap gives micro "
        f"{bootstrap_interval(boot_d, 'micro', scale=100)} percentage points and macro "
        f"{bootstrap_interval(boot_d, 'macro', scale=100)} points. M7-Raw uniquely solves "
        f"{int(paired_d['a_only'])} motions, whereas D-only uniquely solves {int(paired_d['b_only'])}. "
        f"Both solve {int(paired_d['both'])} and neither solves {int(paired_d['neither'])}.",
    )
    switch_orientation(doc, landscape=False)

    doc.add_heading("7. Category, Failure, and Stability Analysis", level=1)
    add_figure(doc, artwork["Fig_6_Final_Test_Category_Deltas"]["preview"], FIGURE_CAPTIONS["Fig_6_Final_Test_Category_Deltas"])
    add_body(
        doc,
        f"M7-Raw improves {int(wtl['wins'])} of 17 categories, ties {int(wtl['ties'])}, and declines in "
        f"{int(wtl['losses'])}. Category differences with small sample sizes are descriptive and should not be "
        "presented as independent headline findings.",
    )
    switch_orientation(doc, landscape=True)
    headers, rows = tables["Table_13_Final_Test_Categories"]
    add_kbs_table(doc, headers, rows, font_size=6.8)
    switch_orientation(doc, landscape=False)
    doc.add_heading("7.1 Termination reasons", level=2)
    headers, rows = tables["Table_14_Termination_Reasons"]
    add_kbs_table(doc, headers, rows, font_size=8.0)
    add_body(
        doc,
        "End-body position is the dominant residual failure mode and accounts for most of the net improvement. "
        "Qualitative figures and videos should therefore include wrist/ankle-intensive actions in addition to "
        "locomotion examples.",
    )
    doc.add_heading("7.2 Dense checkpoint diagnostic", level=2)
    add_figure(doc, artwork["Fig_S1_Dense_Validation_Probe"]["preview"], FIGURE_CAPTIONS["Fig_S1_Dense_Validation_Probe"])
    add_body(
        doc,
        "The 500-motion probe shows a severe collapse around 62k-67k followed by partial recovery. This rules out "
        "a monotonic convergence interpretation and makes independent Validation selection a substantive part of "
        "the experimental design.",
    )

    doc.add_heading("8. Discussion Draft", level=1)
    for heading, body in [
        (
            "8.1 Short-budget diversity benefit",
            "Raw error responds quickly to policy competence, while the frozen cluster budget prevents a large family of similar high-error motions from consuming most rollouts. The common-budget result supports this interpretation because D-only exceeds both GlobalRaw and M4.",
        ),
        (
            "8.2 Slow-onset quality benefit",
            "The gate rejects only 90 of 21,575 segments, so its effect per assignment is small. Across billions of transitions, avoiding repeated starts in unreliable regions can accumulate into a meaningful benefit. This is consistent with the observed crossover but remains a mechanism hypothesis until repeated seeds or sampler-usage traces are available.",
        ),
        (
            "8.3 Negative learning-gap result",
            "Difficulty calibration and relative-gap estimation add delay and noise. M5-M7 do not outperform simpler raw-error variants at 34k, and the joint-gap probe is dominated. The paper should retain this result because it shows that adding knowledge signals without reliable estimators can be counterproductive.",
        ),
        (
            "8.4 Non-monotonic optimization",
            "PPO and the changing sampling distribution form a coupled non-stationary process. The 59k optimum, 60k decline, and later probe collapse show why the final checkpoint should be selected by a predeclared Validation rule rather than by training duration or Test performance.",
        ),
    ]:
        doc.add_heading(heading, level=2)
        add_body(doc, body)

    doc.add_heading("9. Conclusion Draft", level=1)
    add_body(
        doc,
        "This work presents a knowledge-guided hierarchical sampler for training one humanoid whole-body "
        "tracking policy on a large heterogeneous motion library. The method separates offline trajectory quality, "
        "train-only motion diversity, and online policy competence and composes them into a cluster-motion-segment "
        "distribution. Diversity-constrained hierarchical raw-error sampling is strongest at the common 34k "
        "budget. With longer Validation-selected training, adding the quality gate yields a higher optimum: QD-HES "
        "at 59k reaches 0.9231 micro success and 0.9079 macro success on 7,592 held-out Test motions and reduces "
        "failures from 731 to 584 relative to D-only. Performance degradation after the optimum emphasizes that "
        "structured rollout allocation and independent checkpoint selection are both necessary for reliable "
        "large-scale humanoid motion tracking.",
    )

    doc.add_heading("10. Limitations and Required Additions", level=1)
    for item in [
        "Only training seed 42 is available. Add at least two seeds for QD-HES, D-only, and GlobalRaw before making seed-general claims.",
        "M2 and M3 lack 7,636-motion full Validation. Add them only if the paper explicitly claims a complete M0-M7 ablation.",
        "The Test split is held out from training and Validation but was accessed in an older 34k protocol; disclose this and prohibit Test-driven reselection.",
        "Current evidence is simulation-only and does not establish real-robot or sim-to-real performance.",
        "Resolve remaining provisional metadata-profile labels and archive the final configuration/profile hashes.",
        "Record GPU model, wall-clock time, and compute cost for the final methods.",
        "Add representative success, rescued-failure, and persistent-failure frames or videos, especially for end-body-intensive motions.",
        "Do not claim state of the art without compatible external baseline evaluations or explicit protocol caveats.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("11. Manuscript and Artwork Plan", level=1)
    outline = [
        ["1. Introduction", "Problem, KBS framing, research gap, contributions"],
        ["2. Related work", "Humanoid imitation; scalable tracking; curriculum/prioritization; data quality/diversity"],
        ["3. Problem formulation", "Tracking MDP, segmented library, knowledge sources"],
        ["4. Method", "Stage 0, quality gate, online error, diversity budget, QD-HES algorithm"],
        ["5. Experiments", "Data, G1/PPO, baselines, metrics, checkpoint and bootstrap rules"],
        ["6. Results", "34k ablation, budget scaling, Test, categories and failures"],
        ["7. Discussion", "Budget dependence, negative results, instability, limitations"],
        ["8. Conclusion", "Findings and scope"],
    ]
    add_kbs_table(doc, ["Section", "Contents"], outline, font_size=8.5)
    for name, caption in FIGURE_CAPTIONS.items():
        add_bullet(doc, f"{name}: {caption}")

    doc.add_heading("12. Reproducibility and Integrity Record", level=1)
    hashes = [
        ["Train manifest", SPLIT_ROWS[0][3]],
        ["Validation manifest", SPLIT_ROWS[1][3]],
        ["Test manifest", SPLIT_ROWS[3][3]],
        ["Quality metadata", stats["quality_metadata_sha256"]],
        ["Difficulty metadata", stats["difficulty_metadata_sha256"]],
        ["Cluster metadata", stats["cluster_metadata_sha256"]],
        ["Evaluator commit", m7["payload"]["git_commit"]],
        ["QD-HES 59k checkpoint", m7["payload"]["checkpoint_sha256"]],
        ["D-only 54k checkpoint", donly["payload"]["checkpoint_sha256"]],
        ["GlobalRaw 34k checkpoint", global_raw["payload"]["checkpoint_sha256"]],
    ]
    add_kbs_table(doc, ["Artifact", "SHA256 / commit"], hashes, font_size=6.8)
    add_note(
        doc,
        "Supported claim: QD-HES is the best evaluated candidate on the current Validation and held-out Test "
        "splits. Unsupported without more evidence: seed-independent superiority, complete M0-M7 ablation, real-"
        "robot transfer, external state of the art, or a never-accessed blind Test.",
    )

    doc.add_heading("Appendix A. Supplementary and Negative Results", level=1)
    headers, rows = tables["Table_S1_Diagnostic_Full_Validation"]
    add_kbs_table(doc, headers, rows, font_size=6.9)
    headers, rows = tables["Table_S2_Joint_Gap_Probe"]
    add_kbs_table(doc, headers, rows, font_size=7.0)
    add_body(
        doc,
        "The formal-v2 diagnostic combinations remain below D-only at the common budget. The joint-gap pilot is "
        "a 10k Probe500 experiment and must not be compared as if it were a full 34k Validation result.",
    )

    configure_headers(doc)
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def build_highlights_docx() -> None:
    doc = Document()
    configure_document(doc)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Highlights")
    set_run_style(run, size=18, bold=True)
    for item in HIGHLIGHTS:
        add_bullet(doc, item)
    doc.save(HIGHLIGHTS_DOCX_PATH)
    HIGHLIGHTS_TXT_PATH.write_text("\n".join(f"- {item}" for item in HIGHLIGHTS) + "\n", encoding="utf-8")


def write_captions() -> None:
    CAPTIONS_PATH.write_text("\n\n".join(FIGURE_CAPTIONS.values()) + "\n", encoding="utf-8")


def write_algorithm() -> None:
    content = r"""% Requires \usepackage{algorithm} and \usepackage{algpseudocode}.
\begin{algorithm}[t]
\caption{Quality-Gated Diversity-Constrained Hierarchical Error Sampling (QD-HES)}
\label{alg:qdhes}
\begin{algorithmic}[1]
\Require Training motions $\mathcal{M}$; frozen quality labels $q_{m,s}$; frozen clusters $c(m)$
\State Build one-second segment index and valid start-frame masks
\State Mark rejected segments ineligible as episode starts
\State Initialize motion/segment error EMAs and assignment counters
\For{PPO iteration $t=1,\ldots,T$}
  \If{$t$ is after warm-up and is a probability-update iteration}
    \State Update raw-error scores from tracking, completion, and termination statistics
    \State Compute diversity budget $P_{\mathrm{div}}(c)$ from eligible cluster sizes
    \State Compute $P_{\mathrm{err}}(m\mid c)$ with uniform floor, coverage bonus, and caps
    \State Compute $P_{\mathrm{err}}(s\mid m)$ over quality-eligible segments
  \EndIf
  \For{each environment reset}
    \State Sample $c \sim P_{\mathrm{div}}(c)$
    \State Sample $m \sim P_{\mathrm{err}}(m\mid c)$
    \State Sample $s \sim P_{\mathrm{err}}(s\mid m)$ and a valid frame inside $s$
  \EndFor
  \State Collect on-policy rollouts and update PPO
  \State Attribute observed errors and outcomes to active $(m,s)$ entries
\EndFor
\end{algorithmic}
\end{algorithm}
"""
    ALGORITHM_PATH.write_text(content, encoding="utf-8")


def write_references() -> None:
    content = r"""% Verified seed references as of 2026-09-08. Expand before submission.
@article{peng2018deepmimic,
  author  = {Peng, Xue Bin and Abbeel, Pieter and Levine, Sergey and van de Panne, Michiel},
  title   = {DeepMimic: Example-Guided Deep Reinforcement Learning of Physics-Based Character Skills},
  journal = {ACM Transactions on Graphics},
  volume  = {37},
  number  = {4},
  pages   = {1--14},
  year    = {2018},
  doi     = {10.1145/3197517.3201311}
}

@article{schulman2017ppo,
  author  = {Schulman, John and Wolski, Filip and Dhariwal, Prafulla and Radford, Alec and Klimov, Oleg},
  title   = {Proximal Policy Optimization Algorithms},
  journal = {arXiv preprint arXiv:1707.06347},
  year    = {2017},
  doi     = {10.48550/arXiv.1707.06347}
}

@inproceedings{arthur2007kmeanspp,
  author    = {Arthur, David and Vassilvitskii, Sergei},
  title     = {k-means++: The Advantages of Careful Seeding},
  booktitle = {Proceedings of the Eighteenth Annual ACM-SIAM Symposium on Discrete Algorithms},
  pages     = {1027--1035},
  year      = {2007}
}

@inproceedings{bengio2009curriculum,
  author    = {Bengio, Yoshua and Louradour, Jerome and Collobert, Ronan and Weston, Jason},
  title     = {Curriculum Learning},
  booktitle = {Proceedings of the 26th International Conference on Machine Learning},
  pages     = {41--48},
  year      = {2009},
  doi       = {10.1145/1553374.1553380}
}

@inproceedings{schaul2016per,
  author    = {Schaul, Tom and Quan, John and Antonoglou, Ioannis and Silver, David},
  title     = {Prioritized Experience Replay},
  booktitle = {International Conference on Learning Representations},
  year      = {2016},
  eprint    = {1511.05952}
}

@inproceedings{rudin2022minutes,
  author    = {Rudin, Nikita and Hoeller, David and Reist, Philipp and Hutter, Marco},
  title     = {Learning to Walk in Minutes Using Massively Parallel Deep Reinforcement Learning},
  booktitle = {Conference on Robot Learning},
  pages     = {91--100},
  year      = {2022}
}

@article{liao2025beyondmimic,
  author  = {Liao, Qiayuan and Truong, Takara E. and Huang, Xiaoyu and Gao, Yuman and Tevet, Guy and Sreenath, Koushil and Liu, C. Karen},
  title   = {BeyondMimic: From Motion Tracking to Versatile Humanoid Control via Guided Diffusion},
  journal = {arXiv preprint arXiv:2508.08241},
  year    = {2025},
  doi     = {10.48550/arXiv.2508.08241}
}

@article{chen2025gmt,
  author  = {Chen, Zixuan and Ji, Mazeyu and Cheng, Xuxin and Peng, Xuanbin and Peng, Xue Bin and Wang, Xiaolong},
  title   = {GMT: General Motion Tracking for Humanoid Whole-Body Control},
  journal = {arXiv preprint arXiv:2506.14770},
  year    = {2025},
  doi     = {10.48550/arXiv.2506.14770}
}

@article{lee2025phuma,
  author  = {Lee, Kyungmin and Kim, Sibeen and Lee, Youngdo and Park, Minho and Kim, Hyunseung and Hwang, Dongyoon and Kim, Donghu and Lee, Hojoon and Choo, Jaegul},
  title   = {PHUMA: Physically Reliable Humanoid Locomotion Dataset},
  journal = {arXiv preprint arXiv:2510.26236},
  year    = {2025},
  doi     = {10.48550/arXiv.2510.26236}
}

@misc{nvidia2025isaaclab,
  author       = {{NVIDIA}},
  title        = {Isaac Lab: Unified Framework for Robot Learning},
  year         = {2025},
  howpublished = {\url{https://isaac-sim.github.io/IsaacLab/}},
  note         = {Software documentation; verify the release citation before submission}
}
"""
    BIB_PATH.write_text(content, encoding="utf-8")


def write_cover_letter() -> None:
    content = f"""# Cover Letter Draft

Dear Editor-in-Chief,

Please consider our manuscript, **\"{PAPER_TITLE}\"**, for publication in *Knowledge-Based Systems*.

The manuscript studies a knowledge-guided data-allocation problem in large-scale humanoid reinforcement learning. It introduces {SHORT_METHOD_NAME}, which explicitly represents three complementary sources of knowledge: offline trajectory reliability, train-only motion diversity, and online policy competence. These sources are separated by responsibility and composed into a cluster-motion-segment sampling distribution for on-policy PPO training.

The evaluation uses 76,086 converted PHUMA trajectories. Under a common 34k-iteration budget, diversity-constrained hierarchical raw-error sampling is the strongest short-budget variant. Validation-based budget scaling selects {SHORT_METHOD_NAME} at 59k iterations. On a 7,592-motion held-out Test split, it achieves 0.9231 micro success, 0.9079 macro success, and 0.9554 completion, reducing failures from 731 to 584 relative to the strongest diversity-only baseline. The study also identifies severe non-monotonic checkpoint behavior, showing why structured sampling and independent Validation selection must be considered jointly.

We believe the manuscript fits *Knowledge-Based Systems* because its principal contribution is an interpretable machine-learning methodology for integrating heterogeneous offline and online knowledge into data-driven optimization, with humanoid whole-body tracking as a demanding large-scale application.

This manuscript is original, is not under consideration elsewhere, and has been approved by all authors. [VERIFY/EDIT THIS SENTENCE.] All authors' conflicts of interest, funding, data/code availability, and CRediT roles are stated in the submission files.

Suggested reviewers: [ADD 3-5 REVIEWERS WITH INSTITUTIONAL EMAILS AND NO CONFLICTS]

Sincerely,  
[CORRESPONDING AUTHOR]  
[AFFILIATION]  
[EMAIL]
"""
    COVER_LETTER_PATH.write_text(content, encoding="utf-8")


def write_declarations() -> None:
    content = """# Declarations Template

## CRediT authorship contribution statement

[Author 1]: Conceptualization, Methodology, Software, Validation, Formal analysis, Investigation, Data curation, Visualization, Writing - original draft.  
[Author 2]: Supervision, Methodology, Resources, Writing - review and editing.  
[Edit roles to match actual contributions.]

## Funding

This work was supported by [FUNDER, GRANT NUMBER]. / This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.

## Declaration of competing interest

The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper. [VERIFY]

## Data availability

The experiments use the PHUMA motion corpus after local conversion to the Unitree G1/WBT representation. Fixed Train, Validation, and Test manifests and their SHA256 identities are reported in the manuscript. Redistribution of converted trajectories will follow the upstream dataset license. [ADD PUBLIC ARCHIVE DOI/URL OR EXPLAIN RESTRICTION.]

## Code availability

Training, metadata construction, adaptive sampling, deterministic evaluation, and analysis code will be released at [REPOSITORY/ARCHIVE URL] with the final configuration files and artifact hashes. [UPDATE BEFORE SUBMISSION.]

## Ethics statement

The reported experiments use existing motion data and physics simulation; no new human-subject or animal experiment was conducted. [VERIFY AGAINST THE ACTUAL PROJECT.]

## Use of generative AI in scientific writing

[Follow the journal's current disclosure policy. Describe any language-assistance use accurately; authors remain responsible for all content.]

## Test-set provenance statement

Train, Validation, and Test manifests are source-group disjoint. The 59k QD-HES and 54k D-only checkpoints used in the current comparison were selected using Validation only. The Test split had previously been accessed under an older 34k protocol; therefore, the manuscript describes it as held out from Train/Validation but does not characterize it as a pristine never-accessed blind benchmark. No model or checkpoint may be reselected using the current Test outcomes.
"""
    DECLARATIONS_PATH.write_text(content, encoding="utf-8")


def write_submission_checklist() -> None:
    highlight_rows = "\n".join(f"- [x] {len(item):2d}/85 characters: {item}" for item in HIGHLIGHTS)
    content = f"""# KBS Submission Checklist

Checked on {DATE}. Recheck the live journal guide immediately before submission.

## Journal fit

- [x] The manuscript is framed as knowledge-guided machine-learning methodology and data-driven optimization.
- [x] The application demonstrates large-scale humanoid whole-body tracking rather than serving as the sole novelty claim.
- [ ] Tighten the Introduction against the final related-work review and state the exact knowledge representation/optimization gap.

Official scope: https://shop.elsevier.com/journals/knowledge-based-systems/0950-7051

## Front matter

- [x] English title, abstract, six keywords, and five highlights are supplied.
- [x] Highlights satisfy Elsevier's general 3-5 bullet and 85-character limit:
{highlight_rows}
- [ ] Add author names, affiliations, corresponding-author details, and ORCID IDs.
- [ ] Confirm the final article type and any KBS-specific title-page fields in Editorial Manager.

Official highlights guidance: https://www.elsevier.support/publishing/answer/how-do-i-include-highlights-with-my-manuscript

## Figures and artwork

- [x] Main charts and framework artwork are exported as vector PDF with embedded TrueType fonts.
- [x] Every figure also has a 300-dpi PNG preview for review and Word embedding.
- [x] The graphical abstract is supplied as PDF, 300-dpi TIFF, and PNG preview at a 1328:531 aspect ratio.
- [x] Fonts, line weights, colors, panel labels, units, and captions are consistent.
- [x] The palette is color-vision-friendly and does not rely on color alone for line-series identity.
- [ ] Add representative simulator/robot frames as at least 300-dpi photographs/renders.
- [ ] Inspect every final figure at its intended single- or double-column print size.
- [ ] Use vector PDF for line art. If rasterizing, meet Elsevier's type-specific guidance: line art 1000 dpi, combination art 500 dpi, grayscale/color images 300 dpi.

Official artwork guidance:  
https://www.elsevier.com/en-au/about/policies-and-standards/author/artwork-and-media-instructions/artwork-types  
https://www.prod.webpresence.elsevier.com/about/policies-and-standards/author/artwork-and-media-instructions/artwork-sizing

Official graphical-abstract guidance: https://www.elsevier.com/researcher/author/tools-and-resources/graphical-abstract

## Tables

- [x] All tables are editable CSV; key tables also have LaTeX `booktabs` source.
- [x] The Word report uses horizontal-rule academic tables without cell shading or vertical rules.
- [x] Units and higher/lower-is-better directions are explicit.
- [ ] In the final manuscript, bold the best value and optionally underline the second-best value per metric.
- [ ] Keep only the central tables in the main paper; move detailed categories and dense checkpoints to supplementary material.

## Experimental completeness

- [x] Full Validation: principal 34k methods, budget scaling, and selected checkpoints.
- [x] Final Test: GlobalRaw 34k, D-only 54k, and QD-HES/M7-Raw 59k.
- [x] Paired Test differences and 10,000-resample source-group bootstrap intervals.
- [x] Category-wise and termination-reason analyses.
- [x] Dense 50k-70k checkpoint diagnostic.
- [ ] Add at least two training seeds for QD-HES, D-only, and GlobalRaw; this is the highest-priority scientific gap.
- [ ] Add M2/M3 34k full Validation only if claiming a complete M0-M7 ablation.
- [ ] Record GPU, training wall time, and total compute.
- [ ] Add qualitative success, rescued failure, and persistent failure examples.

## Integrity and reproducibility

- [x] Train/Validation/Test manifests and SHA256 values are reported.
- [x] Test summaries record checkpoint SHA256, evaluator commit, seed, deterministic execution, horizon, and randomization status.
- [x] The 59k/54k checkpoints were selected from Validation and must not be changed after Test inspection.
- [x] The report discloses that an earlier 34k Test protocol existed; it does not call the split pristine or never accessed.
- [ ] Archive final configs, metadata profiles, checkpoints or checkpoint access instructions, code commit, and raw per-motion results in a durable repository.
- [ ] Resolve every remaining `provisional` flag in quality/difficulty/cluster profiles.

## Declarations and submission files

- [x] Cover-letter and declaration templates are included.
- [ ] Complete CRediT roles, funding, competing-interest, data/code availability, and AI-use statements.
- [ ] Verify upstream PHUMA license and redistribution terms for converted data.
- [ ] Add a complete literature review and convert references to the journal's final style.
- [ ] Run spelling, grammar, plagiarism, and reference-consistency checks on the final manuscript.
"""
    CHECKLIST_PATH.write_text(content, encoding="utf-8")


def write_statistics(analysis: dict[str, object]) -> None:
    payload = {
        "generated_at": DATE,
        "primary_resampling_unit": "source_group",
        "resamples": 10000,
        "seed": 42,
        "comparisons": {
            "M7Raw59_vs_Donly54": {
                "source_group_bootstrap": analysis["bootstrap_m7_vs_donly"],
                "motion_bootstrap_sensitivity": analysis["motion_bootstrap_m7_vs_donly"],
                "paired_outcomes": analysis["paired_m7_vs_donly"],
            },
            "M7Raw59_vs_GlobalRaw34": {
                "source_group_bootstrap": analysis["bootstrap_m7_vs_globalraw"],
                "motion_bootstrap_sensitivity": analysis["motion_bootstrap_m7_vs_globalraw"],
                "paired_outcomes": analysis["paired_m7_vs_globalraw"],
            },
        },
        "category_wins_ties_losses_vs_Donly": analysis["category_wins_ties_losses"],
        "interpretation": (
            "Intervals quantify uncertainty over the fixed evaluation corpus. They do not quantify PPO training-seed variance."
        ),
    }
    STATISTICS_PATH.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def write_contact_sheet(artwork: dict[str, dict[str, Path]]) -> Path:
    names = [
        "Fig_1_Overall_Framework",
        "Fig_2_Training_Metadata",
        "Fig_3_Validation_Ablation_34k",
        "Fig_4_Validation_Budget_Scaling",
        "Fig_5_Final_Test_Comparison",
        "Fig_6_Final_Test_Category_Deltas",
        "Fig_S1_Dense_Validation_Probe",
        "Graphical_Abstract",
    ]
    fig, axes = plt.subplots(4, 2, figsize=(12, 14))
    for ax, name in zip(axes.flat, names):
        ax.imshow(plt.imread(artwork[name]["preview"]))
        ax.set_title(name, fontsize=9)
        ax.axis("off")
    fig.tight_layout(pad=0.8)
    path = PACKAGE / "Figure_Contact_Sheet.png"
    fig.savefig(path, dpi=160, bbox_inches="tight")
    plt.close(fig)
    return path


def write_readme() -> None:
    content = f"""# PHUMA-WBT KBS Paper Package

Generated from local frozen artifacts on {DATE}. All manuscript prose is in English.

## Start here

- `PHUMA_WBT_KBS_Paper_Materials.docx`: polished paper-material report with modules, protocol, results, discussion, and limitations.
- `PHUMA_WBT_KBS_Paper_Materials.md`: editable source-like version of the same scientific narrative.
- `KBS_Submission_Checklist.md`: journal-format and experiment-completeness checklist.
- `Figure_Contact_Sheet.png`: quick visual quality-control sheet.

## Submission components

- `Highlights.docx` and `Highlights.txt`: five Elsevier-length highlights.
- `Cover_Letter_Draft.md`: KBS cover-letter draft with author placeholders.
- `Declarations_Template.md`: CRediT, funding, conflicts, availability, ethics, AI-use, and Test-provenance templates.
- `Figure_Captions.txt`: standalone English captions.
- `Graphical_Abstract` files in `figures/`: vector PDF, 300-dpi TIFF, and PNG preview.
- `references_seed.bib`: verified seed bibliography; expand it before submission.
- `Algorithm_QD_HES.tex`: manuscript-ready pseudocode.

## Figures

Every scientific figure is supplied as vector PDF plus a 300-dpi PNG preview. Use the PDFs in the final manuscript. PNGs are for Word review. The graphical abstract additionally has a 300-dpi LZW TIFF.

## Tables and statistics

- `tables/Table_*.csv`: editable data for 14 main and 3 supplementary tables.
- `tables/kbs_tables.tex`: key `booktabs` tables without vertical rules.
- `paired_statistics.json`: source-group bootstrap, motion-level sensitivity bootstrap, and paired win/loss counts.

## Scientific status

The strongest current candidate is QD-HES/M7-Raw `model_59000.pt`, selected on full Validation. It reaches Test micro 0.923077, macro 0.907906, completion 0.955444, body error 0.047820 m, joint L2 0.761526 rad, and 584 failures. The most important missing evidence is repeated training seeds. M2/M3 full Validation is needed only for a complete M0-M7 claim.

The Test split is disjoint from Train and Validation, but an older 34k Test protocol exists. Do not describe it as a never-accessed blind set and do not reselect models after the current Test results.
"""
    README_PATH.write_text(content, encoding="utf-8")


def validate_inputs(
    formal: Sequence[dict[str, object]],
    budget: Sequence[dict[str, object]],
    test: Sequence[dict[str, object]],
    categories: Sequence[dict[str, object]],
    probe: Sequence[dict[str, object]],
) -> None:
    if len(formal) != 10:
        raise RuntimeError(f"Expected 10 full-Validation formal rows, found {len(formal)}")
    if len(categories) != 17:
        raise RuntimeError(f"Expected 17 Test categories, found {len(categories)}")
    if not any(int(row["iteration"]) == 59000 for row in budget):
        raise RuntimeError("Budget results do not contain the selected model_59000.pt")
    if not any(int(row["iteration"]) == 59000 for row in probe):
        raise RuntimeError("Dense probe results do not contain model_59000.pt")
    if any(len(item) > 85 for item in HIGHLIGHTS):
        raise RuntimeError(f"Elsevier highlight exceeds 85 characters: {[len(item) for item in HIGHLIGHTS]}")
    manifest_hashes = {str(row["payload"]["manifest_sha256"]) for row in test}
    if manifest_hashes != {SPLIT_ROWS[3][3]}:
        raise RuntimeError(f"Final Test manifest mismatch: {manifest_hashes}")
    for row in test:
        payload = row["payload"]
        if int(payload["num_motions"]) != 7592:
            raise RuntimeError(f"Incomplete Test result for {row['method']}")
        if not payload["deterministic"] or not payload["disable_randomization"]:
            raise RuntimeError(f"Non-deterministic or randomized Test result for {row['method']}")
        if not payload["confirm_final_test"]:
            raise RuntimeError(f"Missing final-Test confirmation for {row['method']}")
    proposed = next(row for row in test if str(row["method"]).startswith("M7-Raw"))
    if int(proposed["successes"]) != 7008 or int(proposed["failures"]) != 584:
        raise RuntimeError("Unexpected proposed-method Test counts")


def validate_outputs(tables: dict[str, tuple[list[str], list[list[object]]]]) -> None:
    required = [DOCX_PATH, MD_PATH, HIGHLIGHTS_DOCX_PATH, CHECKLIST_PATH, BIB_PATH, STATISTICS_PATH]
    required.extend(FIGURES / f"{name}.pdf" for name in FIGURE_CAPTIONS)
    missing = [str(path) for path in required if not path.exists() or path.stat().st_size == 0]
    if missing:
        raise RuntimeError(f"Missing or empty generated artifacts: {missing}")
    if len(list(TABLES.glob("Table_*.csv"))) != len(tables):
        raise RuntimeError("Generated table count does not match the in-memory table set")
    reopened = Document(DOCX_PATH)
    if len(reopened.paragraphs) < 100 or len(reopened.tables) < 15 or len(reopened.inline_shapes) < 7:
        raise RuntimeError(
            f"DOCX validation failed: paragraphs={len(reopened.paragraphs)}, "
            f"tables={len(reopened.tables)}, figures={len(reopened.inline_shapes)}"
        )
    markdown = MD_PATH.read_text(encoding="utf-8")
    for expected in ["0.923077", "0.907906", "0.955444", "584", "source-group bootstrap"]:
        if expected not in markdown:
            raise RuntimeError(f"Expected paper datum missing from Markdown: {expected}")


def write_manifest() -> None:
    files = []
    for path in sorted(PACKAGE.rglob("*")):
        if not path.is_file() or path == MANIFEST_PATH:
            continue
        files.append(
            {
                "path": str(path.relative_to(PACKAGE)),
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
        )
    payload = {
        "generated_at": DATE,
        "project_root": str(ROOT),
        "target_journal": "Knowledge-Based Systems",
        "language": "English",
        "proposed_method": "QD-HES / M7-Raw model_59000.pt",
        "test_scope": "7592 source-group-disjoint held-out motions; prior 34k Test access disclosed",
        "files": files,
    }
    MANIFEST_PATH.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    PACKAGE.mkdir(parents=True, exist_ok=True)
    TABLES.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    formal, budget, test = build_result_sets()
    categories = category_comparison(test)
    probe = dense_probe_rows()
    stats = load_module_statistics()
    validate_inputs(formal, budget, test, categories, probe)
    tables, analysis = build_paper_tables(formal, budget, test, categories, probe, stats)

    configure_plots()
    artwork = {
        "Fig_1_Overall_Framework": plot_framework(),
        "Fig_2_Training_Metadata": plot_training_metadata(stats),
        "Fig_3_Validation_Ablation_34k": plot_formal_ablation(formal),
        "Fig_4_Validation_Budget_Scaling": plot_budget_scaling(budget),
        "Fig_5_Final_Test_Comparison": plot_final_test(test),
        "Fig_6_Final_Test_Category_Deltas": plot_category_deltas(categories),
        "Fig_S1_Dense_Validation_Probe": plot_dense_probe(probe),
        "Graphical_Abstract": plot_graphical_abstract(test),
    }
    write_contact_sheet(artwork)
    write_table_artifacts(tables)
    MD_PATH.write_text(build_markdown(tables, analysis, categories, stats), encoding="utf-8")
    build_docx(tables, analysis, categories, stats, artwork)
    build_highlights_docx()
    write_captions()
    write_algorithm()
    write_references()
    write_cover_letter()
    write_declarations()
    write_submission_checklist()
    write_statistics(analysis)
    write_readme()
    validate_outputs(tables)
    write_manifest()
    print(f"Wrote KBS package: {PACKAGE}")
    print(f"Main DOCX: {DOCX_PATH}")
    print(f"Tables: {len(tables)}; figure sets: {len(artwork)}")


if __name__ == "__main__":
    main()
