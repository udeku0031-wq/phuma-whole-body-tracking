#!/usr/bin/env python3
"""Build the PHUMA-WBT paper-material package from frozen local results."""

from __future__ import annotations

import csv
import json
import os
from pathlib import Path
from typing import Iterable, Sequence


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
ASSETS = REPORTS / "paper_assets_2026-09-08"
TABLES = ASSETS / "tables"
FIGURES = ASSETS / "figures"
DOCX_PATH = REPORTS / "PHUMA_WBT_论文材料汇总_2026-09-08.docx"
MD_PATH = REPORTS / "PHUMA_WBT_论文材料汇总_2026-09-08.md"

os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".matplotlib-cache"))

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Inches, Pt, RGBColor  # noqa: E402


DATE = "2026-09-08"
PROJECT = "PHUMA × Whole-Body Tracking"
CN_FONT_PATH = "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"
CN_FONT_NAME = "Noto Sans CJK SC"


FORMAL_RESULTS = [
    {
        "method": "M0 / Uniform",
        "checkpoint": "model_33999.pt",
        "micro": 0.861053,
        "macro": 0.878095,
        "completion": 0.918896,
        "body": 0.050972,
        "joint": 0.797446,
        "rms": 0.148082,
        "failures": 1061,
    },
    {
        "method": "Q-only / M1",
        "checkpoint": "model_33999.pt",
        "micro": 0.864720,
        "macro": 0.875384,
        "completion": 0.922504,
        "body": 0.051294,
        "joint": 0.828186,
        "rms": 0.153790,
        "failures": 1033,
    },
    {
        "method": "GlobalRaw",
        "checkpoint": "model_33500.pt",
        "micro": 0.863541,
        "macro": 0.881403,
        "completion": 0.922561,
        "body": 0.051413,
        "joint": 0.829823,
        "rms": 0.154094,
        "failures": 1042,
    },
    {
        "method": "GlobalRaw-Q",
        "checkpoint": "model_33999.pt",
        "micro": 0.864065,
        "macro": 0.876390,
        "completion": 0.921520,
        "body": 0.050989,
        "joint": 0.800883,
        "rms": 0.148720,
        "failures": 1038,
    },
    {
        "method": "M4 / Hierarchical Raw",
        "checkpoint": "model_33999.pt",
        "micro": 0.867208,
        "macro": 0.883471,
        "completion": 0.923367,
        "body": 0.052126,
        "joint": 0.845049,
        "rms": 0.156922,
        "failures": 1014,
    },
    {
        "method": "M5 / Learning Gap",
        "checkpoint": "model_25000.pt",
        "micro": 0.853195,
        "macro": 0.871400,
        "completion": 0.915871,
        "body": 0.052667,
        "joint": 0.868623,
        "rms": 0.161299,
        "failures": 1121,
    },
    {
        "method": "M6 / Quality + Gap",
        "checkpoint": "model_33500.pt",
        "micro": 0.861446,
        "macro": 0.877224,
        "completion": 0.922075,
        "body": 0.051718,
        "joint": 0.835751,
        "rms": 0.155195,
        "failures": 1058,
    },
    {
        "method": "M7 / Full Gap + Diversity",
        "checkpoint": "model_33999.pt",
        "micro": 0.873625,
        "macro": 0.883871,
        "completion": 0.928695,
        "body": 0.052157,
        "joint": 0.840422,
        "rms": 0.156062,
        "failures": 965,
    },
    {
        "method": "D-only",
        "checkpoint": "model_33500.pt",
        "micro": 0.884494,
        "macro": 0.899319,
        "completion": 0.934070,
        "body": 0.050515,
        "joint": 0.815035,
        "rms": 0.151348,
        "failures": 882,
    },
    {
        "method": "M7-Raw",
        "checkpoint": "model_33999.pt",
        "micro": 0.884887,
        "macro": 0.893026,
        "completion": 0.934052,
        "body": 0.049918,
        "joint": 0.808584,
        "rms": 0.150150,
        "failures": 879,
    },
]


BUDGET_RESULTS = [
    {
        "method": "D-only",
        "budget": "34k",
        "iteration": 33500,
        "checkpoint": "model_33500.pt",
        "micro": 0.884494,
        "macro": 0.899319,
        "completion": 0.934070,
        "body": 0.050515,
        "joint": 0.815035,
        "rms": 0.151348,
        "failures": 882,
        "phase": "formal/frozen",
    },
    {
        "method": "D-only",
        "budget": "50k",
        "iteration": 49999,
        "checkpoint": "model_49999.pt",
        "micro": 0.887245,
        "macro": 0.887509,
        "completion": 0.934893,
        "body": 0.049836,
        "joint": 0.861048,
        "rms": 0.159893,
        "failures": 861,
        "phase": "post-freeze exploratory",
    },
    {
        "method": "D-only",
        "budget": "54k selected",
        "iteration": 54000,
        "checkpoint": "model_54000.pt",
        "micro": 0.892090,
        "macro": 0.898413,
        "completion": 0.937964,
        "body": 0.050247,
        "joint": 0.823214,
        "rms": 0.152867,
        "failures": 824,
        "phase": "post-freeze exploratory",
    },
    {
        "method": "M7-Raw",
        "budget": "34k",
        "iteration": 33999,
        "checkpoint": "model_33999.pt",
        "micro": 0.884887,
        "macro": 0.893026,
        "completion": 0.934052,
        "body": 0.049918,
        "joint": 0.808584,
        "rms": 0.150150,
        "failures": 879,
        "phase": "formal/frozen",
    },
    {
        "method": "M7-Raw",
        "budget": "45k",
        "iteration": 45000,
        "checkpoint": "model_45000.pt",
        "micro": 0.896674,
        "macro": 0.900510,
        "completion": 0.941407,
        "body": 0.049900,
        "joint": 0.805935,
        "rms": 0.149658,
        "failures": 789,
        "phase": "post-freeze exploratory",
    },
    {
        "method": "M7-Raw",
        "budget": "50k",
        "iteration": 49999,
        "checkpoint": "model_49999.pt",
        "micro": 0.905055,
        "macro": 0.905345,
        "completion": 0.946890,
        "body": 0.049218,
        "joint": 0.784055,
        "rms": 0.145595,
        "failures": 725,
        "phase": "post-freeze exploratory",
    },
    {
        "method": "M7-Raw",
        "budget": "59k selected",
        "iteration": 59000,
        "checkpoint": "model_59000.pt",
        "micro": 0.916841,
        "macro": 0.913028,
        "completion": 0.952066,
        "body": 0.047723,
        "joint": 0.762033,
        "rms": 0.141506,
        "failures": 635,
        "phase": "post-freeze exploratory",
    },
    {
        "method": "M7-Raw",
        "budget": "60k",
        "iteration": 60000,
        "checkpoint": "model_60000.pt",
        "micro": 0.916056,
        "macro": 0.899619,
        "completion": 0.951594,
        "body": 0.048846,
        "joint": 0.793149,
        "rms": 0.147284,
        "failures": 641,
        "phase": "post-freeze exploratory",
    },
    {
        "method": "M7-Raw",
        "budget": "70k",
        "iteration": 69999,
        "checkpoint": "model_69999.pt",
        "micro": 0.895495,
        "macro": 0.898613,
        "completion": 0.939556,
        "body": 0.049019,
        "joint": 0.815119,
        "rms": 0.151364,
        "failures": 798,
        "phase": "post-freeze exploratory",
    },
]


SUPPLEMENTARY_FULL_RESULTS = [
    {
        "method": "V2-Diag-A",
        "design": "Quality + Difficulty + Diversity; motion raw / segment relative gap",
        "budget": "34k",
        "iteration": 25000,
        "checkpoint": "model_25000.pt",
        "micro": 0.8748035620743845,
        "macro": 0.8864142352941177,
        "completion": 0.9281847887637503,
        "body": 0.05134583826610765,
        "joint": 0.7948097173913019,
        "rms": 0.14759245809324253,
        "failures": 956,
        "source": "evaluations/formal_v2/A_raw_motion_gap_segment_seed42/validation_full/model_25000/per_motion.csv",
    },
    {
        "method": "M7-v2 (lambda=0.10)",
        "design": "Quality + Difficulty + Diversity; raw segment error with gap correction",
        "budget": "34k",
        "iteration": 33500,
        "checkpoint": "model_33500.pt",
        "micro": 0.8804347826086957,
        "macro": 0.8834755294117645,
        "completion": 0.9318165838135152,
        "body": 0.050446924305919454,
        "joint": 0.8216253158721831,
        "rms": 0.15257199711891117,
        "failures": 913,
        "source": "evaluations/formal_v2/M7_lambda0p10_seed42/validation_full/model_33500/per_motion.csv",
    },
]


JOINT_GAP_PROBE_RESULTS = [
    {
        "method": "M7-Raw",
        "budget": "10k",
        "checkpoint": "model_10000.pt",
        "micro": 0.828,
        "macro": 0.8398408823529412,
        "completion": 0.8986059120000002,
        "body": 0.053398931999999996,
        "joint": 0.8394617039999999,
        "rms": 0.15588412199999993,
        "failures": 86,
    },
    {
        "method": "M7-JGap (lambda_joint=0.025)",
        "budget": "10k",
        "checkpoint": "model_9999.pt",
        "micro": 0.812,
        "macro": 0.844251411764706,
        "completion": 0.8854216979999999,
        "body": 0.05370323600000002,
        "joint": 0.8483959819999993,
        "rms": 0.15754317799999998,
        "failures": 94,
    },
]


METHOD_MATRIX = [
    ["M0", "Uniform", "Uniform", "Off", "Off", "Off", "完整 Full Val"],
    ["M1 / Q-only", "Uniform", "Uniform", "On", "Off", "Off", "完整 Full Val"],
    ["M2", "Raw error", "Uniform", "Off", "Off", "Off", "仅 pilot；无完整 Full Val"],
    ["M3", "Uniform", "Raw error", "Off", "Off", "Off", "仅 pilot；无完整 Full Val"],
    ["M4", "Raw error", "Raw error", "Off", "Off", "Off", "完整 Full Val"],
    ["M5", "Learning gap", "Relative gap", "Off", "On", "Off", "完整 Full Val"],
    ["M6", "Learning gap", "Relative gap", "On", "On", "Off", "完整 Full Val"],
    ["M7", "Learning gap", "Relative gap", "On", "On", "On", "完整 Full Val"],
    ["GlobalRaw", "全局 segment raw", "扁平抽样", "Off", "Off", "Off", "完整 Full Val"],
    ["GlobalRaw-Q", "全局 segment raw", "扁平抽样", "On", "Off", "Off", "完整 Full Val"],
    ["D-only", "Raw error", "Raw error", "Off", "Off", "On", "34k/50k/54k"],
    ["M7-Raw", "Raw error", "Raw error", "On", "Off", "On", "34k/45k/50k/59k/60k/70k"],
    ["V2-Diag-A", "Raw error", "Relative gap", "On", "On", "On", "补充 34k Full Val"],
    ["M7-v2 lambda=0.10", "Raw error", "Raw + gap correction", "On", "On", "On", "补充 34k Full Val"],
    ["M7-JGap", "Raw error", "Raw + joint-gap correction", "On", "On", "On", "10k Probe; NO-GO"],
]


MODULE_TABLE = [
    [
        "数据转换与多动作库",
        "把 PHUMA G1 .npy 转为 WBT .npz；支持单文件、目录和 manifest；每个环境独立选择 motion 与起始帧。",
        "统一 29-DOF/30-body schema，保留来源信息，建立可扩展 multi-motion tracking 输入。",
    ],
    [
        "Stage 0：Segment 基础设施",
        "按 motion FPS 建立默认 1 s 固定片段，维护 motion/local/global segment 映射、采样统计、概率校验和 checkpoint 恢复。",
        "给所有采样策略提供同一索引与状态接口，并保持 M0 legacy uniform RNG 路径。",
    ],
    [
        "模块 1：Quality Gate",
        "离线检测非有限值、限位、速度一致性、跳变、穿地和足滑，产生 pass/borderline/reject；训练时只屏蔽 reject 起始片段。",
        "把数据缺陷与动作难度分开，避免明显异常轨迹成为 episode assignment 起点。",
    ],
    [
        "模块 2：Intrinsic Difficulty",
        "从最终 G1 轨迹提取 28 项运动学/接触特征，经 Train-only robust 标准化、经验 CDF 和 10 个等频 bin 得到固有难度。",
        "提供策略无关的同难度校准基准；不直接删除样本，也不直接分配训练预算。",
    ],
    [
        "模块 3：Online Error / Learning Gap",
        "按当前 reference frame 归因 body/joint/orientation error 与 traversal outcome，维护 EMA，形成 motion raw error、segment raw error 和 learning gap。",
        "将训练预算集中到当前策略尚未掌握的 motion/segment，同时保留 uniform floor 和欠采样奖励。",
    ],
    [
        "模块 4：Diversity Constraint",
        "用 17 项 segment source feature 聚合成 30 项 motion feature，Train-only KMeans++ 得到 K=8 运动簇，并实行 cluster budget。",
        "防止高误差但相似的动作长期垄断训练；形成 Cluster→Motion→Segment→Start Frame 三级采样。",
    ],
    [
        "PPO 跟踪与评估",
        "Isaac Lab/RSL-RL PPO 训练 29 维关节位置动作；确定性 evaluator 从第 0 帧完整播放每条动作并输出逐 motion、类别和总表。",
        "保持不同 sampler 的网络、reward、termination 和评估器一致，支持 Probe→Full Validation→冻结 Test。",
    ],
]


SPLITS = [
    [
        "Train",
        "random6000_seed42.txt",
        "6000",
        "51f592792f412c5d5e31caf3621162a3b2ef356e05f98dbce43153d08a59e82d",
    ],
    [
        "Validation full",
        "validation_full.txt",
        "7636",
        "fb70e8b444b20f828226a2dea782d6cf24c2db3900a11e8c01f567f24b0bd42d",
    ],
    [
        "Validation probe",
        "validation_probe500_seed42.txt",
        "500",
        "1e44947a99334b104d826776cf1f04a313ad96a08e43de1c21e35569076c8a4d",
    ],
    [
        "Test",
        "test.txt",
        "7592",
        "febffa69e70ed3df966a3d2092b7c0956142babe7de53dc30dc46a8668de570b",
    ],
]


PPO_CONFIG = [
    ["并行环境", "3072（正式训练与评估命令）"],
    ["rollout", "24 steps/env/iteration，即 73,728 transitions/iteration"],
    ["控制频率", "sim dt=0.005 s，decimation=4，控制周期 0.02 s（50 Hz）"],
    ["训练 episode", "10 s"],
    ["评估 horizon", "60 s"],
    ["Actor / Critic", "[512, 256, 128]，ELU；actor obs 160，critic obs 286，action 29"],
    ["PPO", "clip 0.2，5 epochs，4 mini-batches，lr 1e-3 adaptive，gamma 0.99，lambda 0.95"],
    ["正则", "entropy coef 0.005，max grad norm 1.0，desired KL 0.01"],
    ["模型选择", "Validation Macro-first；epsilon=0.002；随后 Micro、Completion、Body Error、较早 checkpoint"],
    ["评估模式", "seed 42，deterministic，randomization off，Fabric off；每条 motion 从 frame 0 评估一次"],
]


REWARDS = [
    ["Global anchor position", "+0.5", "exp error, std=0.3"],
    ["Global anchor orientation", "+0.5", "exp error, std=0.4"],
    ["Relative body position", "+1.0", "exp error, std=0.3"],
    ["Relative body orientation", "+1.0", "exp error, std=0.4"],
    ["Body linear velocity", "+1.0", "exp error, std=1.0"],
    ["Body angular velocity", "+1.0", "exp error, std=3.14"],
    ["Action rate L2", "-0.1", "平滑动作"],
    ["Joint limit", "-10.0", "关节越限惩罚"],
    ["Undesired contacts", "-0.1", "非脚踝/手腕接触惩罚"],
]


METRICS = [
    ["Micro Success ↑", "成功 motion 数 / 7636；容易受样本量大的类别主导。"],
    ["Macro Success ↑", "先计算 17 个类别各自 success rate，再做无权平均；每个类别权重相同，是 checkpoint 首选指标。"],
    ["Completion ↑", "completed_frames / num_frames 的逐 motion 平均，范围 [0,1]。"],
    ["Body Err ↓", "yaw/root 对齐后，配置追踪 body 的平均欧氏位置误差，单位 m。"],
    ["Joint L2 ↓", "每一步 29 个关节角误差向量的 L2 范数，再按 motion/数据集汇总，单位 rad。"],
    ["Joint RMS ↓", "Joint L2 / sqrt(29)，便于理解单关节典型误差，单位 rad。"],
    ["Failures ↓", "未自然到达最后一帧的 motion 数，等于 7636 - successes。"],
]


SOURCE_FILES = [
    "README.md",
    "PROJECT_HANDOFF_PHUMA_WBT.md",
    "docs/stage0_segment_infrastructure.md",
    "docs/module1_quality_gate.md",
    "docs/module2_difficulty_estimation.md",
    "docs/module3_learning_gap_sampling.md",
    "docs/module4_diversity_sampling.md",
    "docs/joint_specific_gap_design.md",
    "docs/final_test_freeze.md",
    "source/whole_body_tracking/whole_body_tracking/tasks/tracking/tracking_env_cfg.py",
    "source/whole_body_tracking/whole_body_tracking/tasks/tracking/config/g1/agents/rsl_rl_ppo_cfg.py",
    "scripts/rsl_rl/evaluation_utils.py",
    "evaluations/formal_v1/*/validation_full",
    "evaluations/final_minimal/*/validation_full",
    "evaluations/postfreeze_M7Raw_extend50k_from33999_seed42_v1",
    "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1",
    "evaluations/postfreeze_Donly_extend50k_from33999_seed42_v1",
    "evaluations/postfreeze_Donly_extend59k_from49999_seed42_v1",
    "evaluations/formal_v2/A_raw_motion_gap_segment_seed42/validation_full",
    "evaluations/formal_v2/M7_lambda0p10_seed42/validation_full",
    "outputs/joint_gap_stage6_probe500/probe_summary.json",
]


def fmt(value: float) -> str:
    return f"{value:.6f}"


def percent_point(delta: float) -> str:
    return f"{delta * 100:+.3f} pp"


def markdown_table(headers: Sequence[str], rows: Iterable[Sequence[object]]) -> str:
    lines = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
    for row in rows:
        lines.append("| " + " | ".join(str(value) for value in row) + " |")
    return "\n".join(lines)


def write_csv(path: Path, headers: Sequence[str], rows: Iterable[Sequence[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        writer.writerows(rows)


def read_category_rows(path: Path) -> dict[str, dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as handle:
        return {row["category"]: row for row in csv.DictReader(handle)}


def assert_close(label: str, observed: float, expected: float, tolerance: float = 7.5e-7) -> None:
    if abs(observed - expected) > tolerance:
        raise RuntimeError(f"{label}: observed={observed}, expected={expected}")


def validate_summary(path: Path, expected: dict[str, object]) -> None:
    payload = json.loads(path.read_text(encoding="utf-8"))
    validate_metric_payload(str(path), payload, expected)


def validate_metric_payload(label: str, payload: dict[str, object], expected: dict[str, object]) -> None:
    mapping = {
        "micro": "micro_success_rate",
        "macro": "macro_success_rate",
        "completion": "mean_completion_ratio",
        "body": "mean_body_position_error_m",
        "joint": "mean_joint_position_error_l2_rad",
        "rms": "mean_joint_position_error_rms_rad",
    }
    for local_key, json_key in mapping.items():
        assert_close(f"{label}:{json_key}", float(payload[json_key]), float(expected[local_key]))
    if int(payload["num_failure"]) != int(expected["failures"]):
        raise RuntimeError(f"{label}: failure count does not match report data")


def summarize_per_motion_csv(path: Path) -> dict[str, object]:
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    if not rows:
        raise RuntimeError(f"No per-motion rows found in {path}")

    success_by_category: dict[str, list[int]] = {}
    for row in rows:
        success_by_category.setdefault(row["category"], []).append(int(row["success"]))
    num_success = sum(int(row["success"]) for row in rows)
    num_motions = len(rows)
    return {
        "num_failure": num_motions - num_success,
        "micro_success_rate": num_success / num_motions,
        "macro_success_rate": sum(sum(values) / len(values) for values in success_by_category.values())
        / len(success_by_category),
        "mean_completion_ratio": sum(float(row["completion_ratio"]) for row in rows) / num_motions,
        "mean_body_position_error_m": sum(float(row["body_position_error_m"]) for row in rows) / num_motions,
        "mean_joint_position_error_l2_rad": sum(float(row["joint_position_error_l2_rad"]) for row in rows)
        / num_motions,
        "mean_joint_position_error_rms_rad": sum(float(row["joint_position_error_rms_rad"]) for row in rows)
        / num_motions,
    }


def validate_source_results() -> None:
    checks = [
        (
            ROOT / "evaluations/postfreeze_Donly_extend50k_from33999_seed42_v1/validation_full/model_49999/summary.json",
            next(row for row in BUDGET_RESULTS if row["method"] == "D-only" and row["iteration"] == 49999),
        ),
        (
            ROOT / "evaluations/postfreeze_Donly_extend59k_from49999_seed42_v1/validation_full/model_54000/summary.json",
            next(row for row in BUDGET_RESULTS if row["method"] == "D-only" and row["iteration"] == 54000),
        ),
        (
            ROOT / "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_full/model_59000/summary.json",
            next(row for row in BUDGET_RESULTS if row["method"] == "M7-Raw" and row["iteration"] == 59000),
        ),
        (
            ROOT / "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_full_sweep/model_60000/summary.json",
            next(row for row in BUDGET_RESULTS if row["method"] == "M7-Raw" and row["iteration"] == 60000),
        ),
        (
            ROOT / "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_full_sweep/model_69999/summary.json",
            next(row for row in BUDGET_RESULTS if row["method"] == "M7-Raw" and row["iteration"] == 69999),
        ),
    ]
    for path, expected in checks:
        validate_summary(path, expected)

    for expected in SUPPLEMENTARY_FULL_RESULTS:
        source = ROOT / str(expected["source"])
        validate_metric_payload(str(source), summarize_per_motion_csv(source), expected)

    joint_gap_path = ROOT / "outputs/joint_gap_stage6_probe500/probe_summary.json"
    joint_gap_payload = json.loads(joint_gap_path.read_text(encoding="utf-8"))
    validate_metric_payload(str(joint_gap_path) + ":m7_raw", joint_gap_payload["m7_raw"], JOINT_GAP_PROBE_RESULTS[0])
    validate_metric_payload(str(joint_gap_path) + ":m7_jgap", joint_gap_payload["m7_jgap"], JOINT_GAP_PROBE_RESULTS[1])
    if joint_gap_payload.get("status") != "NO-GO: DOMINATED" or joint_gap_payload.get("enter_full_34k") is not False:
        raise RuntimeError(f"{joint_gap_path}: unexpected M7-JGap gate decision")


def category_comparison() -> list[dict[str, object]]:
    d_rows = read_category_rows(
        ROOT
        / "evaluations/postfreeze_Donly_extend59k_from49999_seed42_v1/validation_full/model_54000/category_summary.csv"
    )
    m_rows = read_category_rows(
        ROOT
        / "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_full/model_59000/category_summary.csv"
    )
    rows: list[dict[str, object]] = []
    for category in sorted(d_rows):
        d = d_rows[category]
        m = m_rows[category]
        d_success = float(d["success_rate"])
        m_success = float(m["success_rate"])
        rows.append(
            {
                "category": category,
                "motions": int(d["num_motions"]),
                "d_success": d_success,
                "m_success": m_success,
                "delta": m_success - d_success,
                "d_completion": float(d["mean_completion_ratio"]),
                "m_completion": float(m["mean_completion_ratio"]),
                "d_joint": float(d["mean_joint_position_error_l2_rad"]),
                "m_joint": float(m["mean_joint_position_error_l2_rad"]),
                "d_failures": int(d["num_failure"]),
                "m_failures": int(m["num_failure"]),
            }
        )
    return rows


def dense_probe_rows() -> list[dict[str, object]]:
    base = ROOT / "evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1"
    csv_paths = [
        base / "validation_probe500_dense_sweep/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_50500_55000/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_55500_60000/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_60500_65000/checkpoint_comparison.csv",
        base / "validation_probe500_dense_sweep_65500_69999/checkpoint_comparison.csv",
    ]
    by_iteration: dict[int, dict[str, object]] = {}
    for path in csv_paths:
        if not path.exists():
            continue
        with path.open(newline="", encoding="utf-8") as handle:
            for row in csv.DictReader(handle):
                iteration = int(row["iteration"])
                by_iteration[iteration] = {
                    "iteration": iteration,
                    "checkpoint": row["checkpoint"],
                    "micro": float(row["micro_success_rate"]),
                    "macro": float(row["macro_success_rate"]),
                    "completion": float(row["mean_completion_ratio"]),
                    "body": float(row["mean_body_position_error_m"]),
                    "joint": float(row["mean_joint_position_error_l2_rad"]),
                    "failures": int(row["num_failures"]),
                }
    single_path = base / "validation_probe500_dense_single/model_50000/summary.json"
    if single_path.exists():
        payload = json.loads(single_path.read_text(encoding="utf-8"))
        by_iteration[50000] = {
            "iteration": 50000,
            "checkpoint": str(payload["checkpoint"]),
            "micro": float(payload["micro_success_rate"]),
            "macro": float(payload["macro_success_rate"]),
            "completion": float(payload["mean_completion_ratio"]),
            "body": float(payload["mean_body_position_error_m"]),
            "joint": float(payload["mean_joint_position_error_l2_rad"]),
            "failures": int(payload["num_failure"]),
        }
    return [by_iteration[key] for key in sorted(by_iteration)]


def configure_plots() -> font_manager.FontProperties:
    FIGURES.mkdir(parents=True, exist_ok=True)
    Path(os.environ["MPLCONFIGDIR"]).mkdir(parents=True, exist_ok=True)
    font_prop = font_manager.FontProperties(fname=CN_FONT_PATH)
    plt.rcParams.update(
        {
            "font.family": font_prop.get_name(),
            "font.size": 10,
            "axes.unicode_minus": False,
            "axes.edgecolor": "#4B5563",
            "axes.labelcolor": "#111827",
            "xtick.color": "#374151",
            "ytick.color": "#374151",
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "savefig.facecolor": "white",
        }
    )
    return font_prop


def save_figure(fig: plt.Figure, name: str) -> Path:
    path = FIGURES / name
    fig.savefig(path, dpi=220, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_architecture(font_prop: font_manager.FontProperties) -> Path:
    fig, ax = plt.subplots(figsize=(12.4, 5.1))
    ax.set_xlim(0, 12.4)
    ax.set_ylim(0, 5.1)
    ax.axis("off")

    colors = ["#E8F1F8", "#E8F5EC", "#FFF2D8", "#FCE8E6", "#EDE9F7", "#E7F3F4"]
    boxes = [
        (0.2, 3.15, 1.65, 1.05, "PHUMA 数据\n76,086 motions"),
        (2.05, 3.15, 1.65, 1.05, "G1/WBT 转换\n29 DOF · 30 bodies"),
        (3.9, 3.15, 1.65, 1.05, "Train/Val/Test\nmanifest + SHA256"),
        (5.75, 3.15, 1.65, 1.05, "Stage 0\n1 s Segment 索引"),
        (9.75, 3.15, 2.15, 1.05, "Isaac Lab + PPO\n共享 tracking policy"),
        (9.75, 1.00, 2.15, 1.05, "确定性评估\nProbe → Full → Test"),
    ]
    for idx, (x, y, w, h, label) in enumerate(boxes):
        patch = FancyBboxPatch(
            (x, y),
            w,
            h,
            boxstyle="round,pad=0.03,rounding_size=0.05",
            facecolor=colors[idx],
            edgecolor="#334155",
            linewidth=1.2,
        )
        ax.add_patch(patch)
        ax.text(x + w / 2, y + h / 2, label, ha="center", va="center", fontproperties=font_prop, fontsize=10)

    sampler = FancyBboxPatch(
        (5.75, 0.55),
        3.55,
        2.05,
        boxstyle="round,pad=0.04,rounding_size=0.05",
        facecolor="#F8FAFC",
        edgecolor="#0F766E",
        linewidth=1.6,
    )
    ax.add_patch(sampler)
    ax.text(7.525, 2.28, "模块化采样器", ha="center", va="center", fontproperties=font_prop, fontsize=11, weight="bold")
    module_labels = [
        (6.05, 1.55, "Quality Gate"),
        (7.55, 1.55, "Difficulty"),
        (6.05, 0.85, "Online Error / Gap"),
        (7.55, 0.85, "Diversity K=8"),
    ]
    for x, y, label in module_labels:
        p = FancyBboxPatch(
            (x, y),
            1.35,
            0.48,
            boxstyle="round,pad=0.02,rounding_size=0.03",
            facecolor="#FFFFFF",
            edgecolor="#64748B",
            linewidth=0.9,
        )
        ax.add_patch(p)
        ax.text(x + 0.675, y + 0.24, label, ha="center", va="center", fontproperties=font_prop, fontsize=8.6)

    def arrow(x1: float, y1: float, x2: float, y2: float) -> None:
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops={"arrowstyle": "->", "lw": 1.4, "color": "#475569"})

    arrow(1.85, 3.68, 2.05, 3.68)
    arrow(3.70, 3.68, 3.90, 3.68)
    arrow(5.55, 3.68, 5.75, 3.68)
    arrow(7.40, 3.68, 9.75, 3.68)
    arrow(7.525, 2.60, 7.525, 3.15)
    arrow(9.30, 1.58, 9.75, 1.58)
    arrow(10.83, 3.15, 10.83, 2.05)
    ax.text(
        7.53,
        0.22,
        "D-only: Diversity + Raw/Raw    |    M7-Raw: Quality + Diversity + Raw/Raw",
        ha="center",
        va="center",
        fontproperties=font_prop,
        fontsize=9.2,
        color="#0F172A",
    )
    ax.set_title("PHUMA-WBT 系统与采样模块总览", fontproperties=font_prop, fontsize=15, weight="bold", pad=8)
    return save_figure(fig, "figure_1_system_architecture.png")


def plot_formal_results(font_prop: font_manager.FontProperties) -> Path:
    rows = sorted(FORMAL_RESULTS, key=lambda item: item["macro"])
    labels = [str(row["method"]) for row in rows]
    y = list(range(len(rows)))
    fig, ax = plt.subplots(figsize=(10.8, 6.7))
    height = 0.36
    ax.barh([value - height / 2 for value in y], [float(row["micro"]) for row in rows], height, label="Micro", color="#3B82F6")
    ax.barh([value + height / 2 for value in y], [float(row["macro"]) for row in rows], height, label="Macro", color="#E45756")
    ax.set_yticks(y, labels, fontproperties=font_prop)
    ax.set_xlim(0.84, 0.91)
    ax.set_xlabel("Success rate", fontproperties=font_prop)
    ax.set_title("34k 正式 Full Validation：Micro 与 Macro Success", fontproperties=font_prop, fontsize=14, weight="bold")
    ax.grid(axis="x", alpha=0.22)
    ax.legend(frameon=False)
    for idx, row in enumerate(rows):
        ax.text(float(row["micro"]) + 0.0005, idx - height / 2, f"{float(row['micro']):.3f}", va="center", fontsize=7.6)
        ax.text(float(row["macro"]) + 0.0005, idx + height / 2, f"{float(row['macro']):.3f}", va="center", fontsize=7.6)
    fig.tight_layout()
    return save_figure(fig, "figure_2_formal_34k_success.png")


def plot_budget_results(font_prop: font_manager.FontProperties) -> Path:
    fig, axes = plt.subplots(2, 2, figsize=(11.4, 7.5), sharex=True)
    panels = [
        ("macro", "Macro Success ↑"),
        ("micro", "Micro Success ↑"),
        ("completion", "Completion ↑"),
        ("joint", "Joint L2 ↓"),
    ]
    styles = {
        "D-only": {"color": "#2563EB", "marker": "s"},
        "M7-Raw": {"color": "#DC2626", "marker": "o"},
    }
    for ax, (key, title) in zip(axes.flat, panels):
        for method in ["D-only", "M7-Raw"]:
            rows = [row for row in BUDGET_RESULTS if row["method"] == method]
            ax.plot(
                [int(row["iteration"]) / 1000 for row in rows],
                [float(row[key]) for row in rows],
                label=method,
                linewidth=2.0,
                markersize=6,
                **styles[method],
            )
        ax.set_title(title, fontproperties=font_prop, fontsize=11, weight="bold")
        ax.grid(alpha=0.22)
        ax.set_xlabel("Checkpoint iteration (k)", fontproperties=font_prop)
    axes[0, 0].legend(frameon=False)
    axes[0, 0].axvline(59, color="#111827", alpha=0.22, linestyle="--", linewidth=1)
    axes[0, 1].axvline(59, color="#111827", alpha=0.22, linestyle="--", linewidth=1)
    axes[1, 0].axvline(59, color="#111827", alpha=0.22, linestyle="--", linewidth=1)
    axes[1, 1].axvline(59, color="#111827", alpha=0.22, linestyle="--", linewidth=1)
    fig.suptitle("Full Validation 预算扩展：M7-Raw 在 59k 达到当前峰值", fontproperties=font_prop, fontsize=15, weight="bold")
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    return save_figure(fig, "figure_3_budget_scaling.png")


def plot_dense_probe(rows: Sequence[dict[str, object]], font_prop: font_manager.FontProperties) -> Path:
    fig, ax = plt.subplots(figsize=(11.0, 5.3))
    x = [int(row["iteration"]) / 1000 for row in rows]
    ax.plot(x, [float(row["micro"]) for row in rows], color="#2563EB", marker="o", markersize=3.5, linewidth=1.5, label="Micro")
    ax.plot(x, [float(row["macro"]) for row in rows], color="#DC2626", marker="o", markersize=3.5, linewidth=1.5, label="Macro")
    ax.plot(x, [float(row["completion"]) for row in rows], color="#059669", marker="o", markersize=3.5, linewidth=1.5, label="Completion")
    ax.axvline(59, color="#111827", linestyle="--", linewidth=1.2, alpha=0.7)
    ax.axvspan(62, 67, color="#FCA5A5", alpha=0.20)
    ax.text(59.15, 0.10, "59k Probe peak", fontproperties=font_prop, fontsize=9)
    ax.text(62.3, 0.23, "训练不稳定区间", fontproperties=font_prop, fontsize=9, color="#991B1B")
    ax.set_ylim(-0.02, 1.02)
    ax.set_xlabel("Checkpoint iteration (k)", fontproperties=font_prop)
    ax.set_ylabel("Probe500 metric", fontproperties=font_prop)
    ax.set_title("M7-Raw 50k–70k Dense Validation Probe", fontproperties=font_prop, fontsize=14, weight="bold")
    ax.grid(alpha=0.22)
    ax.legend(frameon=False, ncol=3)
    fig.tight_layout()
    return save_figure(fig, "figure_4_m7raw_dense_probe.png")


def plot_category_results(rows: Sequence[dict[str, object]], font_prop: font_manager.FontProperties) -> Path:
    ordered = sorted(rows, key=lambda row: float(row["m_success"]) - float(row["d_success"]))
    labels = [str(row["category"]) for row in ordered]
    y = list(range(len(ordered)))
    fig, ax = plt.subplots(figsize=(10.8, 7.3))
    height = 0.36
    ax.barh([value - height / 2 for value in y], [float(row["d_success"]) for row in ordered], height, color="#2563EB", label="D-only 54k")
    ax.barh([value + height / 2 for value in y], [float(row["m_success"]) for row in ordered], height, color="#DC2626", label="M7-Raw 59k")
    ax.set_yticks(y, labels, fontproperties=font_prop)
    ax.set_xlim(0.68, 1.02)
    ax.set_xlabel("Category success rate", fontproperties=font_prop)
    ax.set_title("17 类动作成功率：M7-Raw 59k vs D-only 54k", fontproperties=font_prop, fontsize=14, weight="bold")
    ax.grid(axis="x", alpha=0.22)
    ax.legend(frameon=False, loc="lower right")
    fig.tight_layout()
    return save_figure(fig, "figure_5_category_success.png")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 55, start: int = 65, bottom: int = 55, end: int = 65) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = CN_FONT_NAME
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT_NAME)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = CN_FONT_NAME
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT_NAME)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    for name, size, color in [
        ("Title", 24, "17365D"),
        ("Heading 1", 16, "17365D"),
        ("Heading 2", 13, "1F4E78"),
        ("Heading 3", 11.5, "2F5597"),
    ]:
        style = doc.styles[name]
        style.font.name = CN_FONT_NAME
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT_NAME)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def configure_headers(doc: Document) -> None:
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.text = f"{PROJECT} · 论文材料汇总 · {DATE}"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header.runs:
            set_run_font(run, 8.5, color="64748B")
        footer = section.footer.paragraphs[0]
        add_page_number(footer)
        for run in footer.runs:
            set_run_font(run, 8.5, color="64748B")


def add_para(doc: Document, text: str, *, bold_prefix: str | None = None, italic: bool = False) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0.74)
    if bold_prefix and text.startswith(bold_prefix):
        first = paragraph.add_run(bold_prefix)
        set_run_font(first, bold=True)
        second = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(second)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
        run.italic = italic


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(2)
    for run in paragraph.runs:
        set_run_font(run)
    if not paragraph.runs:
        run = paragraph.add_run(text)
        set_run_font(run)
    else:
        paragraph.runs[0].text = text


def add_note(doc: Document, text: str, fill: str = "FFF2CC", color: str = "7F6000") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 90, 120, 90, 120)
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    set_run_font(run, 9.5, bold=True, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[object]],
    *,
    font_size: float = 8.5,
    highlight: callable | None = None,
) -> None:
    values = [list(row) for row in rows]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for col, header in enumerate(headers):
        cell = table.rows[0].cells[col]
        cell.text = str(header)
        set_cell_shading(cell, "1F4E78")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, font_size, bold=True, color="FFFFFF")
    for row_index, row in enumerate(values):
        cells = table.add_row().cells
        row_fill = "F4F7FA" if row_index % 2 else "FFFFFF"
        if highlight is not None:
            custom = highlight(row)
            if custom:
                row_fill = custom
        for col, value in enumerate(row):
            cell = cells[col]
            cell.text = str(value)
            set_cell_shading(cell, row_fill)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col in (0, 1) else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_run_font(run, font_size, bold=row_fill in {"E2F0D9", "DDEBF7"})
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure(doc: Document, path: Path, caption: str, width: float = 6.55) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    set_run_font(run, 9, color="475569")


def add_code_block(doc: Document, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F1F5F9")
    set_cell_margins(cell, 80, 120, 80, 120)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "DejaVu Sans Mono"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT_NAME)
    run.font.size = Pt(8.5)


def build_markdown(category_rows: Sequence[dict[str, object]], probe_rows: Sequence[dict[str, object]]) -> str:
    formal_table = [
        [
            row["method"],
            "34k",
            f"`{row['checkpoint']}`",
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in FORMAL_RESULTS
    ]
    budget_table = [
        [
            row["method"],
            row["budget"],
            f"`{row['checkpoint']}`",
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
            row["phase"],
        ]
        for row in BUDGET_RESULTS
    ]
    category_table = [
        [
            row["category"],
            row["motions"],
            fmt(float(row["d_success"])),
            fmt(float(row["m_success"])),
            percent_point(float(row["delta"])),
            row["d_failures"],
            row["m_failures"],
        ]
        for row in category_rows
    ]
    probe_table = [
        [
            row["iteration"],
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in probe_rows
    ]
    supplementary_table = [
        [
            row["method"],
            row["design"],
            row["budget"],
            f"`{row['checkpoint']}`",
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in SUPPLEMENTARY_FULL_RESULTS
    ]
    joint_gap_table = [
        [
            row["method"],
            row["budget"],
            f"`{row['checkpoint']}`",
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in JOINT_GAP_PROBE_RESULTS
    ]

    sections: list[str] = []
    sections.append(
        f"""# PHUMA × Whole-Body Tracking 论文材料汇总

版本：{DATE}  
项目：`/home/l/whole_body_tracking_new`

> 本材料以仓库文档、代码配置和本地评估结果为事实来源。旧 D-only 冻结结论与后续预算扩展实验分开报告；当前没有把 Test 结果用于模型选择。

## 1. 一页结论

- **34k 正式冻结结论**：D-only `model_33500.pt` 的 Macro Success 为 `0.899319`，是当前完整 34k 正式表中的最佳 Macro；因此旧冻结结论成立。
- **预算扩展结论**：M7-Raw 随预算增加显著改善，在 `model_59000.pt` 达到当前最佳 Full Validation：Micro `0.916841`、Macro `0.913028`、Completion `0.952066`、Body Err `0.047723 m`、Joint L2 `0.762033 rad`、Failures `635`。
- **公平对照**：D-only 续训后的 Validation 选择点是 `model_54000.pt`。M7-Raw 59k 相比 D-only 54k，Micro `+2.475 pp`、Macro `+1.462 pp`、Completion `+1.410 pp`，Failures 减少 `189`。
- **训练上限**：M7-Raw 60k 的 Macro 已从 59k 的 `0.913028` 降到 `0.899619`，70k 全面退化；当前证据支持峰值在 59k 附近，而不是“训练越久越好”。
- **论文表述**：D-only 是短预算强方法；Quality Gate + Diversity + Hierarchical Raw Error（M7-Raw）具有更高的长预算上限。该结论基于单 seed Validation，应写成“结果支持/表明”，不要写成普遍定律。

## 2. 论文题目与定位

推荐题目（采用 59k M7-Raw 为最终方法时）：

**面向大规模异质动作的人形机器人全身跟踪：质量门控的多样性约束层级误差采样**  
**Quality-Gated Diversity-Constrained Hierarchical Error Sampling for Large-Scale Humanoid Whole-Body Tracking**

保守题目（保持旧 D-only 冻结方法为主角时）：

**面向大规模人形机器人动作跟踪的多样性预算层级误差采样**  
**Diversity-Budgeted Hierarchical Error Sampling for Large-Scale Humanoid Motion Tracking**

核心研究问题：面对由多个来源构成、难度和质量高度不均衡的大规模动作库，如何分配有限 PPO 训练预算，使策略既能持续覆盖不同运动模式，又能重点学习当前误差较大的动作和片段？

## 3. 可直接改写的摘要初稿

大规模人形机器人动作库包含显著的类别不均衡、轨迹质量差异和动作难度差异，采用均匀采样时，有限的强化学习预算难以同时保证运动多样性覆盖与困难片段学习。本文基于 PHUMA 动作数据构建 Unitree G1 多动作全身跟踪系统，并提出模块化的层级采样框架。该框架首先将动作划分为固定时长片段，在训练集上离线建立轨迹质量标签、固有难度分箱和运动聚类，再根据策略在线跟踪误差在 Cluster、Motion 与 Segment 三个层级分配训练预算。质量模块只限制异常片段成为 episode 起点，固有难度模块只提供策略无关的校准基准，多样性模块只控制运动簇预算，从而保持各因素解耦。本文在 6000 条训练动作和 7636 条独立验证动作上进行消融。34k 固定预算下，多样性预算层级原始误差方法 D-only 取得 0.899319 的 Macro Success；进一步的预算扩展发现，加入质量门控的 M7-Raw 在 59k 达到 0.916841 Micro Success、0.913028 Macro Success 和 0.952066 Completion，失败动作数降至 635。结果表明，多样性约束可提升有限预算下的覆盖，而质量门控的收益具有明显的预算依赖性；同时 60k–70k 的退化说明 PPO checkpoint 性能并非随训练轮数单调增长。

## 4. 建议贡献点

1. 构建 PHUMA 到 WBT/G1 的大规模多动作训练管线，支持 76,086 条转换轨迹、manifest 管理和每环境独立 motion/start-frame assignment。
2. 建立质量、固有难度、在线学习状态与运动多样性相互解耦的模块化采样框架，并提供严格 metadata identity、checkpoint/resume 和概率合法性检查。
3. 提出 Cluster→Motion→Segment→Start Frame 的多样性预算层级采样，并通过 D-only、M7-Raw、GlobalRaw、M0–M7 等消融分离各模块作用。
4. 发现采样器收益具有训练预算依赖性：D-only 在 34k 短预算更强，而 M7-Raw 在 50k 后反超并在 59k 达到更高上限；进一步训练导致明显退化。

## 5. 系统结构

![系统结构](paper_assets_2026-09-08/figures/figure_1_system_architecture.png)

### 5.1 模块总表

{markdown_table(["模块", "设计", "功能"], MODULE_TABLE)}

### 5.2 数据转换与多动作加载

PHUMA G1 原始 `.npy` 经 `scripts/phuma_to_npz.py` 转换为 WBT `.npz`。核心字段包括 `joint_pos`、`joint_vel`、`body_pos_w`、`body_quat_w`、`body_lin_vel_w`、`body_ang_vel_w`、`joint_names` 和 `body_names`。系统将 PHUMA 的关节表示补齐为 WBT 使用的 29 个关节自由度，并在 Isaac 中生成 30 个 body 的位姿与速度。`MotionLoader` 支持单一 `.npz`、目录递归和 `.txt` manifest；`MotionCommand` 为每个并行环境独立采样 motion 和起始帧。

### 5.3 Stage 0：Segment 基础设施

对 motion `m`，默认片段长度为 1 s：

```text
segment_frames[m] = max(1, round(fps[m] × 1.0 s))
num_segments[m]   = ceil(num_frames[m] / segment_frames[m])
global_id         = motion_segment_offsets[m] + local_segment_id
```

random6000 共形成 `21,575` 个 segment。Stage 0 维护局部/全局 ID、合法起点、共享计数、coverage、概率 fallback 和 checkpoint 恢复；关闭研究模块时保持 legacy uniform 采样路径与 RNG 调用结构。

### 5.4 模块 1：Quality Gate

质量模块只判断最终 G1 参考轨迹是否存在数值、转换、几何或明确物理问题，不读取 policy error、reward、success 或 PPO loss。离线审计包含 17 类 metric，包括非有限值、四元数范数、URDF 关节位置/速度限位、速度一致性、加速度/jerk 孤立峰、位姿连续性、穿地和足滑。

质量分：

```text
Q(m,s) = clip(1 - Σ_i w_i severity_i / Σ_i w_i, 0, 1)
```

当前 random6000 统计为 `15,155 pass + 6,330 borderline + 90 reject = 21,575 segments`。开启 gate 且允许 borderline 后，可作为起点的 frame 比例为 `0.995768`；有 2 条 motion 没有任何合法质量起点，因此 runtime effective motions 为 `5998`，但原始 6000 行 manifest 不被删除或替换。Gate 只禁止 reject segment 成为 assignment 起点，随后 rollout 仍可能自然经过 reject 区间。

### 5.5 模块 2：策略无关固有难度

模块二从最终 G1 轨迹提取 28 项运动学与接触特征，包括 root 线/角速度及加速度、关节速度与活动范围、body 高度变化、手脚速度、单/双支撑、腾空比例和接触切换。所有导数先在完整 motion 上计算，再按 segment 聚合，避免切段产生伪零值。

Train-only robust 标准化使用 median 与 `1.4826 × MAD`，MAD 近零时使用 P05–P95 fallback；加权 raw difficulty 再通过经验 CDF 映射到 `[0,1]`，形成 10 个近似等频 bin。Difficulty 不直接删除或重采样样本，只用于 M5/M6/M7 的同难度 error calibration。

### 5.6 模块 3：在线误差与学习缺口

每个 control step 计算 body、joint 和 orientation tracking error，并把它归因到当前播放的 motion/segment。EMA 系数 `ρ=0.95`，默认 warmup 1000 iterations、每 50 iterations 更新采样概率。Segment raw error 综合归一化 tracking error、termination、completion 和 success；Motion raw error再聚合 segment mean/P90 与 episode outcome。

```text
G_global(m,s) = clip((E_segment(m,s) - μ_difficulty_bin) / σ_difficulty_bin, -5, 5)
G_motion(m)   = aggregate(max(G_global, 0), motion outcomes)
G_local(m,s)  = G_global(m,s) - median_s G_global(m,s)
```

所有 adaptive sampler 保留 `0.15` uniform mix、欠采样 bonus 和概率 cap。Raw-error 方法不读取 difficulty；learning-gap 方法才使用难度 bin。

### 5.7 模块 4：运动多样性约束

模块四从模块二的 17 项 segment source feature 聚合出 30 项 motion feature，经 Train-only robust scaling 和 NumPy KMeans++ 得到 `K=8` 个冻结运动簇，cluster sizes 为 `[1902, 456, 503, 948, 603, 557, 728, 303]`。来源目录和人工 category 不参与聚类。

```text
minimum_share = 0.5 / C
size_component(c) = n_c^0.5 / Σ_j n_j^0.5
P(c) = minimum_share + (1 - C×minimum_share) × size_component(c)
P(c,m,s) = P(c) × P(m|c) × P(s|m)
```

`P(c)` 只由 eligible cluster size 决定，不读取 policy error；error/gap 只在 `P(m|c)` 和 `P(s|m)` 内起作用。这一职责分离避免同一运动模式被 cluster 层与 error 层重复放大。

### 5.8 策略网络、奖励与终止

{markdown_table(["配置", "值"], PPO_CONFIG)}

奖励项：

{markdown_table(["项", "权重", "说明"], REWARDS)}

主要终止条件为 timeout、anchor z 误差超过 0.25 m、anchor orientation 误差超过 0.8，以及手腕/脚踝等末端 body z 误差超过 0.25 m。

## 6. 实验设计

### 6.1 数据划分与完整性

{markdown_table(["Split", "Manifest", "Motions", "SHA256"], SPLITS)}

Train 与 Validation/Test 无交集，Validation 与 Test 无交集；Probe500 是 Validation Full 子集。所有正式方法使用同一 random6000、seed 42、PPO、reward、observation、termination 与 evaluator。

### 6.2 方法矩阵

{markdown_table(["Method", "Motion", "Segment", "Quality", "Difficulty", "Diversity", "证据状态"], METHOD_MATRIX)}

其中 D-only 表示在 M4 的 Raw/Raw 层级采样上只加入 Diversity；M7-Raw 表示 D-only 再加入 Quality Gate。它不是标准 M7 的 Learning Gap，而是 `Quality + Diversity + Raw/Raw`。

### 6.3 指标定义

{markdown_table(["指标", "定义"], METRICS)}

注意：失败模型可能在很少的 step 后终止，从而产生表面较低的 Body/Joint 平均误差。因此误差必须与 Success 和 Completion 联合解读；零成功 checkpoint 的低误差不能视为性能提升。

### 6.4 Macro-first checkpoint 规则

评估器按 Macro Success 优先选择。若与当前最优的差值绝对值不超过 `0.002`，依次比较 Micro、Completion、Body Error；各层沿用同一 epsilon，仍无法区分时选更早 checkpoint。Joint L2 不参与选择。Probe500 只筛选候选，论文主结论必须来自 Full Validation。

## 7. 34k 正式 Full Validation

![34k 正式结果](paper_assets_2026-09-08/figures/figure_2_formal_34k_success.png)

{markdown_table(["Method", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body Err ↓", "Joint L2 ↓", "Failures ↓"], formal_table)}

主要观察：D-only 34k 相比 GlobalRaw baseline，Micro 提升 `2.095 pp`、Macro 提升 `1.792 pp`、Completion 提升 `1.151 pp`，失败数减少 `160`。M7-Raw 34k 的 Micro、Body、Joint 和 Failures 略优于 D-only，但 Macro 低 `0.629 pp`，因此 Macro-first 冻结 D-only 是一致且可复现的选择。

M2/M3 已有实现与 pilot，但当前评估归档没有独立的 7636-motion Full Validation，因此不能在论文表格中虚构数值。若论文必须声称“完整 M0–M7 八项正式消融”，M2/M3 是待补项；它们不需要续训到 59k，只需与原 34k 固定预算协议一致。

## 8. 预算扩展实验

![预算扩展](paper_assets_2026-09-08/figures/figure_3_budget_scaling.png)

{markdown_table(["Method", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body Err ↓", "Joint L2 ↓", "Failures ↓", "Phase"], budget_table)}

关键差值：

- M7-Raw 59k vs M7-Raw 34k：Micro `+3.195 pp`，Macro `+2.000 pp`，Completion `+1.801 pp`，Body Err 降低约 `4.40%`，Joint L2 降低约 `5.76%`，Failures 减少 `244`。
- M7-Raw 59k vs D-only 54k：Micro `+2.475 pp`，Macro `+1.462 pp`，Completion `+1.410 pp`，Body Err 低约 `5.02%`，Joint L2 低约 `7.43%`，Failures 少 `189`。
- D-only 54k vs D-only 34k：Micro `+0.760 pp`、Completion `+0.389 pp`、Failures 少 `58`，但 Macro `-0.091 pp`，Joint L2 反而高约 `1.00%`。D-only 的长预算收益有限。
- M7-Raw 60k vs 59k：Macro `-1.341 pp`、Joint L2 变差 `0.031116 rad`；70k 相比 59k 的 Micro `-2.135 pp`、Macro `-1.442 pp`、Failures 增加 `163`。当前上限应报告为 59k 附近。

## 9. Dense Probe 与训练稳定性

![Dense Probe](paper_assets_2026-09-08/figures/figure_4_m7raw_dense_probe.png)

{markdown_table(["Iteration", "Micro", "Macro", "Completion", "Body", "Joint L2", "Failures/500"], probe_table)}

Dense Probe 在 62k–67k 出现接近零成功的崩溃区间，之后有所恢复。这说明 PPO checkpoint 性能高度非单调，保存密集 checkpoint 和 Validation 选择是必要的。该区间的 Body/Joint 误差由于 episode 极早终止而不可单独解释。

## 10. 类别级分析

![类别成功率](paper_assets_2026-09-08/figures/figure_5_category_success.png)

{markdown_table(["Category", "N", "D-only 54k", "M7-Raw 59k", "Δ", "D Fail", "M7-Raw Fail"], category_table)}

M7-Raw 在 17 类中成功率提高 10 类、持平 4 类、降低 3 类。主要提升来自 `fitness (+7.241 pp, 少 142 个失败)`、`dance (+7.143 pp)`、`LAFAN1 (+3.187 pp)`、`kungfu (+2.985 pp)` 和 `aist (+2.762 pp)`。回退出现在 `humman (-1.961 pp)`、`perform (-1.183 pp)` 和 `idea400 (-0.101 pp)`；其中部分类别样本量很小，应避免过度解释。

## 11. 补充消融与负结果

### 11.1 formal_v2 诊断组合（Full Validation）

{markdown_table(["Method", "Design", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body ↓", "Joint ↓", "Failures ↓"], supplementary_table)}

`V2-Diag-A` 使用 Quality、Difficulty 与 Diversity，motion 层采用 raw error，segment 层采用 relative learning gap；`M7-v2 (lambda=0.10)` 则对 raw segment error 施加 gap correction。两者都有完整 7636-motion Validation，但 Macro 分别为 `0.886414` 和 `0.883476`，均低于 D-only 34k 的 `0.899319`，支持“更复杂的 gap 校准在当前预算下没有稳定优于 Raw/Raw”的结论。

### 11.2 M7-JGap 10k 固定 Probe（不可与 Full Validation 主表横比）

{markdown_table(["Method", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body ↓", "Joint ↓", "Failures ↓"], joint_gap_table)}

M7-JGap 在 M7-Raw 的 segment priority 上加入按关节、按难度分箱校准的 top-k joint-gap correction，离线冻结 `lambda_joint=0.025`。与同阶段 M7-Raw 相比，它的 Macro 增加 `0.441 pp`，但 Micro 降低 `1.600 pp`、Completion 降低 `1.318 pp`、Joint L2 变差约 `1.064%`，且多 8 个失败；冻结 gate 因此给出 `NO-GO: DOMINATED`，未进入 34k。另有约 `99.66%` 的关节尺度由 sigma floor 控制，这是该方向的校准限制。

## 12. 论文讨论主线

### 12.1 为什么 D-only 在短预算更强

Raw error 直接响应当前策略跟踪失败，多样性 cluster budget 又限制相似难动作集中占用预算，二者在较短训练期内形成清晰、低延迟的反馈。Learning-gap 路径需要可靠 EMA、motion outcome 和 difficulty-bin calibration，有限预算下可能付出更长 cold-start 与估计噪声成本。

### 12.2 为什么 M7-Raw 后期反超

Quality Gate 只移除约 `0.423%` 的 assignment-start frames，短期影响很小；但在数十亿 transitions 的重复抽样下，少量异常起点会持续消耗 rollout 与优化预算。M7-Raw 保留 diversity 和 raw error 的快速反馈，同时避免 reject 起点，结果与“质量门控是慢热收益”这一解释一致。由于当前只有 seed 42，这应写成受数据支持的解释，而不是已证明的因果定律。

### 12.3 为什么 70k 不是更好

PPO 的策略与 value optimization、持续更新的自适应采样分布以及难动作集中训练共同形成非平稳过程。60k 后 Macro 和 Joint L2 先恶化，62k–67k 甚至出现 Probe collapse，说明后期可能发生策略遗忘或更新不稳定。论文应强调 checkpoint selection，而不是将最后一个 checkpoint 当作最优模型。

## 13. 可写入论文的结论段

本文构建了面向 PHUMA 大规模动作库的 G1 全身动作跟踪与模块化采样系统。固定 34k 预算的消融表明，多样性预算层级原始误差采样 D-only 在类别均衡的 Macro Success 上优于均匀、全局原始误差和 learning-gap 系列方法。进一步的预算扩展显示，在保持层级原始误差与多样性预算的基础上加入质量门控，可在更长训练后获得更高性能：M7-Raw 在 59k 达到 0.913028 Macro Success，并将 7636 条验证动作中的失败数降至 635。与此同时，60k–70k 的性能回落表明强化学习训练并非单调收敛。整体结果支持一种分阶段认识：多样性约束提供短预算效率，质量门控提高长预算上限，而可靠的 Validation checkpoint selection 对大规模动作跟踪不可或缺。

## 14. 论文结构建议

1. Introduction：大规模异质动作库、训练预算分配问题、现有 uniform/hard-example sampling 的不足、本文贡献。
2. Related Work：人形机器人动作模仿、motion retargeting、大规模多动作策略、curriculum/hard-example mining、diversity-aware sampling。
3. System Overview：PHUMA→G1/WBT 转换、multi-motion loader、tracking MDP 与 PPO。
4. Method：Segment infrastructure、Quality、Difficulty、Online Error/Gap、Diversity，以及 D-only/M7-Raw 公式。
5. Experimental Setup：数据 split、网络、reward、训练预算、评估指标、Macro-first 规则。
6. Results：34k 正式消融、50k 同预算对照、59k 上限、类别结果、训练稳定性。
7. Discussion：预算依赖、质量门控慢热、非单调训练、模块失败/负结果。
8. Limitations and Conclusion：单 seed、Validation-only、无 sim-to-real、缺 M2/M3 正式 Full Val。

## 15. 论文诚信与最终 Test 决策

旧 `paper-final-test-v1` 在 34k 冻结 D-only。后续所有 50k–70k 结果必须标为 **post-freeze exploratory**。如果 Test 从未被访问，可以在论文定稿前创建新的 `paper-final-test-v2`：冻结 M7-Raw `model_59000.pt`、D-only `model_54000.pt`、必要 baselines、代码 commit、checkpoint SHA 和唯一一次 Test 协议。若 Test 已经看过，则不能再用 Test 选择 59k；只能把预算扩展作为 post-freeze exploratory Validation study。

当前仍需明确写出的限制：

- 训练随机性只有 seed 42，尚无多 seed 均值/置信区间；7636 条 motion 的统计不等于训练 seed 的重复性。
- M2/M3 没有完整 34k Full Validation 归档；若论文表题使用“M0–M7 完整消融”，必须补齐。
- Quality/Difficulty/Cluster 生成物中的 `provisional=true` 尚未清理。即使实验实际按 SHA 冻结，论文与归档也应说明采用的固定 config/profile/metadata hash。
- 当前结果是 simulation Validation，不应写成真实机器人性能或 sim-to-real 结论。
- Test 结果目前未纳入本材料；正文主表若要求 Test，需要先完成新的冻结协议。

## 16. 复现身份

- Train manifest SHA256：`51f592792f412c5d5e31caf3621162a3b2ef356e05f98dbce43153d08a59e82d`
- Quality metadata SHA256：`0182c17c01f85d14b8d5d0965b539a99069bd6d8d4346ac3b101a3ecef7b9273`
- Quality config SHA256：`cbee6f0a0496f73245c9dbcdce5774d337cbeb88c5580171398d71014c0273da`
- Difficulty metadata SHA256：`4a59f2ec3231d111794a669272ba7ee431418be6bb1ddd0b55ff51b98f1bb97c`
- Difficulty profile SHA256：`39a1f41f0bc39a38cbcf9797fe32fe88703325d672c278ed670d6615ee9a1c6e`
- Cluster metadata SHA256：`43ccf60af2081a010800027c7cfc2d289839b755fd665b0f984a840cf3083b66`
- Cluster profile SHA256：`f24a0bcd6d94f14db8914e57a56510e0f7aa775715c24a7839ab2ef3853cac1d`
- 旧正式冻结 commit：`833438a9a759bff301fc8050ab9937bd40977c10`
- 59k/54k evaluator 记录 commit：`b11ae3a0eea36c6d668e6255b7804bb9ec6f340a`

## 17. 材料来源

"""
    )
    sections.extend(f"- `{path}`" for path in SOURCE_FILES)
    sections.append(
        "\n## 18. 下一阶段最小清单\n\n"
        "- 决定论文最终方法是否从旧 D-only 34k 更新为 M7-Raw 59k。\n"
        "- 若采用 M7-Raw 59k 且 Test 未访问，先建立 `paper-final-test-v2` 冻结文件，再运行一次 Test。\n"
        "- 仅在论文明确需要‘完整 M0–M7’时补 M2/M3 的 34k Full Validation；无需把所有方法训练到 59k。\n"
        "- 建议增加至少 2 个额外训练 seed，优先 M7-Raw 59k、D-only 34k/54k 和 GlobalRaw 34k。\n"
        "- 从逐 motion CSV 做 paired bootstrap、类别显著性和 failure transition 分析。\n"
        "- 选取 success、near-failure、failure 的可视化案例，形成论文定性图。\n"
    )
    return "\n".join(sections)


def build_docx(
    category_rows: Sequence[dict[str, object]],
    probe_rows: Sequence[dict[str, object]],
    figure_paths: dict[str, Path],
) -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(70)
    run = title.add_run("PHUMA × Whole-Body Tracking")
    set_run_font(run, 26, bold=True, color="17365D")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("论文材料汇总")
    set_run_font(run, 21, bold=True, color="1F4E78")
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = line.add_run("模块设计 · 实验协议 · 完整结果 · 论文叙事")
    set_run_font(run, 12, color="475569")
    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run(f"项目路径：{ROOT}\n版本日期：{DATE}\n材料口径：Validation evidence only")
    set_run_font(run, 10.5, color="64748B")
    doc.add_paragraph()
    add_note(
        doc,
        "重要口径：34k D-only 是旧正式冻结结论；50k–70k 为 post-freeze exploratory 预算研究。当前材料没有使用 Test 结果选择模型。",
    )
    doc.add_page_break()

    doc.add_heading("1. 执行摘要", level=1)
    add_para(doc, "项目已完成从 PHUMA 大规模动作到 Unitree G1/WBT 多动作跟踪策略的完整数据、训练、采样与评估链路。当前最强 Validation checkpoint 是 M7-Raw model_59000.pt。")
    summary_rows = [
        ["旧 34k 正式结论", "D-only model_33500.pt", "Macro 0.899319；冻结选择仍成立"],
        ["当前最佳 Validation", "M7-Raw model_59000.pt", "Micro 0.916841；Macro 0.913028；Failures 635"],
        ["D-only 最佳续训点", "D-only model_54000.pt", "Micro 0.892090；Macro 0.898413；Failures 824"],
        ["上限判断", "59k 附近", "60k Macro 回落，70k 全面退化"],
    ]
    add_table(doc, ["结论层", "模型", "证据"], summary_rows, font_size=9.2, highlight=lambda row: "E2F0D9" if "当前最佳" in str(row[0]) else None)
    add_note(doc, "建议将论文实验分成“固定 34k 正式消融”和“预算扩展分析”两部分，避免改变旧冻结语义。", fill="DDEBF7", color="1F4E78")

    doc.add_heading("2. 论文定位", level=1)
    doc.add_heading("2.1 题目建议", level=2)
    add_bullet(doc, "推荐：面向大规模异质动作的人形机器人全身跟踪：质量门控的多样性约束层级误差采样")
    add_bullet(doc, "English: Quality-Gated Diversity-Constrained Hierarchical Error Sampling for Large-Scale Humanoid Whole-Body Tracking")
    add_bullet(doc, "保守题目：面向大规模人形机器人动作跟踪的多样性预算层级误差采样")
    doc.add_heading("2.2 核心研究问题", level=2)
    add_para(doc, "面对由多个来源构成、难度和质量高度不均衡的大规模动作库，如何分配有限 PPO 训练预算，使策略既持续覆盖不同运动模式，又重点学习当前误差较大的动作和局部片段？")
    doc.add_heading("2.3 摘要初稿", level=2)
    add_para(
        doc,
        "大规模人形机器人动作库包含显著的类别不均衡、轨迹质量差异和动作难度差异，采用均匀采样时，有限的强化学习预算难以同时保证运动多样性覆盖与困难片段学习。本文基于 PHUMA 动作数据构建 Unitree G1 多动作全身跟踪系统，并提出模块化层级采样框架。该框架将动作划分为固定时长片段，在训练集上离线建立轨迹质量标签、固有难度分箱和运动聚类，再根据策略在线跟踪误差在 Cluster、Motion 与 Segment 三个层级分配预算。质量模块只限制异常片段成为 episode 起点，固有难度模块只提供策略无关校准，多样性模块只控制运动簇预算，从而保持各因素解耦。在 6000 条训练动作和 7636 条独立验证动作上的消融中，34k 固定预算下 D-only 取得 0.899319 Macro Success；预算扩展后，M7-Raw 在 59k 达到 0.916841 Micro Success、0.913028 Macro Success 和 0.952066 Completion，失败动作数降至 635。结果表明，多样性约束可提升有限预算下的覆盖，质量门控收益具有预算依赖性，而 60k–70k 的退化说明 PPO checkpoint 性能并非随训练轮数单调增长。",
    )
    doc.add_heading("2.4 建议贡献点", level=2)
    for item in [
        "PHUMA→G1/WBT 的大规模多动作训练管线，支持 76,086 条转换轨迹、manifest 管理与每环境独立 assignment。",
        "质量、固有难度、在线学习状态和运动多样性相互解耦的模块化采样与严格 metadata/checkpoint 身份管理。",
        "Cluster→Motion→Segment→Start Frame 的多样性预算层级采样，以及覆盖 M0–M7 与关键附加消融的实验体系。",
        "揭示方法收益的训练预算依赖性与 59k 附近的非单调性能上限。",
    ]:
        add_bullet(doc, item)

    doc.add_heading("3. 系统与模块设计", level=1)
    add_figure(doc, figure_paths["architecture"], "图 1  PHUMA-WBT 系统与模块化采样框架")
    add_table(doc, ["模块", "核心设计", "功能"], MODULE_TABLE, font_size=8.3)

    doc.add_heading("3.1 数据转换与多动作库", level=2)
    add_para(doc, "PHUMA G1 原始 .npy 经 scripts/phuma_to_npz.py 转换为 WBT .npz，统一为 29 个关节自由度和 30 个 body 的轨迹 schema。MotionLoader 支持单文件、目录递归和 txt manifest；MotionCommand 为每个环境分别采样 motion 与起始帧，形成一个共享的 multi-motion tracking policy，而不是每条 motion 单独训练一个策略。")

    doc.add_heading("3.2 Stage 0：Segment 基础设施", level=2)
    add_code_block(doc, "segment_frames[m] = max(1, round(fps[m] × segment_length_seconds))\nnum_segments[m] = ceil(num_frames[m] / segment_frames[m])\nglobal_id = motion_segment_offsets[m] + local_segment_id")
    add_para(doc, "默认 segment 长度为 1 s，random6000 共 21,575 个 segment。Stage 0 维护 motion/local/global ID 双向映射、合法起点、共享采样计数、概率归一化检查、W&B 标量和 checkpoint 恢复。关闭研究模块时，legacy uniform 路径不增加随机调用。")

    doc.add_heading("3.3 模块 1：轨迹质量审计与 Quality Gate", level=2)
    add_para(doc, "Quality 只识别最终 G1 参考轨迹中的数值、转换、几何和明确物理问题，不读取 policy error、reward 或成功率。17 类审计项覆盖 NaN/Inf、四元数范数、URDF 关节限位、速度一致性、加速度/jerk 孤立峰、位姿连续性、穿地和足滑。")
    add_code_block(doc, "Q(m,s) = clip(1 - Σ_i w_i × severity_i / Σ_i w_i, 0, 1)")
    quality_rows = [
        ["Total", "21,575", "100%"],
        ["Pass", "15,155", "70.243%"],
        ["Borderline", "6,330", "29.340%"],
        ["Reject", "90", "0.417%"],
        ["Eligible start frames", "1,016,958", "99.577%"],
        ["Effective motions", "5,998 / 6,000", "2 empty motions runtime-excluded"],
    ]
    add_table(doc, ["Quality item", "Count", "Ratio / note"], quality_rows, font_size=9)
    add_para(doc, "Gate 的作用域是 assignment start：reject segment 不能被选为 episode 起点，但从合法起点开始的 rollout 可以自然经过 reject 区间。原始 manifest 始终保持 6000 行。")

    doc.add_heading("3.4 模块 2：策略无关固有难度", level=2)
    add_para(doc, "模块二提取 28 项运动学与接触特征，包括 root 线/角速度和加速度、关节速度与活动范围、body 高度、手脚速度、单/双支撑、腾空比例和接触切换。所有时间导数先在完整 motion 上计算，再按 segment 聚合。")
    add_code_block(doc, "center_i = median(x_i)\nscale_i = 1.4826 × MAD(x_i)  (near-zero 时使用 P05–P95 fallback)\nz_i = clip((x_i - center_i) / scale_i, -5, 5)\ndifficulty_score = empirical_CDF(weighted_z) → 10 equal-frequency bins")
    add_para(doc, "Difficulty 不删除样本，也不直接分配训练预算；它只给 M5/M6/M7 提供同难度 error calibration。Quality label 不参与难度拟合。")

    doc.add_heading("3.5 模块 3：在线 Error 与 Learning Gap", level=2)
    add_para(doc, "每个 control step 的 body、joint 和 orientation error 被归因到当前播放的 motion/segment。窗口聚合后以 ρ=0.95 更新 EMA；前 1000 iterations 保持 warmup，每 50 iterations 更新一次概率。Raw error 直接反映当前策略误差；learning gap 则把误差与同一 difficulty bin 比较。")
    add_code_block(doc, "G_global(m,s) = clip((E_segment(m,s) - μ_bin) / σ_bin, -5, 5)\nG_motion(m) = aggregate(max(G_global, 0), outcomes)\nG_local(m,s) = G_global(m,s) - median_j G_global(m,j)")
    add_para(doc, "Adaptive probability 使用 0.15 uniform mix、欠采样 bonus、temperature 和 cap，避免完全丢失探索。M2/M3/M4 只用 raw error；M5/M6/M7 才使用 difficulty-calibrated gap。")

    doc.add_heading("3.6 模块 4：多样性约束", level=2)
    add_para(doc, "模块四把 17 项 segment source feature 聚合为 30 项 motion feature，只在 Train 上 robust-scale 并用 NumPy KMeans++ 建立 K=8 冻结运动簇。类别标签与来源目录仅用于事后诊断，不参与聚类。")
    add_code_block(doc, "P(c) = f/C + (1-f) × n_c^α / Σ_j n_j^α,  f=0.5, α=0.5\nP(c,m,s) = P(c) × P(m|c) × P(s|m)")
    cluster_rows = [[f"cluster_{idx}", size, f"{size / 6000:.3%}"] for idx, size in enumerate([1902, 456, 503, 948, 603, 557, 728, 303])]
    add_table(doc, ["Cluster", "Motions", "Share"], cluster_rows, font_size=8.8)
    add_para(doc, "P(c) 只由 eligible cluster size 与 budget floor 决定，不读取 error、gap、reward 或来源；在线难度信号只控制 P(m|c) 与 P(s|m)，避免相同运动模式被重复放大。")

    doc.add_heading("3.7 跟踪 MDP、网络与优化", level=2)
    add_table(doc, ["配置", "值"], PPO_CONFIG, font_size=8.7)
    add_table(doc, ["Reward", "Weight", "Meaning"], REWARDS, font_size=8.5)
    add_para(doc, "Actor 输入由 reference command、anchor target、base velocity、joint state 和上一动作组成；Critic 额外使用 body pose/orientation privileged observations。动作为 29 维关节位置目标。主要提前终止条件为 anchor z、anchor orientation 以及手腕/脚踝末端 body z 误差。")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("4. 实验协议", level=1)
    add_table(doc, ["Split", "Manifest", "Motions", "SHA256"], SPLITS, font_size=7.4)
    add_para(doc, "Train、Validation 和 Test 按 manifest SHA256 固定且互不重叠；Probe500 是 Validation Full 子集。所有正式方法使用同一 random6000、seed 42、网络、PPO、reward、termination 和 evaluator。")
    doc.add_heading("4.1 方法矩阵", level=2)
    add_table(doc, ["Method", "Motion", "Segment", "Quality", "Difficulty", "Diversity", "Evidence"], METHOD_MATRIX, font_size=7.4)
    add_note(doc, "D-only = M4 Raw/Raw + Diversity；M7-Raw = D-only + Quality Gate。M7-Raw 不使用 Learning Gap。", fill="DDEBF7", color="1F4E78")
    doc.add_heading("4.2 评估指标", level=2)
    add_table(doc, ["Metric", "Definition"], METRICS, font_size=8.7)
    add_para(doc, "Macro Success 的实现是 17 个 category success rate 的无权平均，不是 source-group 平均。失败 checkpoint 可能因极早终止而显示较低误差，因此 Body/Joint 只能与 Success/Completion 联合解释。")
    doc.add_heading("4.3 Macro-first 选择规则", level=2)
    add_para(doc, "按 Macro、Micro、Completion、Body Error、较早 checkpoint 依次比较。评估器 epsilon=0.002；只有当前指标差值超过 epsilon 才直接决胜，否则进入下一层。Joint L2 不参与 checkpoint selection。Probe500 只筛选候选，论文主结论使用 Full Validation。")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.sections[-1].orientation = WD_ORIENT.LANDSCAPE
    doc.sections[-1].page_width, doc.sections[-1].page_height = doc.sections[-1].page_height, doc.sections[-1].page_width
    doc.sections[-1].left_margin = Cm(1.3)
    doc.sections[-1].right_margin = Cm(1.3)
    doc.add_heading("5. 34k 正式 Full Validation", level=1)
    formal_rows = [
        [
            row["method"],
            row["checkpoint"],
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in FORMAL_RESULTS
    ]
    add_table(
        doc,
        ["Method", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body Err ↓", "Joint L2 ↓", "Failures ↓"],
        formal_rows,
        font_size=8.2,
        highlight=lambda row: "E2F0D9" if row[0] == "D-only" else None,
    )
    add_figure(doc, figure_paths["formal"], "图 2  34k 正式 Full Validation 的 Micro/Macro Success", width=8.8)
    add_para(doc, "D-only 相比 GlobalRaw：Micro +2.095 pp、Macro +1.792 pp、Completion +1.151 pp、Failures -160。M7-Raw 34k 虽在 Micro、Body、Joint 和 Failures 略优，但 Macro 低 0.629 pp，因此 Macro-first 选择 D-only 合理。")
    add_note(doc, "M2/M3 当前只有 pilot，没有独立 7636-motion Full Validation。若正文声称“完整 M0–M7 消融”，应补两者的 34k 同预算评估；无需全部续训到 59k。")

    doc.add_heading("6. 预算扩展 Full Validation", level=1)
    budget_rows = [
        [
            row["method"],
            row["budget"],
            row["checkpoint"],
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in BUDGET_RESULTS
    ]
    add_table(
        doc,
        ["Method", "Budget", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body Err ↓", "Joint L2 ↓", "Failures ↓"],
        budget_rows,
        font_size=7.7,
        highlight=lambda row: "E2F0D9" if row[0] == "M7-Raw" and "59k" in str(row[1]) else ("DDEBF7" if row[0] == "D-only" and "54k" in str(row[1]) else None),
    )
    add_figure(doc, figure_paths["budget"], "图 3  Full Validation 预算扩展曲线", width=8.9)

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.sections[-1].orientation = WD_ORIENT.PORTRAIT
    doc.sections[-1].page_width, doc.sections[-1].page_height = doc.sections[-1].page_height, doc.sections[-1].page_width
    doc.sections[-1].top_margin = Cm(1.8)
    doc.sections[-1].bottom_margin = Cm(1.7)
    doc.sections[-1].left_margin = Cm(1.8)
    doc.sections[-1].right_margin = Cm(1.8)
    doc.add_heading("7. 关键差值与结论", level=1)
    delta_rows = [
        ["M7-Raw 59k - M7-Raw 34k", "+0.031954", "+0.020002", "+0.018014", "-0.002195", "-0.046551", "-244"],
        ["M7-Raw 59k - D-only 54k", "+0.024751", "+0.014615", "+0.014103", "-0.002523", "-0.061182", "-189"],
        ["D-only 54k - D-only 34k", "+0.007596", "-0.000906", "+0.003894", "-0.000268", "+0.008179", "-58"],
        ["M7-Raw 60k - 59k", "-0.000785", "-0.013409", "-0.000472", "+0.001123", "+0.031116", "+6"],
        ["M7-Raw 70k - 59k", "-0.021346", "-0.014415", "-0.012510", "+0.001296", "+0.053087", "+163"],
    ]
    add_table(doc, ["Comparison", "Δ Micro", "Δ Macro", "Δ Completion", "Δ Body", "Δ Joint", "Δ Failures"], delta_rows, font_size=8.0, highlight=lambda row: "E2F0D9" if "59k - D-only" in str(row[0]) else None)
    add_para(doc, "当前证据支持两个训练区间：D-only 在短预算下效率更高且冻结合理；M7-Raw 在 50k 后明显反超，并在 59k 达到更高上限。60k 与 70k 的退化否定了“更长训练必然更好”。")

    doc.add_heading("8. Dense Probe 与稳定性", level=1)
    add_figure(doc, figure_paths["probe"], "图 4  M7-Raw 50k–70k Dense Probe；阴影为明显崩溃区间")
    selected_probe = [row for row in probe_rows if int(row["iteration"]) in {49999, 50500, 55000, 59000, 60000, 62000, 65000, 67000, 69000, 69999}]
    add_table(
        doc,
        ["Iter", "Micro", "Macro", "Completion", "Body", "Joint", "Fail/500"],
        [[row["iteration"], fmt(float(row["micro"])), fmt(float(row["macro"])), fmt(float(row["completion"])), fmt(float(row["body"])), fmt(float(row["joint"])), row["failures"]] for row in selected_probe],
        font_size=8.2,
        highlight=lambda row: "E2F0D9" if int(row[0]) == 59000 else ("FCE8E6" if 62000 <= int(row[0]) <= 67000 else None),
    )
    add_para(doc, "62k–67k 的 Probe 成功率接近零，之后又部分恢复，说明训练状态明显非单调。极早终止会让误差均值虚低，因此此类 checkpoint 必须首先按 Success/Completion 判为失败。")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.sections[-1].orientation = WD_ORIENT.LANDSCAPE
    doc.sections[-1].page_width, doc.sections[-1].page_height = doc.sections[-1].page_height, doc.sections[-1].page_width
    doc.sections[-1].left_margin = Cm(1.3)
    doc.sections[-1].right_margin = Cm(1.3)
    doc.add_heading("9. 17 类动作结果", level=1)
    category_table = [
        [
            row["category"],
            row["motions"],
            fmt(float(row["d_success"])),
            fmt(float(row["m_success"])),
            percent_point(float(row["delta"])),
            row["d_failures"],
            row["m_failures"],
            fmt(float(row["d_joint"])),
            fmt(float(row["m_joint"])),
        ]
        for row in category_rows
    ]
    add_table(
        doc,
        ["Category", "N", "D Success", "M7-Raw Success", "Δ", "D Fail", "M Fail", "D Joint", "M Joint"],
        category_table,
        font_size=7.5,
        highlight=lambda row: "E2F0D9" if str(row[4]).startswith("+") else ("FCE8E6" if str(row[4]).startswith("-") else None),
    )
    add_figure(doc, figure_paths["category"], "图 5  M7-Raw 59k 与 D-only 54k 的类别成功率", width=8.9)
    add_para(doc, "M7-Raw 在 17 类中提高 10 类、持平 4 类、降低 3 类。最大贡献来自 fitness：失败数从 390 降至 248，减少 142；小样本类别的差异应谨慎解释。")

    doc.add_heading("10. 补充消融与负结果", level=1)
    doc.add_heading("10.1 formal_v2 诊断组合（Full Validation）", level=2)
    supplementary_rows = [
        [
            row["method"],
            row["design"],
            row["checkpoint"],
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in SUPPLEMENTARY_FULL_RESULTS
    ]
    add_table(
        doc,
        ["Method", "Design", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body ↓", "Joint ↓", "Failures ↓"],
        supplementary_rows,
        font_size=7.1,
    )
    add_para(doc, "V2-Diag-A 的 motion 层使用 raw error、segment 层使用 relative learning gap；M7-v2 对 raw segment error 施加 lambda=0.10 的 gap correction。两者的 Macro 均低于 D-only 34k，说明当前预算下更复杂的 gap 校准未稳定优于 Raw/Raw。")

    doc.add_heading("10.2 M7-JGap 10k 固定 Probe", level=2)
    joint_gap_rows = [
        [
            row["method"],
            row["checkpoint"],
            fmt(float(row["micro"])),
            fmt(float(row["macro"])),
            fmt(float(row["completion"])),
            fmt(float(row["body"])),
            fmt(float(row["joint"])),
            row["failures"],
        ]
        for row in JOINT_GAP_PROBE_RESULTS
    ]
    add_table(
        doc,
        ["Method", "Checkpoint", "Micro ↑", "Macro ↑", "Completion ↑", "Body ↓", "Joint ↓", "Fail/500 ↓"],
        joint_gap_rows,
        font_size=7.7,
        highlight=lambda row: "FCE8E6" if str(row[0]).startswith("M7-JGap") else None,
    )
    add_note(doc, "M7-JGap 仅为 10k、16-env 的固定 Probe 对照，不能与 7636-motion Full Validation 主表横向比较。它虽使 Macro +0.441 pp，但 Micro -1.600 pp、Completion -1.318 pp、Joint L2 变差约 1.064%，最终 gate 为 NO-GO: DOMINATED，未进入 34k。")

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.sections[-1].orientation = WD_ORIENT.PORTRAIT
    doc.sections[-1].page_width, doc.sections[-1].page_height = doc.sections[-1].page_height, doc.sections[-1].page_width
    doc.sections[-1].top_margin = Cm(1.8)
    doc.sections[-1].bottom_margin = Cm(1.7)
    doc.sections[-1].left_margin = Cm(1.8)
    doc.sections[-1].right_margin = Cm(1.8)
    doc.add_heading("11. 论文讨论素材", level=1)
    discussion = [
        ("D-only 的短预算优势", "Raw error 反馈直接，多样性 budget 抑制相似高误差动作垄断预算，因此有限迭代内更快体现收益。"),
        ("M7-Raw 的长预算优势", "Quality Gate 只排除很少起点，单次影响小，但数十亿 transitions 下异常起点的累计成本可观；结果与慢热型收益一致。"),
        ("Learning Gap 未成为最强方案", "难度 bin、EMA 和 outcome calibration 增加估计延迟与噪声；M5/M6/M7 在 34k 未超过更简单的 raw-error variants，这是有价值的负结果。"),
        ("59k 后退化", "PPO 与持续变化的采样分布组成非平稳系统，后期可能出现遗忘或不稳定更新；密集 checkpoint 与独立 Validation 选择是方法的一部分。"),
    ]
    for title_text, body in discussion:
        doc.add_heading(title_text, level=2)
        add_para(doc, body)

    doc.add_heading("12. 可直接写入的结论段", level=1)
    add_para(doc, "本文构建了面向 PHUMA 大规模动作库的 G1 全身动作跟踪与模块化采样系统。固定 34k 预算的消融表明，多样性预算层级原始误差采样 D-only 在类别均衡的 Macro Success 上优于均匀、全局原始误差和 learning-gap 系列方法。预算扩展进一步显示，在保持层级原始误差与多样性预算的基础上加入质量门控，可在更长训练后获得更高性能：M7-Raw 在 59k 达到 0.913028 Macro Success，并将 7636 条验证动作中的失败数降至 635。与此同时，60k–70k 的回落表明强化学习训练并非单调收敛。整体结果支持一种分阶段认识：多样性约束提供短预算效率，质量门控提高长预算上限，而可靠的 Validation checkpoint selection 对大规模动作跟踪不可或缺。")

    doc.add_heading("13. 论文结构与图表清单", level=1)
    outline = [
        ["1 Introduction", "问题、动机、贡献"],
        ["2 Related Work", "动作模仿、多动作策略、curriculum/hard mining、多样性采样"],
        ["3 System", "PHUMA→WBT、MotionLoader、tracking MDP"],
        ["4 Method", "Stage 0、Quality、Difficulty、Error/Gap、Diversity"],
        ["5 Experiments", "split、PPO、metrics、Macro-first"],
        ["6 Results", "34k 消融、预算扩展、类别、稳定性"],
        ["7 Discussion", "预算依赖、负结果、非单调性"],
        ["8 Conclusion", "总结与限制"],
    ]
    add_table(doc, ["Section", "内容"], outline, font_size=9)
    for item in [
        "主图：系统与采样框架（已生成）",
        "主表：34k Full Validation 消融（已生成）",
        "主图：预算扩展曲线（已生成）",
        "主图/附录：17 类动作成功率（已生成）",
        "附录：Dense Probe 与 collapse 区间（已生成）",
        "待补：代表性 success/near-failure/failure 视频帧",
        "待补：若执行 Test，则 paired bootstrap 和 per-motion win/loss 分析",
    ]:
        add_bullet(doc, item)

    doc.add_heading("14. 局限、缺口与 Test 冻结", level=1)
    limitations = [
        "目前正式训练随机性只有 seed 42；应优先给最终方法、D-only 和强 baseline 增加至少两个 seed。",
        "M2/M3 缺少独立 34k Full Validation。论文若写“完整 M0–M7 消融”，需要补齐；无需把所有方法续训至 59k。",
        "Quality/Difficulty/Cluster 生成物仍带 provisional 标记。正式论文应保留并报告实际使用的 config/profile/metadata SHA。",
        "当前证据属于 simulation Validation，不代表真实机器人或 sim-to-real 性能。",
        "Test 结果未纳入本材料；不得根据 Test 选择 checkpoint 或修改方法。",
    ]
    for item in limitations:
        add_bullet(doc, item)
    add_note(
        doc,
        "若 Test 尚未访问：先创建 paper-final-test-v2，冻结 M7-Raw model_59000.pt、D-only model_54000.pt、基线、commit、checkpoint SHA 和单次 Test 协议。若 Test 已访问：预算扩展只能作为 post-freeze exploratory Validation study。",
        fill="FCE8E6",
        color="9C0006",
    )

    doc.add_heading("15. 复现身份与材料索引", level=1)
    hashes = [
        ["Train manifest", "51f592792f412c5d5e31caf3621162a3b2ef356e05f98dbce43153d08a59e82d"],
        ["Quality metadata", "0182c17c01f85d14b8d5d0965b539a99069bd6d8d4346ac3b101a3ecef7b9273"],
        ["Quality config", "cbee6f0a0496f73245c9dbcdce5774d337cbeb88c5580171398d71014c0273da"],
        ["Difficulty metadata", "4a59f2ec3231d111794a669272ba7ee431418be6bb1ddd0b55ff51b98f1bb97c"],
        ["Difficulty profile", "39a1f41f0bc39a38cbcf9797fe32fe88703325d672c278ed670d6615ee9a1c6e"],
        ["Cluster metadata", "43ccf60af2081a010800027c7cfc2d289839b755fd665b0f984a840cf3083b66"],
        ["Cluster profile", "f24a0bcd6d94f14db8914e57a56510e0f7aa775715c24a7839ab2ef3853cac1d"],
        ["旧正式冻结 commit", "833438a9a759bff301fc8050ab9937bd40977c10"],
        ["59k/54k evaluator commit", "b11ae3a0eea36c6d668e6255b7804bb9ec6f340a"],
    ]
    add_table(doc, ["Artifact", "SHA / commit"], hashes, font_size=7.6)
    for path in SOURCE_FILES:
        add_bullet(doc, path)

    doc.add_heading("16. 下一阶段最小行动清单", level=1)
    for item in [
        "决定最终论文采用 M7-Raw 59k 还是保持旧 D-only 34k 为主方法。",
        "若采用 M7-Raw 且 Test 未访问，先建立新的 paper-final-test-v2 冻结文件。",
        "只在需要完整八项消融时补 M2/M3 的 34k Full Validation。",
        "增加多 seed，优先 M7-Raw、D-only 和 GlobalRaw。",
        "生成 paired bootstrap、failure reason 和定性可视化材料。",
    ]:
        add_bullet(doc, item)

    configure_headers(doc)
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def write_tables(category_rows: Sequence[dict[str, object]], probe_rows: Sequence[dict[str, object]]) -> None:
    write_csv(
        TABLES / "formal_validation_34k.csv",
        ["method", "checkpoint", "micro", "macro", "completion", "body_error_m", "joint_l2_rad", "joint_rms_rad", "failures"],
        [[row[key] for key in ["method", "checkpoint", "micro", "macro", "completion", "body", "joint", "rms", "failures"]] for row in FORMAL_RESULTS],
    )
    write_csv(
        TABLES / "budget_extension_validation.csv",
        ["method", "budget", "iteration", "checkpoint", "micro", "macro", "completion", "body_error_m", "joint_l2_rad", "joint_rms_rad", "failures", "phase"],
        [[row[key] for key in ["method", "budget", "iteration", "checkpoint", "micro", "macro", "completion", "body", "joint", "rms", "failures", "phase"]] for row in BUDGET_RESULTS],
    )
    write_csv(
        TABLES / "category_comparison_m7raw59k_vs_donly54k.csv",
        ["category", "motions", "donly_success", "m7raw_success", "success_delta", "donly_completion", "m7raw_completion", "donly_joint_l2", "m7raw_joint_l2", "donly_failures", "m7raw_failures"],
        [[row[key] for key in ["category", "motions", "d_success", "m_success", "delta", "d_completion", "m_completion", "d_joint", "m_joint", "d_failures", "m_failures"]] for row in category_rows],
    )
    write_csv(
        TABLES / "m7raw_probe500_dense_50k_70k.csv",
        ["iteration", "checkpoint", "micro", "macro", "completion", "body_error_m", "joint_l2_rad", "failures"],
        [[row[key] for key in ["iteration", "checkpoint", "micro", "macro", "completion", "body", "joint", "failures"]] for row in probe_rows],
    )
    write_csv(
        TABLES / "supplementary_formal_v2_validation.csv",
        ["method", "design", "budget", "iteration", "checkpoint", "micro", "macro", "completion", "body_error_m", "joint_l2_rad", "joint_rms_rad", "failures", "source"],
        [[row[key] for key in ["method", "design", "budget", "iteration", "checkpoint", "micro", "macro", "completion", "body", "joint", "rms", "failures", "source"]] for row in SUPPLEMENTARY_FULL_RESULTS],
    )
    write_csv(
        TABLES / "joint_gap_probe500_10k.csv",
        ["method", "budget", "checkpoint", "micro", "macro", "completion", "body_error_m", "joint_l2_rad", "joint_rms_rad", "failures"],
        [[row[key] for key in ["method", "budget", "checkpoint", "micro", "macro", "completion", "body", "joint", "rms", "failures"]] for row in JOINT_GAP_PROBE_RESULTS],
    )


def write_manifest() -> None:
    payload = {
        "generated_at": DATE,
        "project_root": str(ROOT),
        "docx": str(DOCX_PATH.relative_to(ROOT)),
        "markdown": str(MD_PATH.relative_to(ROOT)),
        "tables": sorted(str(path.relative_to(ROOT)) for path in TABLES.glob("*.csv")),
        "figures": sorted(str(path.relative_to(ROOT)) for path in FIGURES.glob("*.png")),
        "scope": "Local project design and Validation evidence; Test not used for selection.",
    }
    (ASSETS / "manifest.json").write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    REPORTS.mkdir(parents=True, exist_ok=True)
    TABLES.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    validate_source_results()
    category_rows = category_comparison()
    probe_rows = dense_probe_rows()
    if len(category_rows) != 17:
        raise RuntimeError(f"Expected 17 category rows, found {len(category_rows)}")
    if not any(int(row["iteration"]) == 59000 for row in probe_rows):
        raise RuntimeError("Dense probe rows do not contain model_59000.pt")

    font_prop = configure_plots()
    figure_paths = {
        "architecture": plot_architecture(font_prop),
        "formal": plot_formal_results(font_prop),
        "budget": plot_budget_results(font_prop),
        "probe": plot_dense_probe(probe_rows, font_prop),
        "category": plot_category_results(category_rows, font_prop),
    }
    write_tables(category_rows, probe_rows)
    MD_PATH.write_text(build_markdown(category_rows, probe_rows), encoding="utf-8")
    build_docx(category_rows, probe_rows, figure_paths)
    write_manifest()
    print(f"Wrote {DOCX_PATH}")
    print(f"Wrote {MD_PATH}")
    print(f"Wrote {len(list(TABLES.glob('*.csv')))} tables and {len(list(FIGURES.glob('*.png')))} figures")


if __name__ == "__main__":
    main()
