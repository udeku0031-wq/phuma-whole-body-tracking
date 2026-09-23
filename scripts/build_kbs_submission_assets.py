#!/usr/bin/env python3
"""Rebuild print-size English paper assets without launching training or evaluation."""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
import shutil
import subprocess
import textwrap
import zipfile

import build_kbs_paper_package as data
import kbs_paper_text as prose
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.lines import Line2D
from matplotlib.patches import FancyArrowPatch, Rectangle
import numpy as np
from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/kbs_submission_materials_2026-09-11"
FIG = OUT / "figures"
TABLE = OUT / "tables"
EVIDENCE = OUT / "evidence"
DATE = "2026-09-11"
WIDTH_MM = 190
BLUE, VERMILION, GREEN = "#0072B2", "#D55E00", "#009E73"
GRAY, BLACK = "#737373", "#202020"
ART = {}
CAPTIONS = {}
FIGURE_AUDIT = []
BOX_CHECKS = []
SOURCES = {
    "KBS author guide (access blocked; recheck journal-specific rules)":
        "https://www.sciencedirect.com/journal/knowledge-based-systems/publish/guide-for-authors",
    "Elsevier artwork formats":
        "https://www.elsevier.com/about/policies-and-standards/author/artwork-and-media-instructions/artwork-formats-checklist",
    "Elsevier artwork sizing":
        "https://www.elsevier.com/about/policies-and-standards/author/artwork-and-media-instructions/artwork-sizing",
    "Elsevier graphical abstracts":
        "https://www.elsevier.com/researcher/author/tools-and-resources/graphical-abstract",
    "Elsevier highlights":
        "https://www.elsevier.support/publishing/answer/how-do-i-include-highlights-with-my-manuscript",
    "Elsevier AI policy":
        "https://www.elsevier.com/about/policies-and-standards/generative-ai-policies-for-journals",
}
AI_CAPTION = " Diagram layout code was prepared with OpenAI Codex (GPT-5, OpenAI)."


def dump(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=True) + "\n", encoding="utf-8")


def configure():
    for path in (OUT, FIG, TABLE, EVIDENCE, OUT / "source", OUT / "preview"):
        path.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({
        "font.family": "Liberation Sans", "font.size": 9,
        "axes.labelsize": 9, "axes.titlesize": 10, "axes.titleweight": "bold",
        "xtick.labelsize": 8.5, "ytick.labelsize": 8.5, "legend.fontsize": 8.5,
        "axes.linewidth": 0.65, "lines.linewidth": 1.2, "lines.markersize": 4.5,
        "pdf.fonttype": 42, "ps.fonttype": 42, "svg.fonttype": "none",
        "mathtext.fontset": "dejavusans", "axes.unicode_minus": False,
        "savefig.facecolor": "white", "figure.facecolor": "white",
    })


def clean(ax, axis="x"):
    ax.spines[["top", "right"]].set_visible(False)
    ax.grid(axis=axis, color="#dddddd", linewidth=0.5)
    ax.set_axisbelow(True)
    ax.tick_params(length=3, width=0.6)


def box(ax, x, y, w, h, title, body="", color=BLACK):
    patch = Rectangle((x, y), w, h, facecolor="white", edgecolor=color, linewidth=1)
    ax.add_patch(patch)
    items = [ax.text(x + w / 2, y + h * (0.68 if body else 0.5), title,
                     ha="center", va="center", weight="bold", color=color, fontsize=9)]
    if body:
        items.append(ax.text(x + w / 2, y + h * 0.29, body,
                             ha="center", va="center", fontsize=8.5, linespacing=1.3))
    BOX_CHECKS.append((ax.figure, patch, items))


def arrow(ax, start, end, color=BLACK, dashed=False):
    ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=10,
                                linewidth=1, color=color,
                                linestyle="--" if dashed else "-"))


def diagram(height_mm=122):
    fig = plt.figure(figsize=(WIDTH_MM / 25.4, height_mm / 25.4))
    ax = fig.add_axes([0.02, 0.02, 0.96, 0.96])
    ax.set(xlim=(0, 100), ylim=(0, 100))
    ax.axis("off")
    return fig, ax


def save(fig, stem, caption, *, width_mm=WIDTH_MM):
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    clipped = []
    for ax in fig.axes:
        axis_text = [ax.title, ax.xaxis.label, ax.yaxis.label,
                     *ax.get_xticklabels(), *ax.get_yticklabels()] if ax.axison else []
        for item in [*ax.texts, *axis_text]:
            if not item.get_visible() or not item.get_text():
                continue
            b = item.get_window_extent(renderer)
            if b.x0 < -1 or b.y0 < -1 or b.x1 > fig.bbox.width + 1 or b.y1 > fig.bbox.height + 1:
                clipped.append(item.get_text())
    for owner, patch, items in BOX_CHECKS:
        if owner is fig:
            bound = patch.get_window_extent(renderer)
            for item in items:
                b = item.get_window_extent(renderer)
                if not (bound.x0 + 1 <= b.x0 and b.x1 <= bound.x1 - 1 and
                        bound.y0 + 1 <= b.y0 and b.y1 <= bound.y1 - 1):
                    clipped.append("box overflow: " + item.get_text())
    if clipped:
        raise RuntimeError(f"{stem}: text outside canvas/container: {clipped}")
    paths = {}
    for extension in ("pdf", "eps", "svg"):
        paths[extension] = FIG / f"{stem}.{extension}"
        fig.savefig(paths[extension])
    paths["png"] = FIG / f"{stem}.png"
    fig.savefig(paths["png"], dpi=300)
    paths["tiff"] = FIG / f"{stem}.tiff"
    fig.savefig(paths["tiff"], dpi=1000, pil_kwargs={"compression": "tiff_lzw"})
    with Image.open(paths["tiff"]) as raster:
        rgb = raster.convert("RGB")
    rgb.save(paths["tiff"], compression="tiff_lzw", dpi=(1000,1000))
    rgb.close()
    with Image.open(paths["tiff"]) as raster:
        raster_info = {"pixels": list(raster.size), "dpi": [float(d) for d in raster.info["dpi"]], "mode": raster.mode}
    paths["caption"] = FIG / f"{stem}_caption.txt"
    paths["caption"].write_text(caption + "\n", encoding="utf-8")
    text_sizes = [t.get_fontsize() for ax in fig.axes for t in ax.texts if t.get_visible()]
    FIGURE_AUDIT.append({"figure": stem, "width_mm": width_mm,
                         "height_mm": float(fig.get_size_inches()[1] * 25.4),
                         "text_overflow": clipped, "minimum_annotation_font_pt": min(text_sizes or [8.5]),
                         "raster": raster_info})
    ART[stem], CAPTIONS[stem] = paths, caption
    plt.close(fig)


def plot_framework():
    fig, ax = diagram(156)
    ax.set_ylim(-28,100)
    ax.text(1, 97, "(a) Offline reference knowledge: training data only", weight="bold", fontsize=10)
    box(ax, 1, 74, 24, 15, "PHUMA references", "Conversion + 1-s segments")
    box(ax, 37, 78, 27, 11, "Quality audit", "Eligible episode starts", GREEN)
    box(ax, 72, 78, 27, 11, "Motion clustering", "Kinematic features; K = 8", BLUE)
    arrow(ax, (25, 82), (37, 84))
    ax.plot([29, 29, 85.5], [82, 93, 93], color=BLACK, lw=1)
    arrow(ax, (85.5, 93), (85.5, 89))
    ax.text(38, 65, "(b) On-policy training", weight="bold", fontsize=10)
    box(ax, 1, 41, 28, 17, "Hierarchical sampler", "Cluster / motion / segment", BLUE)
    box(ax, 38, 41, 26, 17, "Reference assignment", "Motion, segment, start frame")
    box(ax, 73, 41, 26, 17, "PPO + G1 simulation", "Shared 29-action policy", GREEN)
    arrow(ax, (29, 49.5), (38, 49.5))
    arrow(ax, (64, 49.5), (73, 49.5))
    ax.plot([50.5, 50.5, 12], [78, 69, 69], color=GREEN, lw=1)
    arrow(ax, (12, 69), (12, 58), GREEN)
    ax.plot([85.5, 85.5, 22], [78, 63, 63], color=BLUE, lw=1)
    arrow(ax, (22, 63), (22, 58), BLUE)
    box(ax, 37, 15, 28, 15, "Online error statistics", "Segment and motion EMAs", VERMILION)
    ax.plot([86, 86, 69], [41, 22.5, 22.5], color=VERMILION, lw=1)
    arrow(ax, (69, 22.5), (65, 22.5), VERMILION)
    ax.plot([37, 15, 15], [22.5, 22.5, 35], color=VERMILION, lw=1)
    arrow(ax, (15, 35), (15, 41), VERMILION)
    ax.text(51, 7, "Errors update conditional priorities; they do not change cluster budgets.", ha="center", fontsize=8.5)
    ax.plot([1,99],[1,1],color="#aaaaaa",linewidth=.65)
    ax.text(1,-7,"(c) Separate checkpoint evaluation",weight="bold",fontsize=10)
    box(ax,1,-24,28,12,"Validation","Probe + full Validation",BLUE)
    box(ax,38,-24,26,12,"Freeze selection","Selected checkpoint")
    box(ax,73,-24,26,12,"Final Test","No training feedback",GREEN)
    arrow(ax,(29,-18),(38,-18)); arrow(ax,(64,-18),(73,-18))
    save(fig, "Fig_01_Framework",
         "Fig. 1. Training architecture of quality-gated diversity-constrained hierarchical error "
         "sampling (QD-HES; repository variant M7-Raw). (a) Only training references are used for "
         "offline quality and cluster metadata. (b) A cluster-motion-segment sampler assigns episode "
         "starts to the shared PPO controller. Observed tracking statistics update conditional "
         "priorities, whereas cluster budgets remain error-independent. Difficulty calibration is "
         "an optional ablation branch and is not active in QD-HES. (c) Validation and Test evaluate "
         "saved checkpoints outside the training feedback loop." + AI_CAPTION)


def plot_sampler():
    fig, ax = diagram(120)
    ax.text(2, 96, "(a) Select a cluster", fontsize=10, weight="bold")
    ax.text(35, 96, "(b) Select a motion", fontsize=10, weight="bold")
    ax.text(68, 96, "(c) Select a segment", fontsize=10, weight="bold")
    box(ax, 2, 69, 28, 18, "Diversity budget", "Eligible cluster sizes", BLUE)
    box(ax, 35, 69, 28, 18, "Motion priority", "Aggregated raw error", VERMILION)
    box(ax, 68, 69, 30, 18, "Segment priority", "Raw segment error", VERMILION)
    arrow(ax, (30, 78), (35, 78)); arrow(ax, (63, 78), (68, 78))
    ax.text(16, 57, r"$P(c)=\frac{f}{C}+(1-f)\frac{n_c^\alpha}{\sum_j n_j^\alpha}$", ha="center", fontsize=10)
    ax.text(49, 57, r"$P(m\mid c)$", ha="center", fontsize=12)
    ax.text(83, 57, r"$P(s\mid m)$", ha="center", fontsize=12)
    ax.text(16, 45, r"$f=0.5,\ \alpha=0.5$", ha="center", fontsize=9)
    ax.text(49, 45, "Softmax + exploration\n+ feasible probability cap", ha="center", fontsize=8.5)
    ax.text(83, 45, "Softmax + exploration\n+ quality eligibility mask", ha="center", fontsize=8.5)
    ax.plot([2, 98], [35, 35], color="#aaaaaa", linewidth=.65)
    ax.text(50, 26, r"$P(c,m,s)=P(c)\,P(m\mid c)\,P(s\mid m)$", ha="center", fontsize=13)
    ax.text(50, 13, "Uniform legal frame within the chosen segment; rejected starts have zero probability.",
            ha="center", fontsize=8.5)
    save(fig, "Fig_02_Sampling_Mechanism",
         "Fig. 2. Factorization of the QD-HES episode-start distribution. Cluster probability depends "
         "on eligible motion count n_c, with a uniform-share floor f/C; error priorities act only "
         "inside a cluster and inside a motion. Conditional priorities mix 15% eligible-uniform "
         "exploration with adaptive softmax probabilities and apply feasible probability caps. "
         "The diagram is a mechanism schematic, not a visualization of measured probability values." + AI_CAPTION)


def plot_metadata(stats):
    q = stats["quality"]
    fig, axs = plt.subplots(1, 2, figsize=(WIDTH_MM/25.4, 82/25.4))
    fig.subplots_adjust(left=.09, right=.98, bottom=.20, top=.85, wspace=.43)
    values = [q["pass_count"], q["borderline_count"], q["reject_count"]]
    colors = [GREEN, BLUE, VERMILION]
    axs[0].bar([0,1,2], values, color=colors, width=.6, edgecolor=BLACK, linewidth=.4)
    axs[0].set(xticks=[0,1,2], xticklabels=["Pass", "Borderline", "Reject"], ylim=(0,18500), ylabel="Training segments")
    for i, n in enumerate(values):
        axs[0].text(i, n+450, f"{n:,}\n({n/21575*100:.2f}%)", ha="center", fontsize=8.5)
    axs[0].set_title("(a) Quality audit", loc="left", pad=12)
    sizes = np.array([1902,456,503,948,603,557,728,303])
    axs[1].bar(np.arange(1,9), sizes, color=BLUE, edgecolor=BLACK, linewidth=.4, width=.65)
    axs[1].set(xticks=np.arange(1,9), ylim=(0,2250), xlabel="Cluster identifier", ylabel="Training motions")
    for i,n in enumerate(sizes):
        axs[1].text(i+1, n+55, str(n), ha="center", fontsize=8.5)
    axs[1].set_title("(b) Motion clusters", loc="left", pad=12)
    for ax in axs: clean(ax, "y")
    save(fig, "Fig_03_Training_Metadata",
         "Fig. 3. Offline training metadata. (a) Quality labels for 21,575 nominal one-second "
         "segments from 6,000 training motions; borderline starts remain eligible. (b) Raw sizes "
         "of the eight unsupervised motion clusters before applying reset eligibility. Runtime "
         "cluster budgets use eligible counts rather than these unfiltered counts. These are "
         "observed counts, not semantic class assignments.")


def plot_ablation(formal):
    labels = ["M0: uniform", "M1: quality only", "GlobalRaw", "GlobalRaw-Q", "M4: hierarchical raw",
              "M5: learning gap", "M6: quality + gap", "M7: gap + diversity", "D-only", "M7-Raw / QD-HES"]
    fig, ax = plt.subplots(figsize=(WIDTH_MM/25.4, 111/25.4))
    fig.subplots_adjust(left=.28, right=.86, bottom=.17, top=.92)
    y = np.arange(len(formal))
    ax.scatter([r["micro"]*100 for r in formal], y+.12, color=BLUE, marker="o", s=28, label="Micro success")
    ax.scatter([r["macro"]*100 for r in formal], y-.12, color=VERMILION, marker="s", s=28, label="Macro success")
    ax.set(yticks=y, yticklabels=labels, xlabel="Full-validation success (%)", xlim=(84,91), ylim=(9.7,-.7))
    for i, r in enumerate(formal):
        ax.text(91.25, i, f"{r['macro']*100:.2f}", va="center", fontsize=8.5)
    ax.text(91.25,-.92,"Macro (%)", fontsize=8.5, weight="bold")
    ax.legend(loc="lower left", bbox_to_anchor=(0,1.02), ncol=2, frameon=False)
    clean(ax)
    save(fig, "Fig_04_Validation_Ablation",
         "Fig. 4. Available full-validation ablations after training with a 34k-iteration budget "
         "(7,636 motions; training seed 42). Each point represents an evaluated checkpoint, not a "
         "mean across seeds. Selected checkpoints can precede the run endpoint, notably M5 at "
         "25k. D-only has the largest macro success, while M7-Raw has marginally larger micro "
         "success. M2 and M3 lack formal full-validation results. The dot-plot axis is restricted "
         "to the observed score range; no bar lengths encode these truncated values.")


def plot_budget(budget):
    fig, axs = plt.subplots(1,2,figsize=(WIDTH_MM/25.4,85/25.4))
    fig.subplots_adjust(left=.10,right=.98,bottom=.21,top=.80,wspace=.35)
    for name,color,marker in [("D-only",BLUE,"s"),("M7-Raw",VERMILION,"o")]:
        rows = sorted([r for r in budget if r["method"]==name],key=lambda r:r["iteration"])
        x=[r["iteration"]/1000 for r in rows]
        for ax,key in zip(axs,["macro","micro"]):
            ax.plot(x,[r[key]*100 for r in rows], marker=marker,color=color,label=name)
    for ax,title in zip(axs,["(a) Macro success", "(b) Micro success"]):
        ax.set(xlabel="Checkpoint iteration (k)", ylabel="Full-validation success (%)", xlim=(32,72), ylim=(87,93))
        ax.set_title(title, loc="left",pad=11); ax.set_xticks([34,40,50,59,70]); clean(ax,"y")
        ax.axvline(59, color=GRAY, linestyle=":",linewidth=.8)
    fig.legend(*axs[0].get_legend_handles_labels(),loc="upper center",ncol=2,frameon=False,bbox_to_anchor=(.53,1))
    save(fig,"Fig_05_Validation_Budget_Scaling",
         "Fig. 5. Full-validation (a) macro and (b) micro success of evaluated D-only and M7-Raw "
         "checkpoints. Lines connect available evaluations only and do not represent continuous "
         "training measurements. The vertical reference marks the selected 59k M7-Raw checkpoint. "
         "D-only was explored to 59k but its available selected full-validation checkpoint is 54k; "
         "M7-Raw was explored to 70k. Missing D-only points are not interpolated as measured results.")


def plot_test(test):
    fig, axs = plt.subplots(2,2,figsize=(WIDTH_MM/25.4,127/25.4))
    fig.subplots_adjust(left=.16,right=.98,bottom=.12,top=.86,wspace=.43,hspace=.62)
    colors=[GRAY,BLUE,VERMILION]; markers=["D","s","o"]
    names=["GlobalRaw 33.5k","D-only 54k","M7-Raw 59k"]
    for i,r in enumerate(test):
        axs[0,0].scatter([100*r[k] for k in ["micro","macro","completion"]], np.arange(3)+(i-1)*.16,
                         color=colors[i],marker=markers[i],s=30,label=names[i])
    axs[0,0].set(yticks=[0,1,2],yticklabels=["Micro","Macro","Completion"],xlim=(84,97),ylim=(2.6,-.6),xlabel="Rate (%)")
    for ax,key,scale,label,limit in [(axs[0,1],"body",1000,"Body error (mm)",(46,53)),
                                    (axs[1,0],"joint",1,"Joint L2 error (rad)",(.72,.90)),
                                    (axs[1,1],"failures",1,"Failed motions",(0,1200))]:
        for i,r in enumerate(test):
            value=r[key]*scale
            ax.scatter(value,i,color=colors[i],marker=markers[i],s=32)
            ax.annotate(f"{value:.2f}" if key=="body" else (f"{value:.4f}" if key=="joint" else str(value)),
                        (value,i),xytext=(5,7),textcoords="offset points",fontsize=8.5)
        ax.set(yticks=[0,1,2],yticklabels=names,xlim=limit,ylim=(2.6,-.6),xlabel=label)
    for ax,title in zip(axs.flat,["(a) Success and completion","(b) Body-position error","(c) Joint-position error","(d) Failure count"]):
        ax.set_title(title,loc="left",pad=12); clean(ax)
    fig.legend(*axs[0,0].get_legend_handles_labels(),loc="upper center",ncol=3,frameon=False,bbox_to_anchor=(.5,.99))
    save(fig,"Fig_06_Final_Test",
         "Fig. 6. Final Test comparison on 7,592 motions using validation-selected checkpoints. "
         "(a) Micro success, macro success, and completion are higher-is-better. (b) Body error, "
         "(c) joint L2 error, and (d) failure count are lower-is-better. Error units are millimetres "
         "and radians, respectively. Checkpoints are GlobalRaw 33.5k, D-only 54k, and M7-Raw 59k; "
         "their maximum explored training budgets are 34k, 59k, and 70k. Thus, this is a selected-policy "
         "comparison, not an equal-budget module ablation. Points are single fitted policies; paired "
         "uncertainty is reported in Fig. 7.")


def plot_paired(analysis):
    fig,axs=plt.subplots(1,2,figsize=(WIDTH_MM/25.4,91/25.4))
    fig.subplots_adjust(left=.12,right=.98,bottom=.28,top=.83,wspace=.53)
    specs=[("bootstrap_m7_vs_donly",VERMILION,"o","vs D-only 54k"),
           ("bootstrap_m7_vs_globalraw",BLUE,"s","vs GlobalRaw 33.5k")]
    for index,(key,color,marker,name) in enumerate(specs):
        for ax,metrics,scale in [(axs[0],["micro","macro","completion"],100),(axs[1],["joint"],-1)]:
            for j,metric in enumerate(metrics):
                r=analysis[key][metric]; v=r["estimate"]*scale
                lo,hi=sorted([r["ci_low"]*scale,r["ci_high"]*scale])
                ax.errorbar(v,j+(index-.5)*.22,xerr=[[v-lo],[hi-v]],color=color,fmt=marker,capsize=3,label=name if j==0 else None)
    axs[0].set(yticks=[0,1,2],yticklabels=["Micro","Macro","Completion"],ylim=(2.6,-.6),xlim=(-.2,7.2),xlabel="Rate increase (percentage points)")
    axs[1].set(yticks=[0],yticklabels=["Joint L2"],ylim=(.6,-.6),xlim=(-.002,.085),xticks=[0,.02,.04,.06,.08],xlabel="Joint-error reduction (rad)")
    for ax,title in zip(axs,["(a) Paired rate differences","(b) Paired error reduction"]):
        ax.axvline(0,color=BLACK,linestyle=":",lw=.8); clean(ax)
        ax.set_title(title,loc="left",pad=12)
    fig.legend(*axs[0].get_legend_handles_labels(),loc="lower center",ncol=2,frameon=False,bbox_to_anchor=(.52,.005))
    save(fig,"Fig_07_Paired_Test_Intervals",
         "Fig. 7. Paired differences for the selected M7-Raw 59k policy relative to D-only 54k "
         "and GlobalRaw 33.5k. Positive values favor M7-Raw in both panels: (a) increases in "
         "success/completion and (b) reduction in joint L2 error. Whiskers are 95% percentile "
         "intervals from 10,000 paired source-group bootstrap resamples (3,610 groups; random "
         "seed 42). Macro success uses within-category group resampling. Intervals quantify "
         "evaluation-set variation conditional on selected policies, not training-seed variance "
         "or checkpoint-selection uncertainty. Training budgets are unequal.")


def plot_categories(categories):
    rows=sorted(categories,key=lambda r:r["delta_d"])
    fig,axs=plt.subplots(1,2,figsize=(WIDTH_MM/25.4,147/25.4),sharey=True)
    fig.subplots_adjust(left=.24,right=.98,bottom=.13,top=.9,wspace=.29)
    for i,r in enumerate(rows):
        color=VERMILION if r["delta_d"]>0 else (BLUE if r["delta_d"]<0 else GRAY)
        axs[0].hlines(i,0,100*r["delta_d"],color=color,lw=1.1)
        axs[0].scatter(100*r["delta_d"],i,color=color,s=28)
        gain=r["d_failures"]-r["m_failures"]
        axs[1].hlines(i,0,gain,color=color,lw=1.1); axs[1].scatter(gain,i,color=color,s=28)
        axs[1].annotate(f"{gain:+d}",(gain,i),xytext=(6,0),textcoords="offset points",va="center",fontsize=8.5)
    axs[0].set(yticks=range(len(rows)),yticklabels=[f"{r['category']} (n={r['n']})" for r in rows],
               xlabel="Success difference (percentage points)",xlim=(-2,13),ylim=(-.8,len(rows)-.2))
    axs[1].set(xlabel="Net fewer failures",xlim=(-15,146))
    for ax,title in zip(axs,["(a) M7-Raw minus D-only","(b) Contribution to net gain"]):
        ax.axvline(0,color=BLACK,linestyle=":",lw=.8); clean(ax); ax.set_title(title,loc="left",pad=14)
    save(fig,"Fig_08_Test_Category_Analysis",
         "Fig. 8. Category-wise Test comparison of M7-Raw 59k and D-only 54k. (a) Difference in "
         "micro success within each category; (b) net reduction in failed motions in that category. "
         "Positive values favor M7-Raw. Category sizes appear next to the labels. Fitness accounts "
         "for 119 of the overall 147 fewer failures. There are 10 improved categories, four ties, "
         "and three regressions. Categories are sorted by the observed rate difference; tiny "
         "categories have coarse success-rate resolution. These are descriptive point differences, "
         "not category-wise significance tests.")


def plot_probe(rows):
    fig,ax=plt.subplots(figsize=(WIDTH_MM/25.4,81/25.4))
    fig.subplots_adjust(left=.1,right=.98,bottom=.21,top=.86)
    for key,color,marker,label in [("macro",VERMILION,"o","Macro success"),("micro",BLUE,"s","Micro success"),("completion",GREEN,"^","Completion")]:
        ax.plot([r["iteration"]/1000 for r in rows],[r[key]*100 for r in rows],color=color,marker=marker,label=label,lw=1,ms=3)
    ax.set(xlim=(49,71),ylim=(-3,103),xticks=[50,55,59,60,65,70],xlabel="Checkpoint iteration (k)",ylabel="Validation probe rate (%)")
    ax.axvline(59,color=GRAY,linestyle=":",lw=.8); ax.legend(loc="lower center",bbox_to_anchor=(.5,1.01),ncol=3,frameon=False)
    clean(ax,"y")
    save(fig,"Fig_S1_Dense_Validation_Probe",
         "Fig. S1. Available dense checkpoint evaluations on the fixed 500-motion validation "
         "probe. Lines connect measured points without smoothing. Near-zero-success checkpoints "
         "have very short execution, so their conditional tracking errors are not evidence of "
         "good tracking. The 65k zero-success result is repeated in coarse and dense sweeps with "
         "the same recorded checkpoint hash and evaluation settings. These observations establish "
         "non-monotonic evaluated performance, not its underlying cause. Probe scores are not "
         "substitutes for full-validation or final Test scores.")


def graphical_abstract(test):
    fig,ax=diagram(76)
    ax.text(2,89,"Training references",weight="bold",fontsize=13)
    ax.text(2,73,"6,000 motions\nQuality + diversity",fontsize=11,linespacing=1.5)
    for i,color in enumerate([GREEN,BLUE,GRAY]):
        ax.plot([3,26],[50-i*8]*2,color=color,lw=2)
        ax.plot([7+i*6]*2,[47-i*8,53-i*8],color=color,lw=2)
    arrow(ax,(28,56),(36,56))
    ax.text(39,89,"Hierarchical sampling",weight="bold",fontsize=13)
    ax.text(39,70,"Cluster",color=BLUE,fontsize=12)
    ax.text(45,52,"Motion",color=VERMILION,fontsize=12)
    ax.text(51,34,"Segment",color=VERMILION,fontsize=12)
    arrow(ax,(42,65),(47,58)); arrow(ax,(48,47),(54,40))
    arrow(ax,(65,56),(72,56))
    ax.text(74,89,"Final Test",weight="bold",fontsize=13)
    ax.text(74,64,"92.31%",color=GREEN,weight="bold",fontsize=24)
    ax.text(74,48,"micro success",fontsize=11)
    ax.text(74,28,"7,008 / 7,592",fontsize=11)
    ax.text(50,9,"QD-HES / M7-Raw: validation-selected 59k checkpoint",ha="center",fontsize=10)
    save(fig,"Graphical_Abstract",
         "Graphical abstract. Training-reference knowledge and online errors guide hierarchical "
         "episode-start sampling. The selected M7-Raw 59k policy succeeds on 7,008 of 7,592 Test "
         "motions. Artwork is a reproducible vector schematic and numerical summary, not an "
         "AI-generated robot image." + AI_CAPTION)


def audit_evidence(formal,budget,test):
    metadata=data.read_csv_rows(ROOT/"PHUMA_wbt_motions/manifests/splits_v1/metadata.csv")
    lookup={r["relative_path"]:r for r in metadata}
    groups={}; paths={}; split_root=ROOT/"PHUMA_wbt_motions/manifests"
    manifests={"training":split_root/"experiments/random_seed42/random6000_seed42.txt",
               "validation":split_root/"splits_v1/validation_full.txt", "test":split_root/"splits_v1/test.txt",
               "probe":split_root/"splits_v1/validation_probe500_seed42.txt"}
    for name,path in manifests.items():
        items=[s.strip() for s in path.read_text().splitlines() if s.strip() and not s.startswith("#")]
        normalized=[str(Path(p).relative_to(ROOT)) if Path(p).is_absolute() else p for p in items]
        assert len(normalized)==len(set(normalized)),name
        paths[name]=set(normalized); groups[name]={lookup[p]["source_group"] for p in normalized}
    for a,b in [("training","validation"),("training","test"),("validation","test")]:
        assert not paths[a]&paths[b],(a,b,"path leakage")
        assert not groups[a]&groups[b],(a,b,"source-group leakage")
    assert paths["probe"]<=paths["validation"]
    configs=[]; input_files=set(manifests.values())
    for r in formal+budget+test:
        directory=ROOT/r["directory"]
        for file in ("summary.json","per_motion.csv","evaluation_config.json"):
            path=directory/file
            if path.exists(): input_files.add(path)
    for r in test:
        directory=ROOT/r["directory"]; cfg=data.read_json(directory/"evaluation_config.json")
        rows=data.read_csv_rows(directory/"per_motion.csv")
        by_path={row["motion_path"]:row for row in rows}
        assert len(by_path)==len(rows)==7592
        assert set(by_path)==paths["test"]
        for path,row in by_path.items():
            assert row["source_group"]==lookup[path]["source_group"]
            assert row["category"]==lookup[path]["category"]
        assert data.sha256_file(Path(cfg["checkpoint_path"]))==cfg["checkpoint_sha256"]
        cfg["method"]=r["method"]; configs.append(cfg)
        dest=EVIDENCE/r["method"].split(" ")[0]; dest.mkdir(exist_ok=True)
        for file in ("per_motion.csv","summary.json","evaluation_config.json","category_summary.csv","source_group_summary.csv"):
            shutil.copy2(directory/file,dest/file)
    for key in ("manifest_sha256","git_commit","num_envs","episode_length_s","deterministic","disable_randomization","seed"):
        assert len({str(c[key]) for c in configs})==1,key
    base=ROOT/"evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1"
    coarse=base/"validation_probe500_sweep/model_65000"; dense=base/"validation_probe500_dense_sweep_60500_65000/model_65000"
    c1=data.read_json(coarse/"evaluation_config.json"); c2=data.read_json(dense/"evaluation_config.json")
    check_keys=["checkpoint_sha256","manifest_sha256","git_commit","task","seed","num_envs","episode_length_s","deterministic","disable_randomization"]
    comparison={key:c1[key]==c2[key] for key in check_keys}
    comparison["identical_per_motion_csv"]=data.sha256_file(coarse/"per_motion.csv")==data.sha256_file(dense/"per_motion.csv")
    for d in (coarse,dense):
        input_files.update(d/file for file in ("evaluation_config.json","summary.json","per_motion.csv"))
    source_groups_by_category={}
    for r in data.read_csv_rows(ROOT/test[-1]["directory"]/"per_motion.csv"):
        source_groups_by_category.setdefault(r["source_group"],set()).add(r["category"])
    assert all(len(c)==1 for c in source_groups_by_category.values())
    audit={"date":DATE,"split_counts":{name:{"motions":len(paths[name]),"source_groups":len(groups[name]),"sha256":data.sha256_file(manifests[name])} for name in paths},
           "cross_split_source_group_overlap":0,"test_motion_alignment":"unique motion_path, not row order",
           "checkpoint_hashes_verified":True,"test_protocol_equal":True,"repeat_65000":comparison,
           "repeat_65000_micro":[data.read_json(d/"summary.json")["micro_success_rate"] for d in (coarse,dense)],
           "training_seeds_available":[42],"scope":"retrospective artifact audit, not preregistration"}
    dump(EVIDENCE/"Evidence_Audit.json",audit)
    dump(EVIDENCE/"Final_Test_Configurations.json",configs)
    data.write_csv(EVIDENCE/"Input_Provenance.csv",["Path","SHA256"],[(str(p.relative_to(ROOT)),data.sha256_file(p)) for p in sorted(input_files)])
    return audit


def doc_style(doc,landscape=False):
    sec=doc.sections[0]
    sec.page_width=Cm(29.7 if landscape else 21); sec.page_height=Cm(21 if landscape else 29.7)
    if landscape: sec.orientation=WD_ORIENT.LANDSCAPE
    sec.left_margin=sec.right_margin=Cm(2)
    sec.top_margin=sec.bottom_margin=Cm(1.8)
    if landscape:
        sec.bottom_margin=Cm(2.6)
    normal=doc.styles["Normal"]; normal.font.name="Liberation Serif"; normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
    for name,size in [("Title",20),("Heading 1",14),("Heading 2",12),("Heading 3",11)]:
        st=doc.styles[name]; st.font.name="Liberation Sans"; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
        st.paragraph_format.keep_with_next=True
    foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=foot.add_run("PHUMA-WBT | "); run.font.size=Pt(9)
    field=OxmlElement("w:fldSimple"); field.set(qn("w:instr"),"PAGE"); foot._p.append(field)


def add_table(doc,headers,rows,widths=None,best_cols=None):
    table=doc.add_table(rows=1,cols=len(headers)); table.autofit=False
    width=25.0 if doc.sections[-1].orientation==WD_ORIENT.LANDSCAPE else 17
    widths=widths or [width/len(headers)]*len(headers)
    for col,w in zip(table.columns,widths): col.width=Cm(w)
    for i,value in enumerate(headers): table.rows[0].cells[i].text=str(value)
    for row in rows:
        cells=table.add_row().cells
        for i,value in enumerate(row): cells[i].text=str(value)
    repeat=OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(repeat)
    winners={}
    for col,direction in (best_cols or {}).items():
        values=[float(str(r[col]).replace(",","")) for r in rows]
        winners[col]=(max if direction=="max" else min)(values)
    for ri,row in enumerate(table.rows):
        cant_split=OxmlElement("w:cantSplit"); row._tr.get_or_add_trPr().append(cant_split)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for edge in ("top","bottom","left","right"):
                tag=OxmlElement("w:"+edge)
                on=(ri==0 and edge in ("top","bottom")) or (ri==len(rows) and edge=="bottom")
                tag.set(qn("w:val"),"single" if on else "nil"); tag.set(qn("w:sz"),"8" if ri in (0,len(rows)) else "4")
                borders.append(tag)
            tcPr.append(borders)
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(4); p.paragraph_format.space_before=Pt(4)
                p.paragraph_format.line_spacing=1.0
                for run in p.runs:
                    run.font.name="Liberation Serif"; run.font.size=Pt(9 if len(headers)>5 else 10)
                    run.bold=ri==0 or (ri>0 and ci in winners and float(str(rows[ri-1][ci]).replace(",",""))==winners[ci])
    return table


def latex_escape(value):
    mapping={"\\":r"\textbackslash{}","&":r"\&","%":r"\%","$":r"\$","#":r"\#","_":r"\_",
             "{":r"\{","}":r"\}","↑":r"$\uparrow$","↓":r"$\downarrow$"}
    return "".join(mapping.get(ch,ch) for ch in str(value))


def prepare_tables(tables,formal,budget,test):
    # Long paths and hashes belong in the provenance files, not narrow printed cells.
    h,r=tables["Table_03_Data_Splits"]
    tables["Table_03_Data_Splits"]=(["Split","Manifest","Motions"],[row[:3] for row in r])
    for key in ("Table_08_Validation_Ablation_34k","Table_09_Validation_Budget_Scaling"):
        h,r=tables[key]
        tables[key]=(["Method","Run/selection","Iteration","Micro ↑","Macro ↑","Completion ↑","Body (m) ↓","Joint L2 (rad) ↓","Failures ↓"],
                     [[row[0],row[1],row[2].removeprefix("model_").removesuffix(".pt"),*row[3:]] for row in r])
    tables["Table_10_Final_Test"]=(
        ["Method","Selected / explored (k)","Micro ↑","Macro ↑","Completion ↑","Body (m) ↓","Joint L2 (rad) ↓","Failures ↓"],
        [[r["method"].replace(" (QD-HES)",""),b,f"{r['micro']:.6f}",f"{r['macro']:.6f}",f"{r['completion']:.6f}",
          f"{r['body']:.6f}",f"{r['joint']:.6f}",r["failures"]] for r,b in zip(test,["33.5 / 34","54 / 59","59 / 70"])])
    quality=data.read_json(ROOT/"configs/quality/g1_segment_quality.yaml")
    metrics=quality["metrics"]
    tables["Table_S4_Quality_Thresholds"]=(
        ["Metric","Unit","Weight","Warning","Reject","Hard at reject"],
        [[name.replace("_"," "),cfg.get("unit","").replace("_"," "),cfg.get("weight",""),cfg.get("warning_threshold",""),cfg.get("reject_threshold",""),str(cfg.get("hard_at_reject",False))] for name,cfg in metrics.items()])
    diff=data.read_json(ROOT/"configs/difficulty/g1_segment_difficulty.yaml")
    features=diff.get("features",{})
    if isinstance(features,dict):
        rows=[[name.replace("_"," "),cfg.get("weight",""),cfg.get("description","")] for name,cfg in features.items()]
    else: rows=[[cfg.get("name",""),cfg.get("weight",""),cfg.get("description","")] for cfg in features]
    tables["Table_S5_Difficulty_Features"] = (["Feature","Weight","Definition"],rows)
    tables["Table_S6_Online_Sampling_Parameters"]=(
        ["Parameter","Value / interpretation"],
        [["EMA decay",0.95],["Warm-up / refresh", "1000 / 50 completed iterations"],
         ["Reliability","32 segment observations; 8 motion outcomes"],
         ["Error scales","Body 0.30 m; joint L2 0.50 rad; orientation 0.40 rad"],
         ["Segment weights","Body/joint/orientation/termination/completion/success = 1/1/1/1/0.5/0.5"],
         ["Motion weights","Mean/P90/termination/completion/success = 1/0.25/0.5/0.5/0.5"],
         ["Exploration","Uniform mixture 0.15; count bonus coefficient 0.25"],
         ["Softmax / score clip","Temperature 1.0; [-10,10]"],
         ["Caps","Motion conditional 0.02, relaxed to >=1/n_c; segment conditional 1.0"],
         ["Diversity","K=8; f=0.5; alpha=0.5; active during warm-up"],
         ["Quality","Assignment-start gate; reject excluded; borderline included"],
         ["Final method","Raw/raw + quality + diversity; no learning-gap calibration"]])
    tables["Table_S7_Paired_Test_Outcomes"]=(
        ["Comparison","Both succeed","M7-Raw only","Baseline only","Neither succeeds"],
        [["M7-Raw 59k vs D-only 54k",6711,297,150,434],
         ["M7-Raw 59k vs GlobalRaw 33.5k",6502,506,93,491]])
    tables["Table_01_Module_Design"][1][0][2]="29-joint, 30-body NPZ schema; 14 configured tracking bodies; independent motion/start per environment."
    tables["Table_02_Module_Statistics"][1][1][0]="Training subset"
    return tables


TABLE_NOTES={
    "Table_03_Data_Splits":"SHA256 identities and source-group disjointness are checked separately in evidence/Evidence_Audit.json. Probe is a Validation subset.",
    "Table_08_Validation_Ablation_34k":"All runs have the stated 34k budget; selected checkpoints may be earlier. Full Validation: 7,636 motions. One training seed. Bold indicates the best value in each metric column, not statistical significance. M2/M3 formal results are unavailable.",
    "Table_09_Validation_Budget_Scaling":"Rows are available full-validation checkpoints, not all saved checkpoints. D-only selected at 54k comes from a 59k run; M7-Raw selected at 59k comes from a 70k run. The 34k and 50k endpoint comparisons have different scope from selected-policy comparisons.",
    "Table_10_Final_Test":"Test: 7,592 motions. Selected / explored separates the chosen checkpoint from the maximum training budget searched. All three policies use training seed 42; these rows are not equal-budget ablations. Bold marks best values among tested policies.",
    "Table_11_Final_Test_Paired_Deltas":"Differences are M7-Raw minus baseline; errors favor negative differences. Rate units are percentage points. 95% group-bootstrap intervals condition on selected policies and do not represent multi-seed variance.",
    "Table_S3_Dense_Probe":"Probe500 only; do not mix these scores with Full Validation. Near-zero completion can yield deceptively small errors.",
    "Table_S4_Quality_Thresholds":"Units and metric-specific composite rules are defined in the included quality profile. Thresholds alone do not reproduce the complete decision rule. The profile retains its original provisional flag.",
    "Table_S5_Difficulty_Features":"The complete configuration snapshot defines normalization and aggregation. Difficulty calibration is inactive in M7-Raw; selected raw features are reused for clustering.",
}


def table_best(name):
    if name in ("Table_08_Validation_Ablation_34k","Table_09_Validation_Budget_Scaling"):
        return {3:"max",4:"max",5:"max",6:"min",7:"min",8:"min"}
    if name=="Table_10_Final_Test": return {2:"max",3:"max",4:"max",5:"min",6:"min",7:"min"}
    return {}


def export_tables(tables):
    doc=Document(); doc_style(doc,True)
    tex=[]; md=[]
    for index,(name,(headers,rows)) in enumerate(tables.items()):
        if index: doc.add_page_break()
        title=name.replace("_"," ")
        doc.add_heading(title,1)
        widths={"Table_S4_Quality_Thresholds":[6,7,2,3,3,4],
                "Table_S5_Difficulty_Features":[7.5,2,15.5],
                "Table_S6_Online_Sampling_Parameters":[6,19]}.get(name)
        add_table(doc,headers,rows,widths=widths,best_cols=table_best(name))
        note=TABLE_NOTES.get(name,"")
        if note: doc.add_paragraph(note)
        data.write_csv(TABLE/f"{name}.csv",headers,rows)
        block=[r"\begin{table*}[t]",r"\centering",r"\caption{"+latex_escape(title)+"}",
               r"\label{tab:"+name.lower()+"}",r"\small",r"\setlength{\tabcolsep}{3pt}",
               r"\begin{tabularx}{\textwidth}{"+" ".join("X" for _ in headers)+"}",r"\toprule",
               " & ".join(latex_escape(s) for s in headers)+r" \\",r"\midrule"]
        winners={c:(max if d=="max" else min)(float(row[c]) for row in rows) for c,d in table_best(name).items()}
        for row in rows:
            vals=[]
            for c,v in enumerate(row):
                cell=latex_escape(v)
                if c in winners and float(v)==winners[c]: cell=r"\textbf{"+cell+"}"
                vals.append(cell)
            block.append(" & ".join(vals)+r" \\")
        if len(rows)>20:
            body=block[10:]
            fraction=1/len(headers)
            columns=" ".join(r"p{\dimexpr "+f"{fraction:.6f}"+r"\textwidth-2\tabcolsep\relax}" for _ in headers)
            header = " & ".join(latex_escape(s) for s in headers) + " " + chr(92) * 2
            block=["% Use in a one-column supplementary section; multipage table.",
                   r"\begingroup\small\setlength{\tabcolsep}{3pt}",
                   r"\begin{longtable}{"+columns+"}",
                   r"\caption{"+latex_escape(title)+r"}\label{tab:"+name.lower()+r"}\\",
                   r"\toprule",header,r"\midrule\endfirsthead",r"\toprule",header,
                   r"\midrule\endhead",r"\bottomrule\endfoot",*body,r"\end{longtable}",r"\endgroup"]
        else:
            block.extend([r"\bottomrule",r"\end{tabularx}"])
        if note: block.extend([r"\par\smallskip\begin{minipage}{\textwidth}\footnotesize",latex_escape(note),r"\end{minipage}"])
        if len(rows)<=20: block.append(r"\end{table*}")
        content="\n".join(block)+"\n"; (TABLE/f"{name}.tex").write_text(content,encoding="utf-8"); tex.append(content)
        md.extend(["## "+title,data.markdown_table(headers,rows),note])
    doc.save(OUT/"Editable_Paper_Tables.docx")
    (TABLE/"All_Tables.tex").write_text("% Requires booktabs, tabularx, longtable, amsmath. Long tables need a one-column supplement.\n\n"+"\n".join(tex),encoding="utf-8")
    (TABLE/"All_Tables.md").write_text("\n\n".join(md),encoding="utf-8")


def write_manuscript(tables):
    doc=Document(); doc_style(doc)
    doc.add_heading(prose.TITLE,0)
    doc.add_paragraph("English manuscript materials | Knowledge-Based Systems | " + DATE)
    doc.add_paragraph("Working text with traceable experimental evidence. Author names, affiliations, funding, final declarations, and an expanded related-work review remain for the authors to complete.")
    doc.add_heading("Abstract",1); doc.add_paragraph(prose.ABSTRACT)
    doc.add_paragraph("Keywords: humanoid motion tracking; reinforcement learning; adaptive sampling; trajectory quality; motion diversity.")
    md=["# "+prose.TITLE,"English manuscript materials | KBS | "+DATE,"## Abstract",prose.ABSTRACT]
    for title,paras in prose.SECTIONS:
        doc.add_heading(title,1 if ". " in title[:4] else 2); md.append("## "+title)
        for paragraph in paras: doc.add_paragraph(paragraph); md.append(paragraph)
    doc.add_heading("References",1); md.append("## References")
    for r in prose.REFERENCES: doc.add_paragraph(r); md.append(r)
    doc.add_heading("Reproducible visualization and AI assistance",1)
    note=("Numerical figures are generated directly from the archived CSV/JSON results using Python, NumPy, and Matplotlib. "
          "OpenAI Codex (GPT-5, OpenAI) assisted in organizing the text and preparing analytical and layout code. "
          "No primary robot image or experimental outcome was synthesized. Diagram captions disclose layout assistance. "
          "Authors must review the generated text, code, figures, and tool-version statement before submission and take responsibility for the final content.")
    doc.add_paragraph(note); md.extend(["## Reproducible visualization and AI assistance",note])
    doc.save(OUT/"PHUMA_WBT_KBS_English_Materials.docx")
    (OUT/"PHUMA_WBT_KBS_English_Materials.md").write_text("\n\n".join(md)+"\n",encoding="utf-8")
    highlights=Document(); doc_style(highlights); highlights.add_heading("Highlights",0)
    for h in prose.HIGHLIGHTS:
        assert len(h)<=85; highlights.add_paragraph(h,style="List Bullet")
    highlights.save(OUT/"Highlights.docx")
    (OUT/"Highlights.txt").write_text("\n".join(prose.HIGHLIGHTS)+"\n",encoding="utf-8")


def write_formulae():
    content=r"""% QD-HES method equations. Requires amsmath and amssymb.
\begin{align}
L_m &= \max(1,\operatorname{round}(f_m\Delta)),\quad \Delta=1\,\mathrm{s},
& S_m &= \lceil T_m/L_m\rceil,\\
q_{ms} &= \mathbf{1}[\text{legal start in }(m,s)]\,
          \mathbf{1}[\operatorname{status}_{ms}\ne\mathrm{reject}],\\
\bar x_{ms}^{(t)} &= \rho\bar x_{ms}^{(t-1)}+(1-\rho)x_{ms}^{(t)},
&\rho&=0.95,\\
\mathbf z_{ms} &= \left[\min(\bar e_b/0.30,5),\min(\bar e_j/0.50,5),
\min(\bar e_o/0.40,5),\bar d,1-\bar c,1-\bar y\right],\\
E_{ms} &= \frac{\sum_{k\in A_{ms}} w_k z_{ms,k}}{\sum_{k\in A_{ms}}w_k},
&\mathbf w&=(1,1,1,1,0.5,0.5),\\
P(c) &= \frac{f}{C}+(1-f)\frac{n_c^\alpha}{\sum_{j=1}^C n_j^\alpha},
& (f,\alpha)&=(0.5,0.5),\\
\ell_i &= \frac{\operatorname{clip}(E_i,-10,10)+0.25/\sqrt{N_i+1}}{\tau},
& \tau&=1,\\
\widetilde P_i &= \frac{0.15}{|\mathcal E|}+0.85\frac{e^{\ell_i}}{\sum_{j\in\mathcal E}e^{\ell_j}},
&i&\in\mathcal E,\\
P(c,m,s)&=P(c)P(m\mid c)P(s\mid m).
\end{align}
% EMA is initialized from the first valid observation. A_ms includes observed,
% uncensored components only. P-tilde is followed by proportional water-filling to
% enforce feasible caps; ineligible probabilities are exactly zero. In a cluster,
% the conditional motion cap is max(0.02,1/n_c). Cold items use the code's fallback.

\begin{align}
Q_{ms} &= \operatorname{clip}\left(1-\frac{\sum_k w^Q_k v_{ms,k}}{\sum_k w^Q_k},0,1\right),\\
\overline E_m &= \frac{\sum_{s\in V_m} N_{ms}E_{ms}}{\sum_{s\in V_m}N_{ms}},\\
E_m &= \operatorname{AWMean}\left(
 [\overline E_m,\operatorname{P90}_{s\in V_m}E_{ms},\bar d_m,1-\bar c_m,1-\bar y_m],
 [1,0.25,0.5,0.5,0.5]\right).
\end{align}
% AWMean normalizes over available components only. V_m denotes reliable segments.
% Q severity weights are not the tracking-error weights. Quality status additionally
% applies hard violations, required metric coverage and severe-metric-count rules.

% Optional difficulty/learning-gap branch: NOT part of the selected M7-Raw method.
\begin{align}
G_{ms} &= \operatorname{clip}\left(
 \frac{E_{ms}-\mu_{b(ms)}}{\max(\sigma_{b(ms)},0.10)},-5,5\right),\\
G^{\mathrm{local}}_{ms} &= \operatorname{clip}\left(
 G_{ms}-\operatorname{Median}_{j\in V_m}G_{mj},-5,5\right),\\
G_m &= \operatorname{clip}\left(\operatorname{AWMean}\left(
 [\operatorname{WMean}_s G^+_{ms},\operatorname{P90}_s G^+_{ms},
 \bar d_m,1-\bar c_m,1-\bar y_m],[1,0.5,0.5,0.5,0.5]\right),0,5\right).
\end{align}
% Bin moments use reliable Train segments only; sparse bins use the documented fallback.

\begin{align}
\mathrm{Micro} &= N^{-1}\sum_i y_i,\\
\mathrm{Macro} &= K^{-1}\sum_{k=1}^K N_k^{-1}\sum_{i\in\mathcal C_k}y_i,\\
\mathrm{Completion} &= N^{-1}\sum_i \min(\widehat T_i/T_i,1),\\
\mathrm{JointL2} &= N^{-1}\sum_i H_i^{-1}\sum_{t=1}^{H_i}
                 \|\mathbf q_{it}-\mathbf q^{\mathrm{ref}}_{it}\|_2,\\
\mathrm{JointRMS} &= \mathrm{JointL2}/\sqrt{29},\\
\mathrm{Failures} &= N-\sum_i y_i.
\end{align}
% K=17 source categories; H_i is the number of evaluated steps actually observed.
% Body error uses the code's 14 configured tracked bodies and root/yaw alignment.
\begin{equation}
\mathrm{BodyErr}=\frac1N\sum_{i=1}^N\frac1{H_i}\sum_{t=1}^{H_i}
 \frac1{14}\sum_{b=1}^{14}\|\widetilde{\mathbf p}^{\mathrm{ref}}_{itb}
                              -\mathbf p^{\mathrm{robot}}_{itb}\|_2.
\end{equation}
% The transformed reference uses robot anchor xy, reference anchor z, and yaw
% from robot-anchor quaternion times inverse reference-anchor quaternion.
"""
    (OUT/"Method_and_Metric_Equations.tex").write_text(content,encoding="utf-8")
    algorithm=[
        "Algorithm 1. QD-HES training loop",
        "Input: Train manifest; frozen quality and cluster metadata; policy and PPO state.",
        "1. Build stable motion/segment indices and legal assignment-start masks.",
        "2. Exclude rejected starts and empty motions; compute eligible cluster budgets.",
        "3. Initialize or restore online EMAs, sample counters, completed-window counter and sampler RNG.",
        "4. For each PPO iteration:",
        "   a. At an episode assignment, sample cluster c from P(c).",
        "   b. Sample eligible motion m from P(m|c), segment s from P(s|m), and a legal frame uniformly within s.",
        "   c. Collect on-policy rollout data and attribute observations/outcomes to visited segments.",
        "   d. Commit initialized EMA observations; handle administrative truncations as censored outcomes.",
        "   e. Optimize the unchanged PPO objective using the collected rollout.",
        "   f. During 1000-iteration warm-up, use uniform conditionals with diversity budgets active.",
        "   g. Thereafter, every 50 completed iterations, recompute reliable errors, exploration, softmax and feasible caps.",
        "   h. Checkpoint policy, optimizer and adaptive sampler state at the configured interval.",
        "Output: checkpoint candidates for Validation-only selection; Test is outside the training loop.",
    ]
    (OUT/"Algorithm_QD_HES.txt").write_text("\n".join(algorithm)+"\n",encoding="utf-8")
    doc=Document(); doc_style(doc); doc.add_heading(algorithm[0],1)
    for line in algorithm[1:]: doc.add_paragraph(line)
    doc.save(OUT/"Algorithm_QD_HES.docx")


def write_source_index():
    names={
        "Reference conversion":"scripts/phuma_to_npz.py",
        "Multi-motion loading and integration":"source/whole_body_tracking/whole_body_tracking/tasks/tracking/mdp/commands.py",
        "Stage-0 sampling":"source/whole_body_tracking/whole_body_tracking/utils/sampling.py",
        "Quality audit":"source/whole_body_tracking/whole_body_tracking/utils/quality.py",
        "Quality metadata":"source/whole_body_tracking/whole_body_tracking/utils/quality_metadata.py",
        "Intrinsic difficulty":"source/whole_body_tracking/whole_body_tracking/utils/difficulty.py",
        "Online statistics":"source/whole_body_tracking/whole_body_tracking/utils/online_learning_stats.py",
        "Online sampler integration":"source/whole_body_tracking/whole_body_tracking/utils/online_learning.py",
        "Learning gap":"source/whole_body_tracking/whole_body_tracking/utils/learning_gap.py",
        "Adaptive probability transform":"source/whole_body_tracking/whole_body_tracking/utils/adaptive_sampling.py",
        "Motion features and clustering":"source/whole_body_tracking/whole_body_tracking/utils/motion_clustering.py",
        "Diversity hierarchy":"source/whole_body_tracking/whole_body_tracking/utils/diversity_sampling.py",
        "Joint-gap diagnostics":"source/whole_body_tracking/whole_body_tracking/utils/joint_gap.py",
        "Physical termination":"source/whole_body_tracking/whole_body_tracking/tasks/tracking/mdp/terminations.py",
        "Evaluation":"scripts/rsl_rl/evaluate.py",
        "Checkpoint selection":"scripts/rsl_rl/evaluate_checkpoint_sweep.py",
    }
    rows=[]
    for title,path in names.items():
        assert (ROOT/path).exists(); rows.append([title,path,data.sha256_file(ROOT/path)])
    data.write_csv(EVIDENCE/"Implementation_Index.csv",["Module","Source path","SHA256"],rows)
    snap=EVIDENCE/"configuration_snapshots"; snap.mkdir(exist_ok=True)
    for path in ["configs/quality/g1_segment_quality.yaml","configs/difficulty/g1_segment_difficulty.yaml",
                 "configs/diversity/g1_motion_clustering.yaml","configs/online_learning/g1_module3.yaml"]:
        shutil.copy2(ROOT/path,snap/Path(path).name)
    for spec in data.TEST_SPECS:
        cfg=data.read_json(ROOT/spec[-1]/"evaluation_config.json")
        params=Path(cfg["checkpoint_path"]).parent/"params"
        for filename in ("env.yaml","agent.yaml"):
            if (params/filename).exists():
                shutil.copy2(params/filename,snap/(spec[0].split(" ")[0]+"_"+filename))
    text=["# Implementation evidence index"]
    for title,path,_ in rows: text.append(f"- {title}: `{path}`")
    text.extend(["","All snapshots describe the files inspected on 2026-09-11. Current file hashes do not prove "
                 "that the working tree was clean at historical evaluation time. Recorded evaluation commits "
                 "and checkpoint hashes are retained separately."])
    (OUT/"Implementation_Evidence_Index.md").write_text("\n".join(text),encoding="utf-8")


def write_submission_notes():
    policy=(
        "# Submission preparation notes\n\n"
        "The package is an English evidence-and-artwork kit, not an author-approved final submission. "
        "Elsevier's current general artwork guidance was checked on 2026-09-11. The KBS-specific "
        "ScienceDirect guide returned HTTP 403 in this environment; journal-specific limits, mandatory "
        "items, review format and article type must be confirmed in the submission portal.\n\n"
        "## Artwork supplied\n\n"
        "Each figure has a 190-mm-wide PDF/EPS vector master with embedded TrueType fonts; SVG is "
        "an editable source format. TIFF exports are 1000 dpi with lossless LZW compression. PNGs "
        "are 300-dpi insertion previews. The graphical abstract has a 2.5:1 ratio. Text is English, "
        "normally 8.5-10 pt at the declared size; Liberation Sans is Arial-metric-compatible. "
        "Use PDF/EPS for production, or TIFF when a raster file is requested. Do not shrink the "
        "full-width multi-panel figures to one column. The supplied dimensions follow Elsevier's "
        "general sizing examples, not a verified KBS-specific final layout.\n\n"
        "Elsevier accepts PDF/EPS for vector artwork and recommends separate numbered files and "
        "separate captions. General raster targets are 300 dpi for photographs, 500 for mixed "
        "artwork, and 1000 for line drawings. Tables here remain editable Word/CSV/LaTeX objects "
        "with horizontal rules and no vertical rules or shading.\n\n"
        "## Evidence boundaries\n\n"
        "- Do not label the final Test table an equal-budget ablation: selected/explored budgets differ.\n"
        "- Do not call 59k the global optimum; it is selected among available evaluated candidates.\n"
        "- One training seed is available. Bootstrap intervals are not training-seed confidence intervals.\n"
        "- M2/M3 full-validation results, matched external baselines and real-robot experiments are unavailable.\n"
        "- Prior Test results exist. Do not claim a never-accessed blind Test throughout development.\n"
        "- No extra Test evaluation, model reselection, training, or synthetic experimental image was performed.\n\n"
        "## Author completion\n\n"
        "Finalize author names, affiliations, CRediT roles, funding, conflicts, data/code availability "
        "and upstream dataset licenses. Expand the four core references into a critically reviewed "
        "related-work section. Multi-seed and matched-budget evidence would strengthen broad causal "
        "claims; do not invent it to fill a table. Qualitative policy figures should come from "
        "archived real simulation rollouts with checkpoint and motion provenance.\n\n"
        "## AI-assisted preparation disclosure\n\n"
        "The current Elsevier policy permits AI-assisted explanatory diagrams with caption disclosure, "
        "and data visualizations faithfully derived through reproducible computation with Methods "
        "disclosure. It prohibits fabricating primary research images. This package uses deterministic "
        "Matplotlib plots and schematics, not a general-purpose generative image model. Authors should "
        "verify the tool-version wording and finalize the supplied disclosure before submission.\n\n"
        "Draft declaration: During manuscript preparation, OpenAI Codex (GPT-5, OpenAI) assisted "
        "with English organization, analysis code and figure-layout code. [After review: state "
        "the authors' verification, revisions and responsibility for the final content.]\n\n"
        "## Official sources\n\n"
    )
    policy += "\n".join(f"- [{name}]({url})" for name,url in SOURCES.items())+"\n"
    (OUT/"KBS_Submission_Notes.md").write_text(policy,encoding="utf-8")
    (OUT/"Declarations_Template.md").write_text(
        "# Declarations to be completed by the authors\n\n"
        "- Authors and affiliations: [confirm]\n- CRediT roles: [confirm each author's actual work]\n"
        "- Funding and grant numbers: [confirm; do not assume no funding]\n"
        "- Competing interests: [confirm; do not assume none]\n"
        "- Data availability: [state actual release location and upstream restrictions]\n"
        "- Code availability: [state approved public repository/version; local paths are not a public release]\n"
        "- AI assistance: finalize the truthful statement in KBS_Submission_Notes.md.\n"
        "- Ethics and permissions: [confirm applicability and dataset/third-party licenses]\n",
        encoding="utf-8")
    (OUT/"Cover_Letter_Draft.md").write_text(
        "Dear Editors of Knowledge-Based Systems,\n\n"
        f"Please consider our manuscript, \"{prose.TITLE}\", as a research article. "
        "The study investigates knowledge-guided experience allocation for humanoid motion tracking, "
        "combining reference-quality metadata, unsupervised motion structure and online tracking "
        "statistics in a hierarchical sampler. The manuscript separates fixed-budget validation "
        "evidence from validation-selected Test comparisons and reports source-group-aware paired analysis.\n\n"
        "[Authors: explain the contribution relative to current KBS literature and confirm originality, "
        "exclusive submission, author approval and all required declarations before sending.]\n\n"
        "Sincerely,\n[Corresponding author]\n",encoding="utf-8")


def write_portfolio():
    doc=Document(); doc_style(doc)
    combined=OUT/"Figure_Portfolio.pdf"
    with PdfPages(combined) as pdf:
        for i,(name,paths) in enumerate(ART.items()):
            if i: doc.add_page_break()
            doc.add_heading(name.replace("_"," "),1)
            doc.add_picture(str(paths["png"]),width=Cm(17))
            doc.add_paragraph(CAPTIONS[name])
            fig=plt.figure(figsize=(8.2677,11.6929))
            with Image.open(paths["png"]) as im:
                aspect=im.height/im.width
                h=(190/25.4*aspect)/11.6929
                ax=fig.add_axes([.0476,.9-h,.9048,h]); ax.imshow(im); ax.axis("off")
            fig.text(.065,.96,name.replace("_"," "),fontsize=13,weight="bold",va="top")
            fig.text(.065,.86-h,textwrap.fill(CAPTIONS[name],106),fontsize=10,va="top",linespacing=1.4)
            fig.text(.065,.035,"Review portfolio. Submit the separate vector masters, not this preview document.",fontsize=8.5,color=GRAY)
            pdf.savefig(fig); plt.close(fig)
    doc.save(OUT/"Figure_Portfolio.docx")
    (OUT/"Figure_Captions.txt").write_text("\n\n".join(CAPTIONS.values())+"\n",encoding="utf-8")
    thumbs=[]
    for name,paths in ART.items():
        with Image.open(paths["png"]) as src:
            im=src.convert("RGB"); im.thumbnail((650,470))
        tile=Image.new("RGB",(690,525),"white"); tile.paste(im,((690-im.width)//2,35))
        ImageDraw.Draw(tile).text((15,10),name,fill=BLACK); thumbs.append(tile)
    sheet=Image.new("RGB",(1380,525*((len(thumbs)+1)//2)),"#eeeeee")
    for i,tile in enumerate(thumbs): sheet.paste(tile,((i%2)*690,(i//2)*525))
    sheet.save(OUT/"Figure_Contact_Sheet.png")


def write_bib():
    # Use only the four references verified against primary pages in this build.
    text=(ROOT/"reports/kbs_paper_package_2026-09-08/references_seed.bib").read_text()
    entries=[]
    for key in ("liao2025beyondmimic","lee2025phuma","chen2025gmt","schulman2017ppo"):
        start=text.index("@article{"+key); end=text.index("\n}",start)+2
        entries.append(text[start:end])
    (OUT/"references_core.bib").write_text("% Core primary references verified 2026-09-11; expand related work.\n\n"+"\n\n".join(entries)+"\n",encoding="utf-8")


def complete_probe():
    rows=data.dense_probe_rows()
    single="evaluations/postfreeze_M7Raw_extend70k_from49999_seed42_v1/validation_probe500_dense_single/model_50000"
    if (ROOT/single/"summary.json").exists() and not any(r["iteration"]==50000 for r in rows):
        r=data.metric_record("M7-Raw","50k","model_50000.pt",single)
        r["iteration"]=50000
        rows.append(r)
    return sorted(rows,key=lambda r:r["iteration"])


def write_result_inventory():
    records=[]
    for path in sorted((ROOT/"evaluations").rglob("summary.json")):
        obj=data.read_json(path)
        if not all(k in obj for k in ("num_motions","micro_success_rate","macro_success_rate")):
            continue
        records.append([str(path.parent.relative_to(ROOT)),obj.get("manifest",obj.get("manifest_path","")),
                        obj.get("checkpoint",""),obj["num_motions"],obj["micro_success_rate"],obj["macro_success_rate"],
                        obj.get("mean_completion_ratio",""),obj.get("mean_body_position_error_m",""),
                        obj.get("mean_joint_position_error_l2_rad",""),obj.get("num_failure",""),
                        data.sha256_file(path)])
    data.write_csv(EVIDENCE/"Archival_Evaluation_Inventory.csv",
                   ["Evaluation directory","Manifest","Checkpoint","Motions","Micro","Macro","Completion","Body m","Joint L2 rad","Failures","Summary SHA256"],records)
    (EVIDENCE/"Inventory_Readme.txt").write_text(
        "Archival_Evaluation_Inventory.csv enumerates evaluation summaries, including pilots, probes and older Test phases. "
        "It is an inventory, NOT a harmonized comparison table. Main tables use explicitly selected compatible "
        "result sets. Original final v2 Test CSV/JSON files are supplied in the three method subdirectories.\n",
        encoding="utf-8")


def validate_artifacts():
    font_reports=[]
    for name,paths in ART.items():
        p=paths["pdf"]
        output=subprocess.run(["pdffonts",str(p)],check=True,text=True,capture_output=True).stdout
        lines=[s.split() for s in output.splitlines()[2:] if s.strip()]
        assert lines and all("yes" in line for line in lines),name
        assert "Type 3" not in output,name
        font_reports.append(name+"\n"+output)
        with Image.open(paths["tiff"]) as im:
            assert im.mode=="RGB" and float(im.info["dpi"][0])>=999,name
            assert im.width>=7480,name
        with Image.open(paths["png"]) as im:
            pixels=np.asarray(im.convert("RGB"))
            fraction=float(np.mean(np.min(pixels,axis=2)<230))
            assert .005<fraction<.6,(name,fraction)
    (OUT/"Artwork_Font_Check.txt").write_text("\n".join(font_reports),encoding="utf-8")
    for path in TABLE.glob("Table_*.tex"):
        text=path.read_text()
        for env in ("table*","tabularx","longtable","minipage"):
            assert text.count("\\begin{"+env+"}")==text.count("\\end{"+env+"}"),(path,env)
    dump(OUT/"Build_Verification.json",{
        "date":DATE,"figure_count":len(ART),"raster_blank_check":"passed",
        "figure_canvas_and_box_text_bounds":"passed","pdf_fonts":"embedded; no Type 3",
        "tiff":"RGB, 1000 dpi","latex_environment_balance":"passed",
        "latex_compile":"not run: TeX engine unavailable; use supplied Word tables or compile sources in the author template",
        "training_or_evaluation_launched":False,
    })


def finish(tables,audit):
    for file in (Path(__file__),Path(prose.__file__),Path(data.__file__),
                 ROOT/"scripts/verify_kbs_submission_assets.py"):
        shutil.copy2(file,OUT/"source"/file.name)
    data.write_csv(OUT/"Artwork_Specifications.csv",
                   ["Figure","Width mm","Height mm","TIFF width pixels","TIFF height pixels","TIFF dpi"],
                   [[r["figure"],r["width_mm"],f"{r['height_mm']:.2f}",*r["raster"]["pixels"],r["raster"]["dpi"][0]] for r in FIGURE_AUDIT])
    dump(OUT/"Artwork_Preflight.json",FIGURE_AUDIT)
    readme=(
        "# KBS English manuscript and artwork package\n\n"
        "Prepared 2026-09-11 from local code, saved configurations, per-motion results and summaries. "
        "This supersedes the Sept8 material package's prose/figure interpretation, not its original evidence.\n\n"
        "## Main files\n\n"
        "- PHUMA_WBT_KBS_English_Materials.docx / .md: English draft material covering every module and available experiment.\n"
        "- Editable_Paper_Tables.docx: editable horizontal-rule tables, also supplied as CSV and LaTeX in tables/.\n"
        "- Figure_Portfolio.pdf / .docx: all figures and English captions for inspection.\n"
        "- figures/: 8 main figures, one supplementary probe plot and one graphical abstract; PDF/EPS/SVG/TIFF/PNG.\n"
        "- Figure_Captions.txt, Method_and_Metric_Equations.tex and Algorithm_QD_HES.docx / .txt.\n"
        "- Highlights.docx / .txt, references_core.bib, Cover_Letter_Draft.md and Declarations_Template.md.\n"
        "- KBS_Submission_Notes.md: verified general publisher rules and journal-specific checks still needed.\n"
        "- FINAL_HANDOFF.md, Final_QA.json and Final_PDF_Fonts.txt: final review, author actions and verification results.\n"
        "- evidence/: original final-Test records, hashes, configuration snapshots, paired statistics and source index.\n"
        "- source/: copies of the reproducible builders. Canonical execution is from the project scripts/ directory.\n\n"
        "## Rebuild\n\n```bash\n"
        "/home/l/miniconda3/envs/hybrid_robot/bin/python scripts/build_kbs_submission_assets.py\n```\n\n"
        "No training or simulator process is started. Figures use recorded observations only. "
        "Full Validation and Probe500 remain separate. The selected/explored Test budgets are "
        "GlobalRaw33.5/34k, D-only54/59k, and M7-Raw59/70k. The 59k checkpoint is not claimed "
        "to be a universal optimum. The 14 tracking bodies are distinguished from the 30-body NPZ schema.\n\n"
        "## Recommended manuscript placement\n\n"
        "Figures1-2: Methods. Figure3: data and metadata. Figure4: 34k ablations. Figure5: budget study. "
        "Figures6-8: final Test, paired uncertainty and category behavior. FigureS1: supplementary "
        "checkpoint diagnostics. Main tables:4 (variants),8 (ablations),9 (budget),10 (Test),11 "
        "(uncertainty); other tables may move to the supplement. Author approval and final "
        "journal-specific checks are still required before submission.\n"
    )
    (OUT/"README.md").write_text(readme,encoding="utf-8")
    dump(OUT/"Package_Manifest.json",{"date":DATE,"figures":len(ART),"tables":len(tables),
         "source_group_overlap":audit["cross_split_source_group_overlap"],
         "files":[{"path":str(p.relative_to(OUT)),"bytes":p.stat().st_size,"sha256":data.sha256_file(p)}
                  for p in sorted(OUT.rglob("*")) if p.is_file() and p.name!="Package_Manifest.json"]})
    archive=OUT.with_suffix(".zip")
    with zipfile.ZipFile(archive,"w",compression=zipfile.ZIP_DEFLATED,compresslevel=6) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file(): z.write(p,str(p.relative_to(OUT.parent)))
    print(json.dumps({"package":str(OUT),"archive":str(archive),"figures":len(ART),"tables":len(tables),"audit":audit},indent=2))


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--package-only",action="store_true",help="Refresh source copies, hashes and ZIP after adding reviewed PDF exports.")
    args=parser.parse_args()
    configure()
    if args.package_only:
        FIGURE_AUDIT.extend(data.read_json(OUT/"Artwork_Preflight.json"))
        for item in FIGURE_AUDIT: ART[item["figure"]]={}
        tables={p.stem:None for p in TABLE.glob("Table_*.csv")}
        finish(tables,data.read_json(EVIDENCE/"Evidence_Audit.json"))
        return
    formal,budget,test=data.build_result_sets()
    categories=data.category_comparison(test); probe=complete_probe(); stats=data.load_module_statistics()
    data.validate_inputs(formal,budget,test,categories,probe)
    audit=audit_evidence(formal,budget,test)
    print("Evidence identities and metrics checked.",flush=True)
    tables,analysis=data.build_paper_tables(formal,budget,test,categories,probe,stats)
    dump(EVIDENCE/"Paired_Statistics.json",{k:v for k,v in analysis.items() if k not in ("m7","donly","globalraw")})
    tables=prepare_tables(tables,formal,budget,test)
    export_tables(tables); write_manuscript(tables); write_formulae(); write_source_index(); write_submission_notes(); write_bib(); write_result_inventory()
    print("English text and editable tables exported.",flush=True)
    for func,args in [(plot_framework,()),(plot_sampler,()),(plot_metadata,(stats,)),(plot_ablation,(formal,)),
                      (plot_budget,(budget,)),(plot_test,(test,)),(plot_paired,(analysis,)),
                      (plot_categories,(categories,)),(plot_probe,(probe,)),(graphical_abstract,(test,))]:
        func(*args); print(func.__name__+" exported.",flush=True)
    write_portfolio(); validate_artifacts(); finish(tables,audit)


if __name__=="__main__":
    main()
