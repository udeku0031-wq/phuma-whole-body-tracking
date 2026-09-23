"""Audit the existing KBS package without rerunning experiments or artwork builders."""

import argparse
import csv
import hashlib
import json
import re
import subprocess
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import numpy as np
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "reports/kbs_submission_materials_2026-09-11"
NS = {"h": "http://www.w3.org/1999/xhtml"}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def rows(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.reader(stream))


def repair_pagination():
    path = OUT / "Editable_Paper_Tables.docx"
    doc = Document(path)
    changed = False
    for section in doc.sections:
        if section.bottom_margin != Cm(2.6):
            section.bottom_margin = Cm(2.6)
            changed = True
    if changed:
        doc.save(path)
    print(json.dumps({"table_docx_changed": changed, "next": "Convert only the table DOCX to PDF, then run this audit."}))


def pdf_audit(path):
    fonts = subprocess.check_output(["pdffonts", str(path)], text=True)
    entries = [line.split() for line in fonts.splitlines()[2:] if line.strip()]
    assert entries and all(line[-5] == "yes" for line in entries), path
    assert "Type 3" not in fonts, path
    xml = subprocess.check_output(["pdftotext", "-bbox-layout", str(path), "-"], text=True)
    pages = ET.fromstring(xml).findall(".//h:page", NS)
    page_checks = []
    for number, page in enumerate(pages, 1):
        width, height = float(page.attrib["width"]), float(page.attrib["height"])
        if path.parent.name == "figures":
            assert abs(width * 25.4 / 72 - 190) < .02, path
            if path.stem == "Graphical_Abstract":
                assert abs(width / height - 2.5) < .001, path
        words = page.findall(".//h:word", NS)
        assert words, (path, number, "empty page")
        for word in words:
            x0, y0, x1, y1 = (float(word.attrib[k]) for k in ("xMin", "yMin", "xMax", "yMax"))
            assert x0 >= -0.5 and y0 >= -0.5 and x1 <= width + 0.5 and y1 <= height + 0.5, (path, number, word.text)
        footer = []
        body = []
        for line in page.findall(".//h:line", NS):
            text = " ".join(word.text or "" for word in line.findall("h:word", NS))
            (footer if "PHUMA-WBT |" in text else body).append(line)
        gap = None
        if footer:
            footer_y = min(float(line.attrib["yMin"]) for line in footer)
            body_y = max(float(line.attrib["yMax"]) for line in body)
            gap = round(footer_y - body_y, 3)
            assert gap >= 6, (path, number, "body/footer clearance", gap)
        page_checks.append({"page": number, "words": len(words), "footer_clearance_pt": gap})
    return {"file": str(path.relative_to(OUT)), "pages": len(pages), "fonts_embedded": True,
            "type3_fonts": False, "page_checks": page_checks}, fonts


def table_audit():
    doc = Document(OUT / "Editable_Paper_Tables.docx")
    csvs = sorted((OUT / "tables").glob("Table_*.csv"))
    assert len(doc.tables) == len(csvs) == 21
    checks = []
    for table, path in zip(doc.tables, csvs):
        expected = rows(path)
        actual = [[cell.text for cell in row.cells] for row in table.rows]
        assert actual == expected, path
        assert table.rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None, path
        for row in table.rows:
            assert row._tr.find(".//" + qn("w:cantSplit")) is not None, path
            for cell in row.cells:
                for edge in ("left", "right"):
                    node = cell._tc.find(".//" + qn("w:tcBorders") + "/" + qn("w:" + edge))
                    assert node is not None and node.get(qn("w:val")) == "nil", path
                assert cell._tc.find(".//" + qn("w:shd")) is None, path
        if path.name.startswith(("Table_08_", "Table_09_", "Table_10_")):
            for c, header in enumerate(expected[0]):
                if not header.endswith(("\u2191", "\u2193")):
                    continue
                values = [float(row[c]) for row in expected[1:]]
                best = (max if header.endswith("\u2191") else min)(values)
                for r, value in enumerate(values, 1):
                    assert all(bool(run.bold) == (value == best)
                               for p in table.rows[r].cells[c].paragraphs for run in p.runs), (path, r, c)
        tex = path.with_suffix(".tex").read_text()
        for env in ("table*", "tabularx", "longtable", "minipage"):
            assert tex.count("\\begin{" + env + "}") == tex.count("\\end{" + env + "}"), path
        checks.append({"table": path.stem, "data_rows": len(expected) - 1, "word_matches_csv": True})
    return checks


def test_audit():
    records = {}
    for method in ("GlobalRaw", "D-only", "M7-Raw"):
        folder = OUT / "evidence" / method
        raw = rows(folder / "per_motion.csv")
        values = [dict(zip(raw[0], row)) for row in raw[1:]]
        summary = json.loads((folder / "summary.json").read_text())
        keys = [row["motion_path"] for row in values]
        assert len(keys) == len(set(keys)) == summary["num_motions"] == 7592
        successes = sum(int(row["success"]) for row in values)
        assert successes == summary["num_success"]
        assert 7592 - successes == summary["num_failure"]
        categories = {row["category"] for row in values}
        macro = np.mean([np.mean([int(row["success"]) for row in values if row["category"] == c]) for c in categories])
        assert len(categories) == 17 and abs(macro - summary["macro_success_rate"]) < 1e-6
        for key in ("completion_ratio", "body_position_error_m", "joint_position_error_l2_rad", "joint_position_error_rms_rad"):
            observed = np.mean([float(row[key]) for row in values])
            assert abs(observed - summary["mean_" + key]) < 1e-6, (method, key)
        assert abs(summary["mean_joint_position_error_l2_rad"] / np.sqrt(29)
                   - summary["mean_joint_position_error_rms_rad"]) < 1e-6
        checkpoint = Path(summary["checkpoint_path"])
        assert checkpoint.is_file() and sha(checkpoint) == summary["checkpoint_sha256"]
        records[method] = {row["motion_path"]: row for row in values}
    assert records["GlobalRaw"].keys() == records["D-only"].keys() == records["M7-Raw"].keys()
    paired = {}
    for method in ("GlobalRaw", "D-only"):
        counts = {"both": 0, "m7_only": 0, "baseline_only": 0, "neither": 0}
        for key, row in records["M7-Raw"].items():
            a, b = int(row["success"]), int(records[method][key]["success"])
            counts[{(1, 1): "both", (1, 0): "m7_only", (0, 1): "baseline_only", (0, 0): "neither"}[(a, b)]] += 1
        paired[method] = counts
    assert paired["D-only"] == {"both": 6711, "m7_only": 297, "baseline_only": 150, "neither": 434}
    return {"motions_per_policy": 7592, "matched_motion_sets": True, "checkpoint_hashes": "passed", "paired_outcomes": paired}


def package_check():
    manifest = json.loads((OUT / "Package_Manifest.json").read_text())
    expected = {r["path"] for r in manifest["files"]}
    actual = {str(p.relative_to(OUT)) for p in OUT.rglob("*") if p.is_file() and p.name != "Package_Manifest.json"}
    assert expected == actual, (expected - actual, actual - expected)
    for record in manifest["files"]:
        path = OUT / record["path"]
        assert path.stat().st_size == record["bytes"] and sha(path) == record["sha256"], path
    with zipfile.ZipFile(OUT.with_suffix(".zip")) as archive:
        assert archive.testzip() is None
        paths = [p for p in OUT.rglob("*") if p.is_file()]
        assert len(archive.namelist()) == len(paths)
        for path in paths:
            assert archive.read(str(path.relative_to(OUT.parent))) == path.read_bytes(), path
    print(json.dumps({"manifest_files": len(expected), "archive_matches_current_files": True,
                      "archive_sha256": sha(OUT.with_suffix(".zip"))}, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--repair-table-pagination", action="store_true")
    parser.add_argument("--package-check", action="store_true")
    args = parser.parse_args()
    if args.repair_table_pagination:
        repair_pagination()
        return
    if args.package_check:
        package_check()
        return
    result = {"date": "2026-09-11", "training_or_evaluation_launched": False}
    result["tables"] = table_audit()
    result["final_test"] = test_audit()
    pdfs, fonts = [], []
    for path in sorted(OUT.glob("*.pdf")) + sorted((OUT / "figures").glob("*.pdf")):
        check, report = pdf_audit(path)
        pdfs.append(check)
        fonts.append(str(path.relative_to(OUT)) + "\n" + report)
    result["pdfs"] = pdfs
    figures = []
    for path in sorted((OUT / "figures").glob("*.pdf")):
        for ext in (".eps", ".tiff", ".png", ".svg"):
            assert path.with_suffix(ext).is_file(), path
        eps = path.with_suffix(".eps").read_text()
        assert re.search(r"/FontType\s+42\b", eps) and re.search(r"/sfnts\s*\[", eps), path
        assert "/FontType 3 " not in eps, path
        subprocess.run(["gs", "-q", "-dSAFER", "-dBATCH", "-dNOPAUSE", "-sDEVICE=nullpage",
                        str(path.with_suffix(".eps"))], check=True, capture_output=True)
        with Image.open(path.with_suffix(".tiff")) as image:
            assert image.mode == "RGB" and image.width >= 7480
            assert all(abs(float(dpi) - 1000) < .1 for dpi in image.info["dpi"])
            assert image.info["compression"] == "tiff_lzw"
            pixels = image.size
        with Image.open(path.with_suffix(".png")) as image:
            fraction = float(np.mean(np.min(np.asarray(image.convert("RGB")), axis=2) < 230))
            assert .005 < fraction < .6, path
        figures.append({"figure": path.stem, "tiff_pixels": pixels, "dpi": 1000, "mode": "RGB",
                        "eps_truetype_embedded": True, "eps_ghostscript_validation": "passed"})
    assert len(figures) == 10
    result["figures"] = figures
    text = (OUT / "PHUMA_WBT_KBS_English_Materials.md").read_text()
    assert not re.search("[\u4e00-\u9fff]", text)
    abstract = text.split("## Abstract\n\n", 1)[1].split("\n\n## ", 1)[0]
    highlights = (OUT / "Highlights.txt").read_text().splitlines()
    assert 3 <= len(highlights) <= 5 and all(len(line) <= 85 for line in highlights)
    result["abstract_word_count"] = len(abstract.split())
    result["highlight_character_counts"] = list(map(len, highlights))
    result["latex"] = "Environment balance checked; compilation not performed (no TeX engine)."
    result["layout_scope"] = "PDF word bounds and body/footer clearance; visual review recorded separately in FINAL_HANDOFF.md."
    (OUT / "Final_QA.json").write_text(json.dumps(result, indent=2) + "\n")
    (OUT / "Final_PDF_Fonts.txt").write_text("\n".join(fonts))
    print(json.dumps({"pdfs": len(pdfs), "figures": len(figures), "tables": len(result["tables"]),
                      "abstract_words": result["abstract_word_count"], "status": "passed"}, indent=2))


if __name__ == "__main__":
    main()
